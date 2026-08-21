"""Print a concise structural inspection of an OAS-K config workbook."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook  # noqa: E402

from shared.database.exporting import (  # noqa: E402
    validate_configuration_workbook,
)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Inspect an OAS-K unified configuration workbook."
    )
    parser.add_argument("workbook", type=Path)
    arguments = parser.parse_args()

    result = validate_configuration_workbook(arguments.workbook)
    details: dict[str, object] = {
        "path": str(result.path),
        "is_valid": result.is_valid,
        "sheet_names": result.sheet_names,
        "formula_count": result.formula_count,
        "external_link_count": result.external_link_count,
        "has_macros": result.has_macros,
        "errors": result.errors,
        "warnings": result.warnings,
        "sheets": {},
    }
    if result.path.is_file():
        workbook = load_workbook(
            result.path,
            read_only=False,
            keep_links=False,
        )
        try:
            details["sheets"] = {
                sheet.title: {
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "freeze_panes": str(sheet.freeze_panes or ""),
                    "auto_filter": sheet.auto_filter.ref,
                    "column_widths": {
                        column: dimension.width
                        for column, dimension in sheet.column_dimensions.items()
                        if dimension.width is not None
                    },
                    "validation_ranges": [
                        str(validation.sqref)
                        for validation in sheet.data_validations.dataValidation
                    ],
                    "multiline_cells": [
                        cell.coordinate
                        for row in sheet.iter_rows()
                        for cell in row
                        if isinstance(cell.value, str) and "\n" in cell.value
                    ],
                    "text_formatted_cells": [
                        cell.coordinate
                        for row in sheet.iter_rows()
                        for cell in row
                        if cell.number_format == "@"
                    ][:20],
                }
                for sheet in workbook.worksheets
            }
        finally:
            workbook.close()

    print(json.dumps(details, ensure_ascii=False, indent=2))
    return 0 if result.is_valid else 1


if __name__ == "__main__":
    raise SystemExit(main())
