"""UI6 adapter around the existing HRIS reader, planner, and full engine."""

from __future__ import annotations

import json
import tempfile
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

from hris.click_profile import HRISClickProfileManager
from hris.engine import HRISFullUploadEngine
from hris.assisted_calibrator import HRISAssistedCalibrator
from hris.job_manager import HRISUploadJobManager
from shared.config_manager import HRISConfigurationReader
from shared.database.exporting import export_hris_legacy
from shared.storage.recorder_profile_manager import (
    resolve_profile_path,
    validate_profile_reference,
)
from ui.hris_models import (
    HRISCancelledError,
    HRISCancellationToken,
    HRISFileResult,
    HRISInteractionGate,
    HRISInterventionRequest,
    HRISJobState,
    HRISLogEvent,
    HRISProgressEvent,
    HRISRecorderProfile,
    HRISResolvedRequest,
    HRISRunControlMapping,
    HRISRunResult,
    HRISValidationResult,
)


class HRISAdapter:
    def __init__(
        self,
        *,
        reader_class=HRISConfigurationReader,
        engine_class=HRISFullUploadEngine,
        planner_class=HRISUploadJobManager,
    ) -> None:
        self.reader_class = reader_class
        self.engine_class = engine_class
        self.planner_class = planner_class

    def validate(self, request: HRISResolvedRequest) -> HRISValidationResult:
        errors: list[str] = []
        warnings: list[str] = []
        profile = self._validate_profile(request, errors, warnings)
        mappings: tuple[HRISRunControlMapping, ...] = ()
        try:
            with self.configuration_path(request) as path:
                configuration = self.reader_class(path).read()
                profile = self._validate_runtime_profile(
                    request,
                    configuration,
                    profile,
                    errors,
                )
                plan = self.planner_class().create_upload_plan(
                    configuration,
                    request.source_folder,
                    self._engine_workflow(request.workflow),
                )
            if not plan.is_valid:
                errors.append(plan.message)
            else:
                mappings = tuple(
                    HRISRunControlMapping(
                        item.sequence,
                        Path(item.txt_file_path),
                        str(item.run_control_id),
                        request.workflow,
                    )
                    for item in plan.plan_items
                )
                ids = [item.run_control_id for item in mappings]
                if len(ids) != len(set(ids)):
                    errors.append("Run Control mapping duplikat tidak diizinkan.")
        except Exception as exc:
            errors.append(str(exc))
        return HRISValidationResult(
            valid=not errors,
            browser_allowed=not errors,
            workflow=request.workflow,
            period_start=request.period_start,
            period_end=request.period_end,
            txt_count=len(mappings),
            mappings=mappings,
            profile=profile,
            warnings=tuple(warnings),
            errors=tuple(errors),
        )

    def calibrate_profile(
        self,
        database_path: Path,
        target_profile: Path,
        *,
        instruction_callback: Callable[[str], None],
        coordinate_callback: Callable[[str], tuple[int, int]],
    ) -> Path:
        """Export HRIS DB config temporarily and run assisted calibration."""

        with tempfile.TemporaryDirectory(prefix="oask-hris-calibration-") as temp:
            configuration_path = Path(temp) / "HRIS_Configuration.xlsx"
            export_hris_legacy(
                database_path,
                configuration_path,
                overwrite=False,
                create_parent=False,
            )
            self._set_legacy_click_profile_path(
                configuration_path,
                target_profile,
            )
            return HRISAssistedCalibrator(
                configuration_path,
                instruction_callback=instruction_callback,
                coordinate_callback=coordinate_callback,
            ).run()

    def run(
        self,
        request: HRISResolvedRequest,
        *,
        cancellation: HRISCancellationToken,
        gate: HRISInteractionGate,
        progress: Callable[[HRISProgressEvent], None],
        log: Callable[[HRISLogEvent], None],
        intervention: Callable[[HRISInterventionRequest], None],
    ) -> HRISRunResult:
        started = self._now()
        total = len(tuple(request.source_folder.glob("*.txt")))

        def emit(state: HRISJobState, message: str, completed: int = 0) -> None:
            progress(HRISProgressEvent(state, message, completed, total))
            log(HRISLogEvent(self._now(), "INFO", state, message))

        def wait_login() -> None:
            emit(HRISJobState.WAITING_FOR_LOGIN, "Menunggu login manual HRIS.")
            request_value = HRISInterventionRequest(
                request.job_id,
                HRISJobState.WAITING_FOR_LOGIN,
                "Silakan login pada browser HRIS, lalu kembali ke OAS-K.",
            )
            gate.current_intervention = request_value
            intervention(request_value)
            gate.wait(gate.login_event, cancellation)
            emit(HRISJobState.NAVIGATING, "Login dikonfirmasi; melanjutkan navigasi.")

        def wait_checkpoint(message: str) -> None:
            intervention_value = HRISInterventionRequest(
                request.job_id,
                HRISJobState.WAITING_FOR_CHECKPOINT,
                message,
                actions=("continue", "stop"),
            )
            gate.prepare_decision(intervention_value)
            emit(
                HRISJobState.WAITING_FOR_CHECKPOINT,
                "Automation HRIS memerlukan checkpoint operator.",
            )
            intervention(intervention_value)
            gate.wait(gate.decision_event, cancellation)
            if gate.decision_action != "continue":
                raise RuntimeError("HRIS dihentikan oleh operator pada checkpoint.")
            emit(HRISJobState.RESUMING, "Checkpoint dikonfirmasi; automation dilanjutkan.")

        def wait_verification(message: str) -> str:
            intervention_value = HRISInterventionRequest(
                request.job_id,
                HRISJobState.WAITING_FOR_VERIFICATION,
                message,
                actions=("submitted", "failed", "retry", "stop"),
            )
            gate.prepare_decision(intervention_value)
            emit(
                HRISJobState.WAITING_FOR_VERIFICATION,
                "Hasil submission HRIS memerlukan verifikasi operator.",
            )
            intervention(intervention_value)
            gate.wait(gate.decision_event, cancellation)
            action = gate.decision_action or "stop"
            emit(HRISJobState.RESUMING, f"Verifikasi operator: {action}.")
            return action

        emit(HRISJobState.STARTING_BROWSER, "Membuka Microsoft Edge untuk HRIS.")
        try:
            with self.configuration_path(request) as configuration_path:
                result = self.engine_class(
                    configuration_file=configuration_path,
                    txt_folder=request.source_folder,
                    output_root=request.output_root,
                    workflow=self._engine_workflow(request.workflow),
                    start_date=self._hris_date(request.period_start),
                    end_date=self._hris_date(request.period_end),
                    wait_for_manual_login=request.manual_login,
                    manual_login_callback=wait_login,
                    manual_checkpoint_callback=wait_checkpoint,
                    hris_username=request.login_id,
                    hris_password=request.login_secret,
                    manual_verification_callback=wait_verification,
                    profile_path_override=self._absolute_profile_path(request),
                    move_failed_files=False,
                ).run()
        except HRISCancelledError as exc:
            emit(HRISJobState.CANCELLED, str(exc))
            return self._cancelled(request, started, str(exc))
        except Exception as exc:
            if cancellation.requested:
                return self._cancelled(request, started, str(exc))
            log(HRISLogEvent(self._now(), "ERROR", HRISJobState.FAILED, str(exc)))
            return HRISRunResult(
                False,
                False,
                request.job_id,
                request.workflow,
                started,
                self._now(),
                (),
                error_summary=str(exc),
            )
        files = self._file_results(request, result)
        cancelled = cancellation.requested
        state = (
            HRISJobState.CANCELLED
            if cancelled
            else HRISJobState.COMPLETED
            if result.success
            else HRISJobState.FAILED
        )
        emit(state, str(result.message), sum(item.status == "SUCCESS" for item in files))
        return HRISRunResult(
            bool(result.success and not cancelled),
            cancelled,
            request.job_id,
            request.workflow,
            started,
            self._now(),
            files,
            self._existing(result.report_file),
            self._existing(result.summary_json_file),
            self._existing(result.process_log_file),
            self._existing(result.report_folder),
            None if result.success else str(result.message),
        )

    @contextmanager
    def configuration_path(self, request: HRISResolvedRequest) -> Iterator[Path]:
        if request.configuration_path is not None:
            yield request.configuration_path
            return
        with tempfile.TemporaryDirectory(prefix="oask-hris-config-") as temp:
            path = Path(temp) / "HRIS_Configuration.xlsx"
            export_hris_legacy(
                request.database_path,
                path,
                overwrite=False,
                create_parent=False,
            )
            yield path

    @staticmethod
    def _set_legacy_click_profile_path(
        configuration_path: Path,
        target_profile: Path,
    ) -> None:
        from openpyxl import load_workbook

        workbook = load_workbook(configuration_path)
        try:
            sheet = workbook["Upload"]
            for row in sheet.iter_rows(min_row=2):
                if str(row[0].value or "").strip() == "Click_Profile_Path":
                    row[1].value = str(target_profile)
                    break
            else:
                sheet.append(("Click_Profile_Path", str(target_profile), ""))
            workbook.save(configuration_path)
        finally:
            workbook.close()

    def _validate_profile(self, request, errors, warnings):
        result = validate_profile_reference(request.data_root, request.recorder_profile)
        profile_errors: list[str] = list(result.errors)
        errors.extend(result.errors)
        warnings.extend(result.warnings)
        reference = result.reference
        if reference is None:
            return None
        if not result.exists:
            errors.append(
                "Recorder profile HRIS tidak ditemukan. Jalankan Calibrate Click Profile dulu."
            )
            return HRISRecorderProfile(
                reference.absolute_path.stem,
                reference.relative_path,
                reference.absolute_path,
                None,
                False,
                False,
                None,
                "ERROR",
                tuple(result.errors),
            )
        workflow = None
        if result.json_object_valid:
            try:
                payload = HRISClickProfileManager.load_profile(reference.absolute_path)
                workflow_value = payload.get("workflow")
                workflow = str(workflow_value).upper() if workflow_value else None
                if workflow and workflow != request.workflow:
                    message = "Recorder profile workflow tidak sesuai dengan job."
                    profile_errors.append(message)
                    errors.append(message)
            except (OSError, ValueError, json.JSONDecodeError) as exc:
                profile_errors.append(str(exc))
                errors.append(str(exc))
        return HRISRecorderProfile(
            reference.absolute_path.stem,
            reference.relative_path,
            reference.absolute_path,
            workflow,
            result.exists,
            result.json_object_valid,
            reference.absolute_path.stat().st_mtime if result.exists else None,
            "READY" if result.json_object_valid and not profile_errors else "ERROR",
            tuple(profile_errors),
        )

    def _validate_runtime_profile(
        self,
        request,
        configuration,
        profile: HRISRecorderProfile | None,
        errors,
    ):
        if profile is None or not profile.exists or not profile.json_valid:
            return profile
        profile_errors = list(profile.errors)
        try:
            payload = HRISClickProfileManager.load_profile(profile.absolute_path)
            validation = HRISClickProfileManager.validate_profile(
                payload,
                self._screen_size(),
                configuration.upload,
                current_scale_percent=self._display_scale_percent(),
                required_steps=configuration.assisted_steps,
            )
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            message = str(exc)
            profile_errors.append(message)
            errors.append(message)
            return replace(profile, validation_status="ERROR", errors=tuple(profile_errors))
        if validation.valid:
            return profile
        message = (
            "Recorder profile HRIS tidak cocok dengan runtime: "
            f"{validation.message}"
        )
        profile_errors.append(message)
        errors.append(message)
        return replace(profile, validation_status="ERROR", errors=tuple(profile_errors))

    @staticmethod
    def _absolute_profile_path(request: HRISResolvedRequest) -> Path:
        if request.recorder_profile.is_absolute():
            return request.recorder_profile
        return resolve_profile_path(request.data_root, request.recorder_profile)

    @staticmethod
    def _screen_size() -> tuple[int, int]:
        import pyautogui

        screen = pyautogui.size()
        return int(screen.width), int(screen.height)

    @staticmethod
    def _display_scale_percent() -> int:
        try:
            import ctypes

            return round(ctypes.windll.shcore.GetScaleFactorForDevice(0))
        except Exception:
            return 100

    @staticmethod
    def _file_results(request, result) -> tuple[HRISFileResult, ...]:
        summary_path = HRISAdapter._existing(result.summary_json_file)
        values = []
        if summary_path:
            try:
                payload = json.loads(summary_path.read_text(encoding="utf-8"))
                values = payload.get("items") or payload.get("plan_items") or []
            except (OSError, json.JSONDecodeError):
                values = []
        return tuple(
            HRISFileResult(
                request.source_folder / str(item.get("txt_file_name", "")),
                str(item.get("run_control_id", "")),
                str(item.get("status", "UNKNOWN")),
                str(item.get("message", "")),
                Path(item["moved_to"])
                if item.get("moved_to")
                else None,
            )
            for item in values
            if item.get("txt_file_name")
        )

    @staticmethod
    def _engine_workflow(workflow: str) -> str:
        return "Branch" if workflow == "BRANCH" else "HO"

    @staticmethod
    def _hris_date(value: str) -> str:
        return datetime.strptime(value, "%Y-%m-%d").strftime("%m/%d/%Y")

    @staticmethod
    def _existing(value) -> Path | None:
        path = Path(value) if value else None
        return path if path and path.exists() else None

    @classmethod
    def _cancelled(cls, request, started, message) -> HRISRunResult:
        return HRISRunResult(
            False,
            True,
            request.job_id,
            request.workflow,
            started,
            cls._now(),
            (),
            error_summary=message,
        )

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).isoformat(timespec="seconds")
