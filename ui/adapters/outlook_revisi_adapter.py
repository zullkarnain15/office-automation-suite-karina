"""Typed adapter for the legacy Outlook Revisi reader, COM client, and engine."""

from __future__ import annotations

import re
import tempfile
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from outlook.downloader import OutlookComClient
from outlook.engine import OutlookRevisiEngine
from shared.config_manager import OutlookRevisiConfigurationReader
from shared.database.exporting import export_outlook_legacy
from shared.hris_txt_staging import stage_hris_txt_files
from ui.outlook_revisi_models import (
    TARGET_FOLDER,
    TARGET_MAILBOX_SMTP,
    MailboxValidationResult,
    OutboundSafetyState,
    OutlookRevisiCancellationToken,
    OutlookRevisiLogEvent,
    OutlookRevisiOutputFile,
    OutlookRevisiProgressEvent,
    OutlookRevisiResolvedRequest,
    OutlookRevisiRunResult,
    OutlookRevisiValidationResult,
)

ProgressCallback = Callable[[OutlookRevisiProgressEvent], None]
LogCallback = Callable[[OutlookRevisiLogEvent], None]

_KNOWN_PLACEHOLDERS = {
    "ANOMALY_ROW_COUNT",
    "ERROR_REASON",
    "EXPECTED_SUBJECT",
    "FAILED_EMAIL",
    "ORIGINAL_SUBJECT",
    "OUTPUT_FOLDER",
    "OUTPUT_TXT_COUNT",
    "PERIOD",
    "REQUIRED_CC_EMAIL",
    "RESUBMIT_DEADLINE",
    "SENDER_NAME",
    "SUCCESS_EMAIL",
    "TOTAL_EMAIL",
    "VALID_ROW_COUNT",
}


class OutlookRevisiAdapter:
    def __init__(
        self,
        reader_class=OutlookRevisiConfigurationReader,
        engine_class=OutlookRevisiEngine,
        client_factory=None,
    ) -> None:
        self.reader_class = reader_class
        self.engine_class = engine_class
        self.client_factory = client_factory or self._build_client

    def validate(
        self,
        request: OutlookRevisiResolvedRequest,
        *,
        check_mailbox: bool,
    ) -> OutlookRevisiValidationResult:
        errors: list[str] = []
        warnings: list[str] = []
        try:
            with self._configuration_path(request) as path:
                configuration = self.reader_class(path).read()
        except Exception as exc:
            return self._invalid_validation(
                request,
                [f"Configuration tidak valid: {exc}"],
            )

        general = configuration.general
        configured_mailbox = str(general.get("Mailbox_SMTP", "")).strip()
        configured_folder = str(general.get("Source_Folder", "Inbox")).strip()
        if configured_mailbox.casefold() != TARGET_MAILBOX_SMTP:
            errors.append(
                "Mailbox target wajib karina.hr.1@oto.co.id; "
                f"configuration berisi {configured_mailbox or '(empty)'}."
            )
        if configured_folder.casefold() != TARGET_FOLDER.casefold():
            errors.append(
                f"Source folder wajib Inbox; configuration berisi "
                f"{configured_folder or '(empty)'}."
            )

        workflow = "Branch" if request.workflow == "BRANCH" else "HO"
        senders = (
            configuration.branch_senders
            if workflow == "Branch"
            else configuration.ho_senders
        )
        subject_rules = self._workflow_items(configuration.subject_rules, workflow)
        attachment_rules = self._workflow_items(
            configuration.attachment_rules, workflow
        )
        validation_rules = self._workflow_items(
            configuration.validation_rules, workflow
        )
        if not senders:
            errors.append(
                f"Sender master aktif untuk {request.workflow} tidak tersedia."
            )
        if not subject_rules:
            errors.append(
                f"Subject rule aktif untuk {request.workflow} tidak tersedia."
            )
        if not attachment_rules:
            errors.append(
                f"Attachment rule aktif untuk {request.workflow} tidak tersedia."
            )

        template_previews = tuple(
            (item.reply_code, item.subject_template, item.body_template)
            for item in configuration.reply_templates
        )
        unknown = sorted(
            {
                placeholder
                for _, subject, body in template_previews
                for placeholder in re.findall(r"\{([A-Za-z0-9_]+)\}", subject + body)
                if placeholder.upper() not in _KNOWN_PLACEHOLDERS
            }
        )
        if unknown:
            warnings.append("Unknown template placeholder: " + ", ".join(unknown))

        outbound = self._outbound_safety(configuration, request.dry_run, senders)
        warnings.extend(outbound.warnings)
        if (
            not request.dry_run
            and str(general.get("Send_Transport", "OUTLOOK")).upper() == "SMTP"
            and outbound.send_mode == "DRAFT"
        ):
            errors.append("SMTP transport tidak mendukung Send_Mode DRAFT.")

        mailbox = MailboxValidationResult(
            valid=not errors and not check_mailbox,
            configured_mailbox=configured_mailbox,
            configured_folder=configured_folder,
            outlook_status="NOT_CHECKED",
            mailbox_found=False,
            folder_found=False,
        )
        configuration_valid = not errors
        if check_mailbox and configuration_valid:
            mailbox = self._validate_mailbox(configuration)
            if not mailbox.valid:
                errors.append(mailbox.error or "Mailbox validation gagal.")

        return OutlookRevisiValidationResult(
            valid=not errors,
            configuration_valid=configuration_valid,
            workflow=request.workflow,
            period_start=request.period_start,
            period_end=request.period_end,
            output_root=request.output_root,
            mailbox=mailbox,
            outbound=outbound,
            sender_count=len(senders),
            subject_rule_count=len(subject_rules),
            attachment_rule_count=len(attachment_rules),
            validation_rule_count=len(validation_rules),
            reply_template_count=len(configuration.reply_templates),
            summary_recipient_count=len(configuration.get_pic_hr_emails())
            + len(configuration.get_spv_pic_hr_emails()),
            template_previews=template_previews,
            warnings=tuple(warnings),
            errors=tuple(errors),
        )

    def run(
        self,
        request: OutlookRevisiResolvedRequest,
        *,
        cancellation: OutlookRevisiCancellationToken,
        progress: ProgressCallback,
        log: LogCallback,
    ) -> OutlookRevisiRunResult:
        started = self._now()
        if cancellation.requested:
            return self._cancelled(
                request, started, "Dibatalkan sebelum Outlook connect."
            )
        progress(
            OutlookRevisiProgressEvent(
                "VALIDATING_CONFIGURATION", "Validating Configuration"
            )
        )
        try:
            with self._configuration_path(request) as configuration_path:
                log(
                    self._log(
                        "INFO",
                        f"Reading {request.configuration_source} configuration.",
                    )
                )
                configuration = self.reader_class(configuration_path).read()
                self._require_target_mailbox(configuration)
                client = self.client_factory(configuration)
                with self._runtime_configuration(
                    request, configuration_path
                ) as runtime_path:
                    engine = self.engine_class(
                        configuration_file=runtime_path,
                        workflow="Branch" if request.workflow == "BRANCH" else "HO",
                        dry_run=request.dry_run,
                        message_limit=request.message_limit,
                        client=client,
                        progress_callback=lambda stage, message: progress(
                            OutlookRevisiProgressEvent(stage, message)
                        ),
                        cancellation_requested=lambda: cancellation.requested,
                    )
                    log(
                        self._log(
                            "INFO",
                            f"Running Outlook Revisi for {request.workflow} in "
                            f"{'PREVIEW' if request.dry_run else 'LIVE'} mode.",
                        )
                    )
                    with self._com_initialized():
                        engine_result = engine.run()
            return self._normalize_result(request, engine_result, started, log)
        except Exception as exc:
            log(self._log("ERROR", f"Outlook Revisi run failed: {exc}"))
            return OutlookRevisiRunResult(
                success=False,
                cancelled=cancellation.requested,
                job_id=request.job_id,
                workflow=request.workflow,
                mailbox=TARGET_MAILBOX_SMTP,
                started_at=started,
                ended_at=self._now(),
                output_root=request.output_root,
                job_folder=None,
                message_counts={},
                attachment_counts={},
                reply_counts={},
                output_files=(),
                warning_count=int(cancellation.requested),
                error_summary=str(exc),
            )

    def _validate_mailbox(self, configuration) -> MailboxValidationResult:
        mailbox = str(configuration.general.get("Mailbox_SMTP", "")).strip()
        folder = str(configuration.general.get("Source_Folder", "Inbox")).strip()
        try:
            client = self.client_factory(configuration)
            with self._com_initialized():
                display, folder_name = client.validate_mailbox()
            return MailboxValidationResult(
                valid=True,
                configured_mailbox=mailbox,
                configured_folder=folder,
                outlook_status="READY",
                mailbox_found=True,
                folder_found=True,
                mailbox_display=display,
            )
        except Exception as exc:
            return MailboxValidationResult(
                valid=False,
                configured_mailbox=mailbox,
                configured_folder=folder,
                outlook_status="UNAVAILABLE",
                mailbox_found=False,
                folder_found=False,
                error=str(exc),
            )

    @staticmethod
    def _outbound_safety(
        configuration, dry_run: bool, workflow_senders
    ) -> OutboundSafetyState:
        general = configuration.general
        auto_reply = OutlookRevisiAdapter._to_bool(general.get("Auto_Reply_Enabled"))
        send_mode = str(general.get("Send_Mode", "SEND")).strip().upper()
        recipients = tuple(
            dict.fromkeys(
                configuration.get_pic_hr_emails()
                + configuration.get_spv_pic_hr_emails()
                + [item.sender_email for item in workflow_senders]
            )
        )
        summary_enabled = any(
            item.reply_code == "SUMMARY_PIC" for item in configuration.reply_templates
        ) and bool(configuration.get_pic_hr_emails())
        warnings: list[str] = []
        if auto_reply:
            warnings.append("Auto Reply aktif pada configuration.")
        if send_mode == "SEND":
            warnings.append("Send Mode adalah SEND.")
        effective = "PREVIEW" if dry_run else send_mode
        return OutboundSafetyState(
            auto_reply_enabled=auto_reply,
            send_mode=send_mode,
            effective_mode=effective,
            reply_enabled=auto_reply,
            summary_email_enabled=summary_enabled,
            recipients=recipients,
            requires_confirmation=(not dry_run and (auto_reply or send_mode == "SEND")),
            warnings=tuple(warnings),
        )

    def _normalize_result(self, request, result, started, log):
        cancelled = bool(getattr(result, "cancelled", False))
        outputs = self._outputs(result)
        staging_warnings = 0
        if bool(result.success) and not cancelled:
            staging_warnings = self._stage_hris_txt(request, outputs, log)
        messages = list(getattr(result, "message_results", ()))
        attachment_results = [
            attachment
            for message in messages
            for attachment in getattr(message, "attachment_results", ())
        ]
        reply_counts = {
            "sent": sum(item.reply_result == "SENT" for item in messages),
            "drafted": sum(item.reply_result == "DRAFTED" for item in messages),
            "failed": sum(item.reply_result == "FAILED" for item in messages),
        }
        warnings = len(getattr(result, "reconciliation_issues", ())) + int(
            getattr(result, "anomaly_row_count", 0)
        ) + staging_warnings
        if cancelled:
            warnings += 1
        error_summary = None
        if cancelled:
            error_summary = (
                "Pembatalan diterapkan pada safe checkpoint; partial output "
                "dipertahankan."
            )
        elif not result.success:
            error_summary = f"{int(result.failed_email)} email gagal diproses."
        log(
            self._log(
                "WARNING" if cancelled or warnings else "INFO",
                "Outlook Revisi cancelled at a safe checkpoint."
                if cancelled
                else "Outlook Revisi engine completed.",
            )
        )
        attachment_folder = self._existing_path(
            Path(result.output_folder) / "Attachments"
        )
        return OutlookRevisiRunResult(
            success=bool(result.success) and not cancelled,
            cancelled=cancelled,
            job_id=request.job_id,
            workflow=request.workflow,
            mailbox=TARGET_MAILBOX_SMTP,
            started_at=started,
            ended_at=self._now(),
            output_root=request.output_root,
            job_folder=self._existing_path(result.output_folder),
            message_counts={
                "total": int(result.total_email),
                "target": int(result.target_email),
                "success": int(result.success_email),
                "failed": int(result.failed_email),
                "skipped": int(result.skipped_other_workflow),
            },
            attachment_counts={
                "total": sum(
                    int(getattr(item, "attachment_count", 0)) for item in messages
                ),
                "accepted": sum(
                    item.file_status == "ACCEPTED" for item in attachment_results
                ),
                "rejected": sum(
                    item.file_status != "ACCEPTED" for item in attachment_results
                ),
            },
            reply_counts=reply_counts,
            output_files=outputs,
            warning_count=warnings,
            error_summary=error_summary,
            process_log_path=self._existing_path(result.process_log),
            summary_json_path=self._existing_path(result.summary_json),
            attachment_folder=attachment_folder,
        )

    @classmethod
    def _outputs(cls, result) -> tuple[OutlookRevisiOutputFile, ...]:
        candidates: list[tuple[str, object]] = [
            ("OUTPUT_FOLDER", getattr(result, "output_folder", None)),
            (
                "ATTACHMENT_FOLDER",
                Path(result.output_folder) / "Attachments"
                if getattr(result, "output_folder", None)
                else None,
            ),
            ("EXCEL_REPORT", getattr(result, "report_file", None)),
            ("PROCESS_LOG", getattr(result, "process_log", None)),
            ("SUMMARY_JSON", getattr(result, "summary_json", None)),
        ]
        candidates.extend(
            ("HRIS_TXT", path)
            for message in getattr(result, "message_results", ())
            for path in getattr(message, "output_files", ())
        )
        seen: set[tuple[str, Path]] = set()
        outputs: list[OutlookRevisiOutputFile] = []
        for role, value in candidates:
            path = cls._existing_path(value)
            key = (role, path) if path is not None else None
            if path is not None and key not in seen:
                seen.add(key)
                outputs.append(OutlookRevisiOutputFile(role, path))
        return tuple(outputs)

    @classmethod
    def _stage_hris_txt(cls, request, outputs, log) -> int:
        txt_files = tuple(
            item.path for item in outputs if item.file_type == "HRIS_TXT"
        )
        if not txt_files:
            return 0
        try:
            staged = stage_hris_txt_files(
                txt_files,
                request.output_root,
                request.workflow,
            )
        except Exception as exc:
            log(cls._log("ERROR", f"HRIS staging gagal: {exc}"))
            return 1
        copied = sum(item.status != "SKIPPED_IDENTICAL" for item in staged)
        skipped = len(staged) - copied
        log(
            cls._log(
                "INFO",
                f"HRIS staging selesai: {copied} TXT disalin, "
                f"{skipped} duplikat identik dilewati.",
            )
        )
        return 0

    @contextmanager
    def _runtime_configuration(
        self,
        request: OutlookRevisiResolvedRequest,
        source_path: Path,
    ) -> Iterator[Path]:
        with tempfile.TemporaryDirectory(prefix="oask-outlook-runtime-") as temp:
            runtime_path = Path(temp) / source_path.name
            book = load_workbook(source_path)
            try:
                sheet = book["General"]
                self._set_general_value(sheet, "Output_Root", str(request.output_root))
                self._set_general_value(sheet, "Payroll_Period", request.payroll_period)
                book.save(runtime_path)
            finally:
                book.close()
            yield runtime_path

    @staticmethod
    @contextmanager
    def _configuration_path(request) -> Iterator[Path]:
        if request.configuration_path is not None:
            path = request.configuration_path
            if path.suffix.casefold() != ".xlsx" or not path.is_file():
                raise ValueError("Outlook Revisi Configuration .xlsx tidak ditemukan.")
            yield path
            return
        if request.database_path is None:
            raise RuntimeError("Database konfigurasi aktif tidak tersedia.")
        with tempfile.TemporaryDirectory(prefix="oask-outlook-config-") as temp:
            path = Path(temp) / "Outlook_Revisi_Configuration.xlsx"
            export_outlook_legacy(
                request.database_path,
                path,
                overwrite=False,
                create_parent=False,
            )
            yield path

    @staticmethod
    def _set_general_value(sheet, key: str, value: str) -> None:
        for row in sheet.iter_rows():
            if str(row[0].value or "").strip() == key:
                row[1].value = value
                return
        sheet.append([key, value, "Unified UI runtime override"])

    @staticmethod
    def _workflow_items(items, workflow: str):
        return [item for item in items if item.workflow in {workflow, "All"}]

    @staticmethod
    def _require_target_mailbox(configuration) -> None:
        mailbox = str(configuration.general.get("Mailbox_SMTP", "")).strip()
        folder = str(configuration.general.get("Source_Folder", "Inbox")).strip()
        if mailbox.casefold() != TARGET_MAILBOX_SMTP or folder.casefold() != "inbox":
            raise ValueError(
                "Configuration mailbox/folder tidak sesuai locked contract."
            )

    @staticmethod
    def _build_client(configuration):
        general = configuration.general
        return OutlookComClient(
            mailbox_smtp=str(general.get("Mailbox_SMTP", "")),
            source_folder=str(general.get("Source_Folder", "Inbox")),
            reply_from_smtp=str(
                general.get("Reply_From_SMTP", "") or general.get("Mailbox_SMTP", "")
            ),
            send_transport=str(general.get("Send_Transport", "OUTLOOK")),
            smtp_server=str(general.get("SMTP_Server", "")),
            smtp_port=int(general.get("SMTP_Port") or 25),
            smtp_timeout=int(general.get("SMTP_Timeout_Seconds") or 30),
            save_smtp_copy_to_sent=OutlookRevisiAdapter._to_bool_default_true(
                general.get("Save_SMTP_Copy_To_Sent")
            ),
        )

    @staticmethod
    @contextmanager
    def _com_initialized() -> Iterator[None]:
        pythoncom = None
        try:
            import pythoncom as module  # type: ignore[import-not-found]

            pythoncom = module
            pythoncom.CoInitialize()
        except ImportError:
            pass
        try:
            yield
        finally:
            if pythoncom is not None:
                pythoncom.CoUninitialize()

    @staticmethod
    def _invalid_validation(request, errors):
        mailbox = MailboxValidationResult(False, "", "", "NOT_CHECKED", False, False)
        outbound = OutboundSafetyState(False, "UNKNOWN", "PREVIEW", False, False)
        return OutlookRevisiValidationResult(
            False,
            False,
            request.workflow,
            request.period_start,
            request.period_end,
            request.output_root,
            mailbox,
            outbound,
            errors=tuple(errors),
        )

    @classmethod
    def _cancelled(cls, request, started, message):
        return OutlookRevisiRunResult(
            False,
            True,
            request.job_id,
            request.workflow,
            TARGET_MAILBOX_SMTP,
            started,
            cls._now(),
            request.output_root,
            None,
            {},
            {},
            {},
            (),
            warning_count=1,
            error_summary=message,
        )

    @staticmethod
    def _existing_path(value) -> Path | None:
        path = Path(str(value)) if value else None
        return path if path is not None and path.exists() else None

    @staticmethod
    def _to_bool(value) -> bool:
        return str(value or "").strip().casefold() in {
            "1",
            "true",
            "yes",
            "y",
            "active",
        }

    @staticmethod
    def _to_bool_default_true(value) -> bool:
        if value is None or str(value).strip() == "":
            return True
        return OutlookRevisiAdapter._to_bool(value)

    @classmethod
    def _log(cls, level: str, message: str) -> OutlookRevisiLogEvent:
        return OutlookRevisiLogEvent(cls._now(), level, message)

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).isoformat(timespec="seconds")
