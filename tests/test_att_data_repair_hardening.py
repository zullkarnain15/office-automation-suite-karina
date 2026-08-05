from __future__ import annotations

import hashlib
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from shared.database import SQLiteConnectionFactory, SchemaManager
from ui.adapters.att_data_repair_adapter import AttDataRepairAdapter
from ui.services.att_data_repair_service import AttDataRepairService
from ui.services.protocols import StorageStatusView
from ui.utilities_models import AttDataRepairCancellationToken, AttDataRepairRunRequest
from utilities.att_data_repair.constants import (
    INVALID_RECORDS_COLUMNS,
    VALID_RECORDS_COLUMNS,
)
from utilities.att_data_repair.engine import AttDataRepairEngine
from utilities.att_data_repair.statuses import JobStatus
from utilities.att_data_repair.txt_writer import AttDataRepairTxtWriter


class Storage:
    def __init__(self, root: Path, database: Path) -> None:
        self.root = root
        self.database = database

    def resolve_status(self):
        return StorageStatusView(
            "READY",
            self.root,
            self.database,
            True,
            True,
            3,
            "test",
            "Fake Registry",
            self.root / "profiles",
            self.root / "backup",
            self.root / "output",
            self.root / "logs",
            self.root / "diagnostics",
        )


class FixedClock:
    def __call__(self) -> datetime:
        return datetime(2026, 8, 3, 9, 0, 0)


class CodeProvider:
    def __call__(self) -> int:
        return 4837


def test_ui_e2e_preserves_source_and_keeps_status_layers_consistent(
    tmp_path: Path,
) -> None:
    database = _database(tmp_path)
    source = tmp_path / "source.xlsx"
    _write_source(source, invalid_nik=True)
    before_hash = _sha256(source)
    service = AttDataRepairService(
        Storage(tmp_path, database),
        AttDataRepairAdapter(
            AttDataRepairEngine(
                txt_writer=AttDataRepairTxtWriter(CodeProvider()),
                clock=FixedClock(),
            )
        ),
    )
    resolved, validation = service.preflight(
        AttDataRepairRunRequest(
            source,
            False,
            "2026-08-01",
            "2026-08-31",
            False,
            tmp_path / "output",
            True,
            True,
        ),
        cancellation=AttDataRepairCancellationToken(),
    )

    result = service.run_job(
        resolved,
        validation,
        cancellation=AttDataRepairCancellationToken(),
        progress=lambda event: None,
        log=lambda event: None,
    )

    assert _sha256(source) == before_hash
    assert result.status == JobStatus.PARTIAL_SUCCESS
    payload = json.loads(
        next(item.path for item in result.outputs if str(item.role) == "SUMMARY_JSON")
        .read_text(encoding="utf-8")
    )
    report_path = next(item.path for item in result.outputs if str(item.role) == "EXCEL_REPORT")
    process_log = next(item.path for item in result.outputs if str(item.role) == "PROCESS_LOG")
    workbook = load_workbook(report_path, read_only=True, data_only=True)
    try:
        summary_rows = list(workbook["Process_Summary"].iter_rows(values_only=True))
        report_status = next(row[1] for row in summary_rows if row[0] == "Job Status")
    finally:
        workbook.close()
    with SQLiteConnectionFactory().connect(database, read_only=True) as connection:
        history = connection.execute("SELECT * FROM job_history").fetchone()
        files = {
            row["file_role"]
            for row in connection.execute("SELECT file_role FROM job_files")
        }
    assert payload["status"] == result.status == report_status
    assert history["feature_code"] == "Att Data Repair"
    assert history["legacy_status"] == result.status
    assert history["unified_status"] == "COMPLETED_WITH_WARNING"
    assert {"HRIS_TXT", "EXCEL_REPORT", "PROCESS_LOG", "SUMMARY_JSON"} <= files
    assert "Job status: PARTIAL_SUCCESS" in process_log.read_text(encoding="utf-8")


def test_report_disabled_no_valid_records_keeps_summary_and_history_consistent(
    tmp_path: Path,
) -> None:
    database = _database(tmp_path)
    source = tmp_path / "source.xlsx"
    _write_source(source, invalid_nik=True, only_invalid=True)
    service = AttDataRepairService(
        Storage(tmp_path, database),
        AttDataRepairAdapter(AttDataRepairEngine(clock=FixedClock())),
    )
    resolved, validation = service.preflight(
        AttDataRepairRunRequest(
            source,
            False,
            "2026-08-01",
            "2026-08-31",
            False,
            tmp_path / "output",
            True,
            False,
        ),
        cancellation=AttDataRepairCancellationToken(),
    )

    result = service.run_job(
        resolved,
        validation,
        cancellation=AttDataRepairCancellationToken(),
        progress=lambda event: None,
        log=lambda event: None,
    )

    assert result.status == JobStatus.NO_VALID_RECORDS
    assert result.report_generated is False
    assert result.txt_file_count == 0
    roles = {str(item.role) for item in result.outputs}
    assert "EXCEL_REPORT" not in roles
    summary = json.loads(
        next(item.path for item in result.outputs if str(item.role) == "SUMMARY_JSON")
        .read_text(encoding="utf-8")
    )
    assert summary["status"] == "NO_VALID_RECORDS"
    assert summary["report"]["generated"] is False
    assert summary["txt"]["generated"] is False
    with SQLiteConnectionFactory().connect(database, read_only=True) as connection:
        history = connection.execute("SELECT * FROM job_history").fetchone()
        assert history["legacy_status"] == "NO_VALID_RECORDS"
        assert history["unified_status"] == "COMPLETED_WITH_WARNING"


def _database(tmp_path: Path) -> Path:
    database = tmp_path / "OAS-K.db"
    SchemaManager().initialize_database(database, "hardening")
    return database


def _write_source(
    path: Path,
    *,
    invalid_nik: bool = False,
    only_invalid: bool = False,
) -> None:
    workbook = Workbook()
    workbook.active.title = "Valid_Records"
    valid = workbook["Valid_Records"]
    invalid = workbook.create_sheet("Invalid_Records")
    valid.append(VALID_RECORDS_COLUMNS)
    invalid.append(INVALID_RECORDS_COLUMNS)
    if not only_invalid:
        valid.append(
            [
                1,
                "revision.xlsx",
                "nested/revision.xlsx",
                2,
                "HO",
                "000001234",
                "08/03/2026",
                "09:30",
                "08/03/2026",
                "17:00",
                "Outlook_Revisi_HO_001.txt",
                "VALID",
            ]
        )
    if invalid_nik:
        valid.append(
            [
                2,
                "revision.xlsx",
                "nested/revision.xlsx",
                3,
                "Branch",
                "#REF!",
                "08/03/2026",
                "09:30",
                "08/03/2026",
                "17:00",
                "Outlook_Revisi_Branch_001.txt",
                "VALID",
            ]
        )
    workbook.save(path)
    workbook.close()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
