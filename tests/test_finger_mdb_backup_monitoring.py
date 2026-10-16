"""Backup MDB monitoring rules without network shares or real Outlook."""

from __future__ import annotations

import os
import sqlite3
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from shared.finger_monitoring.mdb_backup_service import (
    MDBBackupMonitoringService,
    MDBBackupScanner,
    MDBRequirement,
)


class _Sender:
    def __init__(self) -> None:
        self.messages = []

    def send_html(self, recipient, subject, body_html) -> None:
        self.messages.append((recipient, subject, body_html))


def _mdb(root: Path, unit: str, branch: str, name: str, modified: datetime) -> Path:
    path = root / "MDB" / unit / branch / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(name.encode("ascii"))
    timestamp = modified.timestamp()
    os.utime(path, (timestamp, timestamp))
    return path


def test_cutoff_h_minus_three_is_inclusive(tmp_path):
    boundary = _mdb(
        tmp_path, "SOF", "BDU", "boundary.mdb", datetime(2026, 10, 13, 0, 0),
    )
    stale = _mdb(
        tmp_path, "SOF", "BDU", "stale.mdb", datetime(2026, 10, 12, 23, 59),
    )

    progress = []
    snapshot = MDBBackupScanner().scan(
        tmp_path, date(2026, 10, 16),
        progress_callback=lambda current, total, message: progress.append(
            (current, total, message)
        ),
    )
    rows = {row.path: row for row in snapshot.files}

    assert snapshot.cutoff_date == date(2026, 10, 13)
    assert rows[boundary].status == "Memenuhi"
    assert rows[stale].status == "Kedaluwarsa"
    assert progress[-1][:2] == (2, 2)


def test_multiple_mdb_are_independent_and_branch_can_be_partial(tmp_path):
    _mdb(tmp_path, "OTO", "MDU", "fresh.mdb", datetime(2026, 10, 16, 8, 0))
    _mdb(tmp_path, "OTO", "MDU", "old.mdb", datetime(2026, 10, 1, 8, 0))

    snapshot = MDBBackupScanner().scan(tmp_path, date(2026, 10, 16))

    assert [row.status for row in snapshot.files] == ["Memenuhi", "Kedaluwarsa"]
    assert snapshot.branches[0].status == "Sebagian diperbarui"
    assert snapshot.branches[0].total_files == 2


def test_required_missing_archive_exclusion_and_access_failure(tmp_path):
    fresh = _mdb(tmp_path, "SOF", "BDU", "fresh.mdb", datetime(2026, 10, 16, 8, 0))
    ignored = tmp_path / "MDB" / "SOF" / "BDU" / "arsip" / "old.mdb"
    ignored.parent.mkdir(parents=True)
    ignored.write_bytes(b"archive")
    locked = _mdb(tmp_path, "OTO", "MDU", "locked.mdb", datetime(2026, 10, 16, 8, 0))

    def stat_provider(path: Path):
        if path == locked:
            raise PermissionError("access denied")
        return path.stat()

    scanner = MDBBackupScanner(stat_provider=stat_provider)
    snapshot = scanner.scan(
        tmp_path,
        date(2026, 10, 16),
        required=(MDBRequirement("OTO", "MDU", "required.mdb", "Manual"),),
        archive_exclusions=("arsip",),
    )

    by_file = {row.file_name: row for row in snapshot.files}
    assert fresh.name in by_file
    assert ignored.name not in by_file
    assert by_file["required.mdb"].status == "File wajib hilang"
    assert by_file["locked.mdb"].status == "Gagal diakses"


def test_export_contains_branch_summary_detail_and_scan_metadata(tmp_path):
    _mdb(tmp_path, "OTO", "MDU", "fresh.mdb", datetime(2026, 10, 16, 8, 0))
    _mdb(tmp_path, "OTO", "MDU", "old.mdb", datetime(2026, 10, 1, 8, 0))
    service = MDBBackupMonitoringService(now=lambda: datetime(2026, 10, 16, 9, 30))
    snapshot = service.scan(tmp_path, date(2026, 10, 16))
    output = tmp_path / "backup-report.xlsx"

    service.export(snapshot, output)

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["Ringkasan Cabang", "Detail File"]
    detail = workbook["Detail File"]
    assert detail["B1"].value == snapshot.checked_at.isoformat(sep=" ")
    assert detail["B2"].value == "2026-10-16"
    assert detail["B3"].value == "2026-10-13"
    assert detail.max_row == 8
    assert {detail.cell(row=row, column=6).value for row in (7, 8)} == {
        "Memenuhi", "Kedaluwarsa",
    }
    workbook.close()


def test_email_contains_only_problem_files_and_skips_all_current(tmp_path):
    _mdb(tmp_path, "OTO", "MDU", "fresh.mdb", datetime(2026, 10, 16, 8, 0))
    _mdb(tmp_path, "OTO", "MDU", "old.mdb", datetime(2026, 10, 1, 8, 0))
    sender = _Sender()
    service = MDBBackupMonitoringService(
        sender=sender, now=lambda: datetime(2026, 10, 16, 9, 30),
    )
    snapshot = service.scan(
        tmp_path,
        date(2026, 10, 16),
        required=(MDBRequirement("SOF", "BDU", "missing.mdb", "Manual"),),
    )

    assert service.send_email(snapshot, "ops@example.test")
    recipient, subject, body = sender.messages[0]
    assert recipient == "ops@example.test"
    assert "2 file bermasalah" in subject
    assert "old.mdb" in body and "missing.mdb" in body
    assert "fresh.mdb" not in body
    assert "OTO" in body and "MDU" in body and "SOF" in body and "BDU" in body

    clean_root = tmp_path / "clean"
    _mdb(clean_root, "OTO", "MDU", "good.mdb", datetime(2026, 10, 16, 8, 0))
    clean = service.scan(clean_root, date(2026, 10, 16))
    assert not service.send_email(clean, "ops@example.test")
    assert len(sender.messages) == 1


def test_required_parser_supports_manual_and_attendance_paths():
    manual = MDBBackupMonitoringService.parse_required_entries(
        "OTO/MDU/one.mdb; SOF/BDU/two.mdb"
    )
    attendance = MDBBackupMonitoringService.requirements_from_attendance_paths(
        (Path("Z:/share/MDB/CF/JKT/three.mdb"),)
    )

    assert [(item.unit, item.branch, item.file_name) for item in manual] == [
        ("OTO", "MDU", "one.mdb"), ("SOF", "BDU", "two.mdb"),
    ]
    assert [(item.unit, item.branch, item.file_name) for item in attendance] == [
        ("CF", "JKT", "three.mdb"),
    ]


def test_active_attendance_sources_are_loaded_as_required_mdb(tmp_path):
    database = tmp_path / "core.db"
    with sqlite3.connect(database) as connection:
        connection.execute(
            """CREATE TABLE attendance_sources (
                workflow TEXT, source_code TEXT, mdb_path TEXT,
                is_active INTEGER, sort_order INTEGER
            )"""
        )
        connection.executemany(
            "INSERT INTO attendance_sources VALUES (?,?,?,?,?)",
            (
                ("BRANCH", "MDU", "Z:/share/MDB/OTO/MDU/main.mdb", 1, 1),
                ("BRANCH", "OLD", "Z:/share/MDB/OTO/OLD/archive.mdb", 0, 2),
            ),
        )

    class _Storage:
        def resolve_status(self):
            return type("Status", (), {
                "database_valid": True, "database_path": database,
            })()

    required = MDBBackupMonitoringService(_Storage()).load_attendance_requirements()

    assert [(item.unit, item.branch, item.file_name, item.source) for item in required] == [
        ("OTO", "MDU", "main.mdb", "Attendance"),
    ]
