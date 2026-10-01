"""Sprint 4 cross-view correctness, export, and destructive-operation safety."""

import sqlite3
import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.test_attendance_ot_enrichment import _environment, _seed
from shared.attendance_ot import (
    AttendanceOTEnrichmentService,
    AttendanceOTAnalyticsQueryService,
)
from shared.attendance_ot.reporting import (
    AnalyticsFilters,
    DETAIL_COLUMNS,
    DETAIL_LABELS,
    MONTHS,
    MASTER_COLUMNS,
)
from shared.attendance_ot.export_service import AttendanceOTExportService
from shared.attendance_ot.maintenance import AttendanceOTMaintenanceService


@pytest.fixture
def reports(tmp_path):
    storage, database = _environment(tmp_path)
    _seed(database)
    AttendanceOTEnrichmentService(batch_size=2).rebuild(database)
    return storage, database, AttendanceOTAnalyticsQueryService(storage)


@pytest.mark.parametrize(
    "filters,expected",
    [
        (AnalyticsFilters(), 8),
        (AnalyticsFilters(period="2026-01"), 3),
        (AnalyticsFilters(date_from="2026-01-03", date_to="2026-02-01"), 2),
        (
            AnalyticsFilters(
                branches=("Branch Bandung", "Branch Surabaya"), pay_groups=("PG-BR",)
            ),
            2,
        ),
        (AnalyticsFilters(departments=("Finance",), status="LATE"), 1),
        (AnalyticsFilters(divisions=("Corporate",)), 3),
        (AnalyticsFilters(search="E1", classification="HO"), 3),
        (AnalyticsFilters(search="' OR 1=1 --"), 0),
        (AnalyticsFilters(status="ON_TIME"), 2),
    ],
)
def test_filters_match_dashboard_details_matrix_and_source_amounts(
    reports, filters, expected
):
    storage, database, query = reports
    dash = query.view(filters)
    detail = query.view(filters, tab="Detail Data")
    matrix = query.view(filters, tab="Data Terlambat", year=2026)
    assert detail.total == expected == dash["kpi"]["records"]
    assert dash["kpi"]["late_records"] == sum(
        r["late_status"] == "LATE" for r in detail.items
    )
    assert dash["kpi"]["ot_amount"] == sum(r["ot_amount"] for r in detail.items)
    assert dash["kpi"]["meal_amount"] == sum(r["meal_ot_amount"] for r in detail.items)
    assert sum(r["Total"] for r in matrix.items) == dash["kpi"]["late_records"]
    assert all(sum(r[m] for m in MONTHS) == r["Total"] for r in matrix.items)
    assert dash["kpi"]["employees"] == len({r["emplid"] for r in detail.items})


def test_rankings_count_records_not_minutes_and_query_uses_aggregate(reports):
    _, _, query = reports
    with query.snapshot() as session:
        statements = []
        session.connection.set_trace_callback(statements.append)
        result = session.dashboard(AnalyticsFilters())
    assert all(
        "attendance_records" not in s and "enriched_attendance" not in s
        for s in statements
    )
    assert sum("LIMIT 5" in statement.upper() for statement in statements) == 4
    assert result["rankings"]["HO Employees"][0]["late_count"] == 1
    assert result["rankings"]["HO Departments"][0]["department"] == "Finance"
    assert len(result["rankings"]["BRANCH Employees"]) == 3
    assert all(r["late_count"] == 1 for r in result["rankings"]["BRANCH Employees"])


def test_detail_and_master_pagination_are_bounded(reports):
    _, _, query = reports
    with query.snapshot() as session:
        first = session.details(AnalyticsFilters(), limit=3)
        second = session.details(AnalyticsFilters(), limit=3, offset=3)
        assert first.total == 8
        assert len(first.items) == len(second.items) == 3
        assert first.items[0] != second.items[0]
        assert tuple(first.items[0]) == DETAIL_COLUMNS
        for kind in MASTER_COLUMNS:
            page = session.master(kind, limit=2)
            assert len(page.items) <= 2 and page.total >= 1
            assert tuple(page.items[0]) == MASTER_COLUMNS[kind]
        assert session.master("Employee Snapshot", search="Alice").total == 1
        assert session.sources(limit=2).total == 3
    assert query.filter_options()["pay_groups"] == ("PG-BR", "PG-HO")
    assert query.filter_options()["divisions"] == (
        "Corporate",
        "Operations",
        "Retail",
        "Technology",
    )


def test_employee_master_keeps_one_current_row_and_marks_missing_inactive(reports):
    _, database, query = reports
    with sqlite3.connect(database) as connection:
        connection.execute(
            """
            INSERT INTO source_registry (
                source_id, source_type, source_path, filename, size_bytes,
                modified_time_ns, row_count, status, imported_at, fingerprint,
                error_detail, last_seen_at, created_at, updated_at
            ) VALUES (4, 'EMPLOYEE', 'C:/fixture/latest-employee.xlsx',
                      'latest-employee.xlsx', 1, 2, 1, 'NEW',
                      '2026-10-01T00:00:00', 'latest', NULL,
                      '2026-10-01T00:00:00', '2026-10-01T00:00:00',
                      '2026-10-01T00:00:00')
            """
        )
        connection.execute(
            """
            INSERT INTO employee_records (
                source_id, source_file, source_sheet, source_row, emplid,
                name, pay_group, status, location_descr, regional, dept_desc,
                organization_json, snapshot_date, division, job_title
            ) VALUES (4, 'latest-employee.xlsx', 'Data', 2, 'E1',
                      'Alice Latest', 'PG-NEW', 'A', 'HEAD OFFICE', 'WEST',
                      'Finance', '{}', '2026-10-01', 'Corporate', 'Manager')
            """
        )
    with query.snapshot() as session:
        page = session.master("Employee Snapshot")
        monthly = session.master("Employee Monthly Count")
    assert page.total == 5
    rows = {row["emplid"]: row for row in page.items}
    assert rows["E1"]["empl_status"] == "ACTIVE"
    assert rows["E1"]["pay_group"] == "PG-NEW"
    assert rows["E1"]["job_title"] == "Manager"
    assert rows["E1"]["last_seen_period"] == "2026-10"
    assert rows["E3"]["empl_status"] == "INACTIVE"
    assert tuple(page.items[0]) == MASTER_COLUMNS["Employee Snapshot"]
    counts = {row["snapshot_period"]: row["employee_count"] for row in monthly.items}
    assert counts["2026-10"] == 1
    assert counts["2026-01"] == 4


def test_schedule_master_is_current_per_employee_and_detail_honors_end_effdt(reports):
    _, _, query = reports
    with query.snapshot() as session:
        schedules = session.master("Schedule Assignment")
        january = session.details(AnalyticsFilters(period="2026-01"))
        february = session.details(AnalyticsFilters(period="2026-02"))
    assert schedules.total == 5
    current = {row["emplid"]: row for row in schedules.items}
    assert current["E1"]["schedule_id"] == "WS_CCD"
    assert current["E1"]["assignment_status"] == "ACTIVE"
    assert {row["resolved_schedule_id"] for row in january.items if row["emplid"] == "E1"} == {"WS_HO 3"}
    assert {row["resolved_schedule_id"] for row in february.items if row["emplid"] == "E1"} == {"WS_CCD"}


def test_excel_all_filtered_and_header_aware_split(reports, tmp_path):
    _, database, query = reports
    service = AttendanceOTExportService(query, sheet_row_limit=4)
    target = tmp_path / "all.xlsx"
    service.export(target)
    book = load_workbook(target, read_only=True, data_only=False)
    try:
        assert "Overview" in book.sheetnames
        assert "Source_Files" in book.sheetnames
        for prefix, first_header in (
            ("Late_Summary", "Year"),
            ("Employee_Summary", "EMPLID"),
        ):
            summaries = [s for s in book if s.title.startswith(prefix)]
            assert summaries
            for sheet in summaries:
                values = list(sheet.values)
                assert len(values) <= 4
                assert values[0][0] == first_header
        details = [s for s in book if s.title.startswith("Detail_Data")]
        assert [s.title for s in details] == [
            "Detail_Data_01",
            "Detail_Data_02",
            "Detail_Data_03",
        ]
        rows = []
        for sheet in details:
            values = list(sheet.values)
            assert len(values) <= 4
            assert values[0][1] == "EMPLID"
            assert values[0] == DETAIL_LABELS
            rows.extend(values[1:])
        assert len(rows) == 8
        assert sum(row[16] for row in rows) == 310
        assert sum(row[17] for row in rows) == 62
    finally:
        book.close()
    filters = AnalyticsFilters(
        date_from="2026-02-02", date_to="2026-02-02", pay_groups=("PG-BR",), search="E1"
    )
    service.export(tmp_path / "filtered.xlsx", filters)
    book = load_workbook(tmp_path / "filtered.xlsx", read_only=True)
    try:
        rows = list(book["Detail_Data"].values)
        assert len(rows) == 2 and rows[1][1] == "E1"
        overview = dict(list(book["Overview"].values)[1:])
        assert overview["records"] == 1
    finally:
        book.close()


def test_export_preserves_literal_formula_like_source_text(reports, tmp_path):
    _, database, query = reports
    with sqlite3.connect(database) as c:
        c.execute(
            "UPDATE enriched_attendance SET employee_name='=HYPERLINK(\"x\")' WHERE emplid='E1'"
        )
    target = tmp_path / "text.xlsx"
    AttendanceOTExportService(query).export(target)
    book = load_workbook(target, read_only=True, data_only=False)
    try:
        cell = next(book["Detail_Data"].iter_rows(min_row=2, max_row=2))[2]
        assert cell.value.startswith("=") and cell.data_type == "s"
    finally:
        book.close()


def test_purge_reset_vacuum_core_and_config_preserved(reports):
    storage, database, query = reports
    core = storage.storage_service.database
    before = hashlib.sha256(core.read_bytes()).hexdigest()
    service = AttendanceOTMaintenanceService(storage)
    preview = service.preview("PURGE", "2026-01-01", "2026-01-31")
    assert dict(preview.counts)["attendance_records"] == 3
    with pytest.raises(ValueError, match="Konfirmasi"):
        service.execute(preview, confirmation="yes")
    assert query.view()["kpi"]["records"] == 8
    result = service.execute(preview, confirmation="PURGE")
    assert result["deleted_records"] == 3
    assert Path(result["backup"]).is_file()
    with sqlite3.connect(result["backup"]) as backup:
        assert (
            backup.execute("SELECT COUNT(*) FROM attendance_records").fetchone()[0] == 8
        )
    assert query.view()["kpi"]["records"] == 5
    assert query.view()["kpi"]["late_records"] == 3
    assert query.view(tab="Master Data").total == 5
    info = service.info()
    assert info["schema_version"] == 5 and info["last_cleanup"]
    assert info["period_from"] == info["period_to"] == "2026-02-02"
    service.vacuum()
    preview = service.preview("RESET")
    service.execute(preview, confirmation="RESET")
    assert query.view()["kpi"]["records"] == 0
    assert query.view(tab="Source Files").total == 0
    assert query.view(tab="Master Data").total == 0
    assert service.info()["database_uuid"] == info["database_uuid"]
    assert hashlib.sha256(core.read_bytes()).hexdigest() == before
    with sqlite3.connect(database) as c:
        assert c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert not c.execute("PRAGMA foreign_key_check").fetchall()


def test_duplicate_preview_and_delete_use_full_in_out_key(reports):
    storage, database, query = reports
    with sqlite3.connect(database) as connection:
        original = connection.execute(
            """
            SELECT source_id, source_file, source_sheet, emplid, date_in,
                   time_in, date_out, time_out, overtime_hour,
                   overtime_minute, ot_amount, meal_ot_amount,
                   location_descr, name, paylink
            FROM attendance_records WHERE attendance_record_id=1
            """
        ).fetchone()
        connection.execute(
            """
            INSERT INTO attendance_records (
                source_id, source_file, source_sheet, source_row, emplid,
                date_in, time_in, date_out, time_out, overtime_hour,
                overtime_minute, ot_amount, meal_ot_amount,
                location_descr, name, paylink
            ) VALUES (?, ?, ?, 99, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            original,
        )
        # Same NIK/date/time-in but a different Date Out is not a duplicate.
        changed_out = list(original)
        changed_out[6] = "2026-01-03"
        connection.execute(
            """
            INSERT INTO attendance_records (
                source_id, source_file, source_sheet, source_row, emplid,
                date_in, time_in, date_out, time_out, overtime_hour,
                overtime_minute, ot_amount, meal_ot_amount,
                location_descr, name, paylink
            ) VALUES (?, ?, ?, 100, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            changed_out,
        )
    AttendanceOTEnrichmentService(batch_size=2).rebuild(database)

    service = AttendanceOTMaintenanceService(storage)
    preview = service.preview_duplicates()
    assert preview.duplicate_groups == 1
    assert preview.duplicate_rows == 1
    assert preview.samples[0]["source_row"] == 99

    with pytest.raises(ValueError, match="DELETE DUPLICATES"):
        service.delete_duplicates(preview, confirmation="DELETE")

    result = service.delete_duplicates(
        preview, confirmation="DELETE DUPLICATES"
    )
    assert result["deleted_records"] == 1
    assert Path(result["backup"]).is_file()
    assert service.preview_duplicates().duplicate_rows == 0
    assert query.view()["kpi"]["records"] == 9
    detail = query.view(tab="Detail Data")
    assert detail.total == 9
    assert {"date_in", "date_out"}.issubset(detail.items[0])
    with sqlite3.connect(database) as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM attendance_records"
        ).fetchone()[0] == 9
        assert connection.execute("PRAGMA quick_check").fetchone()[0] == "ok"


def test_stale_preview_rejected_and_failed_cleanup_rolls_back(reports):
    storage, database, query = reports
    service = AttendanceOTMaintenanceService(storage)
    old = service.preview("RESET")
    service.vacuum()
    with pytest.raises(RuntimeError, match="Data berubah"):
        service.execute(old, confirmation="RESET")
    preview = service.preview("PURGE", "2026-01-01", "2026-01-31")
    with sqlite3.connect(database) as c:
        c.execute(
            "CREATE TRIGGER refuse_purge BEFORE DELETE ON attendance_records BEGIN SELECT RAISE(ABORT,'fixture failure'); END"
        )
    with pytest.raises(sqlite3.IntegrityError):
        service.execute(preview, confirmation="PURGE")
    assert query.view()["kpi"]["records"] == 8
    assert query.view(tab="Detail Data").total == 8


def test_dirty_analytics_refuse_report_but_allow_registry_audit(reports):
    _, database, query = reports
    with sqlite3.connect(database) as c:
        c.execute("UPDATE analytics_metadata SET derived_dirty=1")
    with pytest.raises(RuntimeError, match="Refresh Data"):
        query.view()
    assert query.view(tab="Source Files").total == 3
