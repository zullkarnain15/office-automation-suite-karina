from __future__ import annotations

import hashlib
from datetime import date, datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook

from utilities.att_data_repair.constants import (
    INVALID_RECORDS_COLUMNS,
    VALID_RECORDS_COLUMNS,
)
from utilities.att_data_repair.engine import AttDataRepairEngine
from utilities.att_data_repair.models import (
    AttDataRepairRequest,
    MissingRequiredColumnError,
    MissingRequiredSheetError,
)
from utilities.att_data_repair.normalizer import normalize_time
from utilities.att_data_repair.report_discovery import AttDataRepairReportDiscovery
from utilities.att_data_repair.report_reader import AttDataRepairReportReader
from utilities.att_data_repair.statuses import AnomalyCode, ChangeCode, FinalStatus


def _default_valid(**overrides):
    row = {
        "No": 1,
        "Source_File": "revision.xlsx",
        "Relative_Path": "nested/revision.xlsx",
        "Source_Row": 2,
        "Workflow": "HO",
        "NIK": "000001234",
        "Date_In": "08/03/2026",
        "Time_In": "09:30",
        "Date_Out": "08/03/2026",
        "Time_Out": "17:00",
        "Output_TXT": "Outlook_Revisi_HO_001.txt",
        "Status": "VALID",
    }
    row.update(overrides)
    return row


def _default_invalid(**overrides):
    row = {
        "No": 1,
        "Source_File": "bad.xlsx",
        "Relative_Path": "bad.xlsx",
        "Source_Row": 3,
        "Workflow": "Branch",
        "NIK": "000009999",
        "Date_In": "08/04/2026",
        "Time_In": "",
        "Date_Out": "08/04/2026",
        "Time_Out": "",
        "Status_Code": "TIME_FORMAT",
        "Reason": "Invalid time",
        "Raw_Value": "raw bad row",
    }
    row.update(overrides)
    return row


def _write_report(
    path: Path,
    *,
    valid_rows=(),
    invalid_rows=(),
    valid_sheet: str = "Valid_Records",
    invalid_sheet: str = "Invalid_Records",
) -> None:
    workbook = Workbook()
    workbook.active.title = valid_sheet
    valid = workbook[valid_sheet]
    invalid = workbook.create_sheet(invalid_sheet)
    valid.append(VALID_RECORDS_COLUMNS)
    invalid.append(INVALID_RECORDS_COLUMNS)
    for item in valid_rows:
        valid.append([item.get(column) for column in VALID_RECORDS_COLUMNS])
    valid.append([None for _column in VALID_RECORDS_COLUMNS])
    for item in invalid_rows:
        invalid.append([item.get(column) for column in INVALID_RECORDS_COLUMNS])
    workbook.save(path)
    workbook.close()


def _request(path: Path, start=date(2026, 8, 1), end=date(2026, 8, 31)):
    return AttDataRepairRequest(path, start, end)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _single_result(tmp_path: Path, row: dict) -> object:
    report = tmp_path / "source.xlsx"
    _write_report(report, valid_rows=(row,))
    return AttDataRepairEngine().analyze(_request(report))


def test_reader_reads_both_sheets_preserves_lineage_and_does_not_modify_source(
    tmp_path: Path,
) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(_default_valid(No=7),),
        invalid_rows=(_default_invalid(No=8),),
        valid_sheet=" Valid_Records ",
        invalid_sheet="invalid records",
    )
    before = _sha256(report)

    records = AttDataRepairReportReader().read(report)

    assert _sha256(report) == before
    assert [record.source_record_no for record in records] == [7, 8]
    assert [record.record_id for record in records] == ["ADR-000001", "ADR-000002"]
    assert records[0].source_sheet == " Valid_Records "
    assert records[0].source_workbook == report
    assert records[0].source_file == "revision.xlsx"
    assert records[0].relative_path == "nested/revision.xlsx"
    assert records[1].source_reason == "Invalid time"
    assert records[1].raw_value == "raw bad row"


def test_reader_missing_required_sheet_is_fatal(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Valid_Records"
    workbook["Valid_Records"].append(VALID_RECORDS_COLUMNS)
    workbook.save(report)
    workbook.close()

    with pytest.raises(MissingRequiredSheetError):
        AttDataRepairReportReader().read(report)


def test_report_discovery_scans_folder_and_detects_valid_report(
    tmp_path: Path,
) -> None:
    source = tmp_path / "reports"
    nested = source / "2026-08" / "HO"
    nested.mkdir(parents=True)
    valid_report = nested / "attachment_consolidation_report.xlsx"
    invalid_report = source / "wrong.xlsx"
    temp_report = source / "~$open.xlsx"
    _write_report(valid_report, valid_rows=(_default_valid(),))
    temp_report.write_text("temporary lock file", encoding="utf-8")
    workbook = Workbook()
    workbook.active.title = "Only_One_Sheet"
    workbook.save(invalid_report)
    workbook.close()

    result = AttDataRepairReportDiscovery().discover(source, recursive=True)

    assert result.files_scanned == 2
    assert result.skipped_temp_files == 1
    assert [candidate.path for candidate in result.valid_candidates] == [valid_report]
    assert len(result.invalid_candidates) == 1
    assert result.valid_candidates[0].record_count == 1
    assert "Valid_Records" in result.invalid_candidates[0].reason


def test_reader_missing_required_header_is_fatal(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(report)
    workbook = Workbook()
    workbook.active.title = "Valid_Records"
    workbook["Valid_Records"].append(
        [column for column in VALID_RECORDS_COLUMNS if column != "NIK"]
    )
    invalid = workbook.create_sheet("Invalid_Records")
    invalid.append(INVALID_RECORDS_COLUMNS)
    workbook.save(report)
    workbook.close()

    with pytest.raises(MissingRequiredColumnError, match="NIK"):
        AttDataRepairReportReader().read(report)


def test_nik_leading_zero_is_preserved_and_spreadsheet_errors_are_anomaly(
    tmp_path: Path,
) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(NIK=" 000001234 "),
            _default_valid(No=2, NIK="#REF!"),
            _default_valid(No=3, NIK="#N/A"),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert [record.nik for record in result.final_records] == ["000001234"]
    assert [item.anomaly_code for item in result.anomalies] == [
        AnomalyCode.INVALID_NIK,
        AnomalyCode.INVALID_NIK,
    ]


@pytest.mark.parametrize(
    ("original", "expected", "change_code"),
    [
        ("0000031815", "20000031815", ChangeCode.NIK_LEADING_TWO_RESTORED),
        ("2000031815", "20000031815", ChangeCode.NIK_MISSING_ZERO_RESTORED),
        ("2000001815", "20000001815", ChangeCode.NIK_MISSING_ZERO_RESTORED),
        ("2000000815", "20000000815", ChangeCode.NIK_MISSING_ZERO_RESTORED),
        ("2000000015", "20000000015", ChangeCode.NIK_MISSING_ZERO_RESTORED),
        ("2000104012", "000104012", ChangeCode.NIK_EXTRA_LEADING_TWO_REMOVED),
        ("00000031815", "20000031815", ChangeCode.NIK_LEADING_TWO_REPLACED),
        ("00000001815", "20000001815", ChangeCode.NIK_LEADING_TWO_REPLACED),
        ("00000000815", "20000000815", ChangeCode.NIK_LEADING_TWO_REPLACED),
        ("00000000015", "20000000015", ChangeCode.NIK_LEADING_TWO_REPLACED),
    ],
)
def test_nik_known_mistype_patterns_are_repaired_and_audited(
    tmp_path: Path,
    original: str,
    expected: str,
    change_code: ChangeCode,
) -> None:
    result = _single_result(tmp_path, _default_valid(NIK=original))

    assert [record.nik for record in result.final_records] == [expected]
    nik_changes = [
        change for change in result.change_log if change.field_name == "NIK"
    ]
    assert [(change.original_value, change.final_value, change.change_code) for change in nik_changes] == [
        (original, expected, change_code)
    ]
    assert result.final_records[0].final_status == FinalStatus.REPAIRED


@pytest.mark.parametrize(
    "nik",
    [
        "000104012",
        "1100609084",
        "2360411174",
        "9070609050",
        "20000031815",
        "20000001815",
        "20000000815",
        "20000000015",
    ],
)
def test_complete_nik_patterns_are_never_changed(tmp_path: Path, nik: str) -> None:
    result = _single_result(tmp_path, _default_valid(NIK=nik))

    assert [record.nik for record in result.final_records] == [nik]
    assert not [change for change in result.change_log if change.field_name == "NIK"]
    assert result.final_records[0].final_status == FinalStatus.VALID_UNCHANGED


@pytest.mark.parametrize(
    "nik",
    [
        "12345678",
        "123456789012",
        "00010A012",
        "2E10",
        "30000031815",
        "20000318155",
    ],
)
def test_uncertain_or_malformed_nik_becomes_anomaly(tmp_path: Path, nik: str) -> None:
    result = _single_result(tmp_path, _default_valid(NIK=nik))

    assert not result.final_records
    assert [item.anomaly_code for item in result.anomalies] == [
        AnomalyCode.INVALID_NIK
    ]


def test_duplicate_detection_uses_repaired_final_nik(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(NIK="0000031815"),
            _default_valid(No=2, NIK="20000031815"),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert [record.nik for record in result.final_records] == ["20000031815"]
    assert [item.anomaly_code for item in result.anomalies] == [
        AnomalyCode.DUPLICATE_RECORD
    ]
    assert any(
        change.change_code == ChangeCode.NIK_LEADING_TWO_RESTORED
        for change in result.change_log
    )


def test_workflow_normalization_and_invalid_workflow(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Workflow="HO"),
            _default_valid(No=2, Workflow="BRANCH", NIK="000001235"),
            _default_valid(No=3, Workflow=""),
            _default_valid(No=4, Workflow="Store"),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert [record.workflow for record in result.final_records] == ["HO", "Branch"]
    assert [item.anomaly_code for item in result.anomalies] == [
        AnomalyCode.INVALID_WORKFLOW,
        AnomalyCode.INVALID_WORKFLOW,
    ]


def test_date_parsing_alignment_and_period_rules(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Date_In=datetime(2026, 8, 3), Date_Out="#VALUE!"),
            _default_valid(No=2, Date_In="13/08/2026", Date_Out="08/13/2026"),
            _default_valid(No=3, Date_In="03/08", Date_Out="03/08"),
            _default_valid(No=4, Date_In="07/31/2026", Date_Out="07/31/2026"),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert [record.date_in for record in result.final_records] == [
        date(2026, 8, 3),
        date(2026, 8, 13),
        date(2026, 8, 3),
    ]
    assert result.anomalies[0].anomaly_code == AnomalyCode.OUTSIDE_REPORT_PERIOD
    assert any(
        change.change_code == ChangeCode.DATE_OUT_ALIGNED_TO_DATE_IN
        for change in result.change_log
    )


def test_ambiguous_dates_are_not_guessed(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Date_In="03/04/2026", Date_Out="03/04/2026"),
            _default_valid(No=2, Date_In="03/08", Date_Out="03/08"),
        ),
    )

    ambiguous = AttDataRepairEngine().analyze(
        AttDataRepairRequest(report, date(2026, 3, 1), date(2026, 4, 30))
    )
    cross_year = AttDataRepairEngine().analyze(
        AttDataRepairRequest(report, date(2025, 12, 1), date(2026, 4, 30))
    )

    assert ambiguous.anomalies[0].anomaly_code == AnomalyCode.INVALID_DATE
    assert cross_year.anomalies[1].anomaly_code == AnomalyCode.INVALID_DATE


@pytest.mark.parametrize(
    ("value", "expected", "normalized"),
    [
        ("7;20", time(7, 20), True),
        ("7.20", time(7, 20), True),
        ("7:2", time(7, 2), True),
        (time(8, 15), time(8, 15), False),
    ],
)
def test_time_normalization_accepts_safe_values(value, expected, normalized) -> None:
    result = normalize_time(value)

    assert result.value == expected
    assert result.normalized is normalized


@pytest.mark.parametrize("value", ["25:90", "#VALUE!", ""])
def test_time_normalization_rejects_invalid_values(value) -> None:
    assert normalize_time(value).value is None


def test_repair_swaps_time_and_applies_minimum_duration(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Time_In="18:00", Time_Out="09:00"),
            _default_valid(No=2, NIK="000001235", Time_In="09:30", Time_Out="10:30"),
            _default_valid(No=3, NIK="000001236", Time_In="09:30", Time_Out="10:31"),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert result.final_records[0].time_in == time(9, 0)
    assert result.final_records[0].time_out == time(18, 0)
    assert result.final_records[1].duration_minutes == 61
    assert result.final_records[2].final_status == FinalStatus.VALID_UNCHANGED
    assert ChangeCode.TIME_IN_OUT_SWAPPED in {
        change.change_code for change in result.final_records[0].changes
    }
    assert ChangeCode.MINIMUM_DURATION_APPLIED in {
        change.change_code for change in result.final_records[1].changes
    }


def test_repair_partial_and_default_time_rules(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Time_In="14:00", Time_Out="#VALUE!"),
            _default_valid(No=2, Time_In="#REF!", Time_Out="10:00"),
            _default_valid(No=3, Time_In="#REF!", Time_Out="#VALUE!"),
            _default_valid(No=4, Time_In="23:30", Time_Out=""),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert result.final_records[0].time_out == time(17, 0)
    assert result.final_records[1].time_in == time(8, 59)
    assert (result.final_records[2].time_in, result.final_records[2].time_out) == (
        time(9, 30),
        time(17, 0),
    )
    assert (result.final_records[3].time_in, result.final_records[3].time_out) == (
        time(9, 30),
        time(17, 0),
    )


def test_weekday_saturday_and_sunday_policy_rules(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Date_In="08/03/2026", Time_In="", Time_Out=""),
            _default_valid(No=2, Date_In="08/08/2026", Date_Out="08/08/2026", Time_In="", Time_Out=""),
            _default_valid(No=3, Date_In="08/09/2026", Date_Out="08/09/2026", Time_In="", Time_Out=""),
            _default_valid(No=4, Date_In="08/09/2026", Date_Out="08/09/2026", Time_In="10:00", Time_Out="12:00"),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert (result.final_records[0].time_in, result.final_records[0].time_out) == (
        time(9, 30),
        time(17, 0),
    )
    assert (result.final_records[1].time_in, result.final_records[1].time_out) == (
        time(9, 30),
        time(11, 0),
    )
    assert [item.anomaly_code for item in result.anomalies] == [
        AnomalyCode.SUNDAY_NOT_MONTH_END,
        AnomalyCode.SUNDAY_NOT_MONTH_END,
    ]


def test_month_end_sunday_is_retained(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(
                Date_In="2026-05-31",
                Date_Out="2026-05-31",
                Time_In="10:00",
                Time_Out="12:00",
            ),
            _default_valid(
                No=2,
                Date_In="2026-05-24",
                Date_Out="2026-05-24",
            ),
        ),
    )

    result = AttDataRepairEngine(
        clock=lambda: datetime(2026, 5, 31, 9, 0)
    ).analyze(AttDataRepairRequest(report, date(2026, 5, 1), date(2026, 5, 31)))

    assert [record.date_in for record in result.final_records] == [date(2026, 5, 31)]
    assert result.anomalies[0].anomaly_code == AnomalyCode.SUNDAY_NOT_MONTH_END


def test_only_complete_duplicates_move_to_anomaly(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(NIK="000001234", Date_In="08/03/2026"),
            _default_valid(No=2, NIK="000001234", Date_In="08/03/2026"),
            _default_valid(No=3, NIK="000005555", Time_Out=""),
            _default_valid(No=4, NIK="000005555", Time_Out=""),
            _default_valid(No=5, NIK="", Date_In="08/03/2026"),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert len(result.final_records) == 3
    assert [record.nik for record in result.final_records] == [
        "000001234",
        "000005555",
        "000005555",
    ]
    assert {item.anomaly_code for item in result.anomalies} == {
        AnomalyCode.DUPLICATE_RECORD,
        AnomalyCode.INVALID_NIK,
    }
    assert result.source_counts["total"] == 5
    assert result.status_counts[FinalStatus.VALID_UNCHANGED] == 1
    assert result.status_counts[AnomalyCode.INVALID_NIK] == 1
    assert result.status_counts[AnomalyCode.DUPLICATE_RECORD] == 1


def test_midnight_and_saturday_missing_out_rules(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Time_In="00:00", Time_Out="17:00"),
            _default_valid(No=2, Time_In="09:30", Time_Out="00:00"),
            _default_valid(
                No=3,
                Date_In="08/08/2026",
                Date_Out="08/08/2026",
                Time_In="09:30",
                Time_Out="",
            ),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert result.final_records[0].time_in == time(9, 30)
    assert result.final_records[1].time_out == time(23, 59)
    assert result.final_records[2].time_out == time(11, 0)
    codes = {change.change_code for change in result.change_log}
    assert ChangeCode.MIDNIGHT_TIME_IN_DEFAULTED in codes
    assert ChangeCode.MIDNIGHT_TIME_OUT_DEFAULTED in codes
    assert ChangeCode.SATURDAY_MISSING_OUT_DEFAULTED in codes


def test_date_year_is_aligned_before_period_validation(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(Date_In="2025-08-04", Date_Out="2025-08-04"),
        ),
    )

    result = AttDataRepairEngine(
        clock=lambda: datetime(2026, 8, 4, 9, 0)
    ).analyze(_request(report))

    assert result.final_records[0].date_in == date(2026, 8, 4)
    assert result.final_records[0].date_out == date(2026, 8, 4)
    assert {
        change.field_name
        for change in result.change_log
        if change.change_code == ChangeCode.DATE_YEAR_ALIGNED_TO_CURRENT_YEAR
    } == {"Date_In", "Date_Out"}


def test_changed_records_and_change_log_allow_multiple_changes(tmp_path: Path) -> None:
    report = tmp_path / "source.xlsx"
    _write_report(
        report,
        valid_rows=(
            _default_valid(
                Date_Out="08/04/2026",
                Time_In="7;20",
                Time_Out="7.40",
            ),
        ),
    )

    result = AttDataRepairEngine().analyze(_request(report))

    assert len(result.changed_records) == 1
    assert len(result.anomalies) == 0
    assert len(result.change_log) >= 3
    assert {change.record_id for change in result.change_log} == {"ADR-000001"}
