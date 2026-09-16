"""Semantic output checks against the pre-stage-two style/mapping behavior."""
from __future__ import annotations

import json
from copy import copy
from datetime import date, datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import PatternFill

from tools.benchmark_utilities import reference_data_style, reference_output_mapping
from utilities._excel_styles import apply_data_style


def workbook_snapshot(path, replacements=()):
    """Compare values and semantic formatting, not workbook-local style IDs."""
    workbook = load_workbook(path, data_only=False)
    try:
        result = {}
        for sheet in workbook:
            rows = []
            for row in sheet:
                cells = []
                for cell in row:
                    value = cell.value
                    if isinstance(value, str):
                        for old, new in replacements:
                            value = value.replace(old, new)
                    cells.append((value, cell.data_type, cell.number_format, copy(cell.font),
                                  copy(cell.fill), copy(cell.border), copy(cell.alignment), copy(cell.protection)))
                rows.append(cells)
            result[sheet.title] = (rows, sheet.freeze_panes, sheet.auto_filter.ref,
                                   {key: value.width for key, value in sheet.column_dimensions.items()})
        return result
    finally:
        workbook.close()


def test_cached_styles_preserve_types_formats_and_workbook_isolation(tmp_path):
    values = [None, "00000123", "=1+1", "#REF!", "text", 42, 3.25, True,
              date(2026, 8, 3), time(9, 30), datetime(2026, 8, 3, 9, 30)]
    fills = [None, PatternFill("solid", fgColor="FFC7CE"), PatternFill("solid", fgColor="C6EFCE")]
    for workbook_index in range(2):
        snapshots = []
        for label, apply in (("reference", reference_data_style), ("cached", apply_data_style)):
            book = Workbook(write_only=True)
            sheet = book.create_sheet("Data")
            # Different registration order must not reuse IDs from another book.
            if workbook_index:
                warmup = WriteOnlyCell(sheet, value="warmup")
                warmup.fill = PatternFill("solid", fgColor="112233")
                sheet.append([warmup])
            for _ in range(2):
                for fill in fills:
                    for number_format in (None, "@", "mm/dd/yyyy", "hh:mm"):
                        row = []
                        for value in values:
                            cell = WriteOnlyCell(sheet, value=value)
                            apply(cell, number_format=number_format, fill=fill)
                            row.append(cell)
                        sheet.append(row)
            path = tmp_path / f"{label}-{workbook_index}.xlsx"
            book.save(path)
            book.close()
            snapshots.append(workbook_snapshot(path))
        assert snapshots[0] == snapshots[1]


def test_changing_a_styled_cell_does_not_mutate_cached_template():
    book = Workbook()
    sheet = book.active
    fill = PatternFill("solid", fgColor="C6EFCE")
    first, second, third = [WriteOnlyCell(sheet, value="0001") for _ in range(3)]
    apply_data_style(first, number_format="@", fill=fill)
    apply_data_style(second, number_format="@", fill=fill)
    first.number_format = "0.00"
    second.fill = PatternFill("solid", fgColor="FF0000")
    apply_data_style(third, number_format="@", fill=fill)
    assert third.number_format == "@"
    assert third.fill == fill
    assert second.number_format == "@"
    book.close()


def test_all_repair_report_sheets_match_reference_styles(tmp_path, monkeypatch):
    from tests.test_att_data_repair_report import _run_fixture, _request, FixedClock
    from utilities.att_data_repair.artifacts import AttDataRepairJobRequest
    from utilities.att_data_repair import report_writer

    source, result = _run_fixture(tmp_path)
    path = result.report_artifact.file_path
    optimized = workbook_snapshot(path)
    txt_before = {artifact.file_name: artifact.file_path.read_bytes() for artifact in result.txt_artifacts}
    path.unlink()
    monkeypatch.setattr(report_writer, "apply_data_style", reference_data_style)
    report_writer.AttDataRepairReportWriter().write(
        report_folder=result.paths.report_folder, job_id=result.job_id, status=result.status,
        request=AttDataRepairJobRequest(_request(source), tmp_path), paths=result.paths,
        analysis=result.analysis_result, txt_artifacts=result.txt_artifacts,
        record_txt_assignment=result.record_txt_assignment, generated_at=FixedClock()(),
    )
    assert len(optimized) == 8
    assert optimized == workbook_snapshot(path)
    assert txt_before == {artifact.file_name: artifact.file_path.read_bytes() for artifact in result.txt_artifacts}


def test_comparison_report_rerender_preserves_values_and_styles(tmp_path, monkeypatch):
    from tests.test_attendance_reconciliation_engine import request
    from utilities.attendance_reconciliation.engine import ReconciliationEngine
    from utilities.attendance_reconciliation import excel_writer

    class FixedDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 8, 20, 9, 0)

    monkeypatch.setattr(excel_writer, "datetime", FixedDatetime)
    selected = request(tmp_path)
    engine = ReconciliationEngine()
    result = engine.run(selected)
    assert result.success
    summary = json.loads(result.summary_json.read_text(encoding="utf-8"))
    args = (result.report_file, result.job_id, selected, result.scan, result.comparisons,
            result.conflicts, result.duplicates, summary)
    engine.excel_writer.write(*args)
    optimized = workbook_snapshot(result.report_file)
    # Comparison's existing writer is retained after profiling. Re-render the
    # same validated result to guard deterministic values and formatting.
    engine.excel_writer.write(*args)
    assert optimized == workbook_snapshot(result.report_file)


def test_consolidation_txt_and_report_match_reference_mapping(tmp_path, monkeypatch):
    from tests.test_attachment_consolidation import _excel_attachment
    from utilities.attachment_consolidation import engine as module
    from utilities.attachment_consolidation.models import ConsolidationRequest, MODE_EXCEL

    class FixedDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 8, 20, 9, 0)

    monkeypatch.setattr(module, "datetime", FixedDatetime)
    monkeypatch.setattr(module.OutlookTxtWriter, "_unique_random_suffix", staticmethod(
        lambda used: f"{len(used):03d}"
    ))
    source = tmp_path / "source"
    source.mkdir()
    for name in ("a.xlsx", "b.xlsx"):
        _excel_attachment(source / name)
    before_sources = {p.name: p.read_bytes() for p in source.iterdir()}
    snapshots = []
    for name in ("optimized", "reference"):
        if name == "reference":
            monkeypatch.setattr(module.AttachmentConsolidationEngine, "_assign_output_files", staticmethod(reference_output_mapping))
        output = tmp_path / name
        result = module.AttachmentConsolidationEngine().run(
            ConsolidationRequest(MODE_EXCEL, "HO", source, output), txt_max_lines=1,
        )
        assert result.success
        snapshots.append((
            {p.name: p.read_bytes() for p in result.output_files},
            workbook_snapshot(result.artifacts.report_file, [(str(output), "<OUTPUT>")]),
            [(r.source_row, r.nik, r.date_in, r.time_in, r.date_out, r.time_out) for r in result.records],
            [(a.source_row, a.code, a.reason) for a in result.anomalies],
        ))
    assert snapshots[0] == snapshots[1]
    assert before_sources == {p.name: p.read_bytes() for p in source.iterdir()}


@pytest.mark.parametrize("case", ["attachment_excel", "attachment_txt", "comparison", "repair"])
def test_benchmark_jobs_validate_record_counts_and_stages(tmp_path, case):
    from tools.benchmark_utilities import make_fixtures, run_case

    make_fixtures(tmp_path / "source", 8)
    sample = run_case(case, tmp_path / "source", tmp_path / "output", 8)
    assert sample["records"] == 8
    assert sample["seconds"] > 0
    assert sample["stages_seconds"]["excel_write"] > 0
    assert sample["stage_calls"]["excel_write"] == 1
