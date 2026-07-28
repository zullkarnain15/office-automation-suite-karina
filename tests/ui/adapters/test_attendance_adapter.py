from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from ui.adapters.attendance_adapter import AttendanceAdapter
from ui.attendance_models import AttendanceCancellationToken, AttendanceResolvedRequest


def workbook(path: Path, ho: Path, branch: Path) -> Path:
    book = Workbook()
    general = book.active
    general.title = "General"
    general.append(["Parameter", "Value", "Description"])
    general.append(["Split_TXT_Rows", 10000, ""])
    output = book.create_sheet("Output")
    output.append(["Parameter", "Value", "Description"])
    output.append(["Output_Root", str(path.parent / "output"), ""])
    for title, code, source in (("MDB_HO", "HO1", ho), ("MDB_Branch", "BR1", branch)):
        sheet = book.create_sheet(title)
        sheet.append(["Active", "Code", "Description", "MDB_Path"])
        sheet.append(["Y", code, code, str(source)])
    book.save(path)
    return path


def resolved(config: Path, output: Path, workflow: str = "HO"):
    return AttendanceResolvedRequest(
        "UI-JOB",
        None,
        config,
        workflow,
        output,
        "2026-07-01",
        "2026-07-31",
        False,
        False,
        True,
        True,
    )


def test_configuration_validation_selects_ho_and_branch(tmp_path: Path) -> None:
    ho, branch = tmp_path / "ho.mdb", tmp_path / "branch.mdb"
    ho.touch()
    branch.touch()
    config = workbook(tmp_path / "attendance.xlsx", ho, branch)
    adapter = AttendanceAdapter()
    ho_result = adapter.validate_configuration(resolved(config, tmp_path, "HO"))
    branch_result = adapter.validate_configuration(resolved(config, tmp_path, "BRANCH"))
    assert ho_result.valid and ho_result.sources[0].mdb_path == ho
    assert branch_result.valid and branch_result.sources[0].mdb_path == branch


def test_configuration_missing_and_invalid_extension(tmp_path: Path) -> None:
    adapter = AttendanceAdapter()
    assert not adapter.validate_configuration(
        resolved(tmp_path / "missing.xlsx", tmp_path)
    ).valid
    wrong = tmp_path / "config.xls"
    wrong.touch()
    assert not adapter.validate_configuration(resolved(wrong, tmp_path)).valid


def test_adapter_normalizes_engine_outputs_and_progress(tmp_path: Path) -> None:
    config = tmp_path / "config.xlsx"
    config.touch()
    txt = tmp_path / "one.txt"
    report = tmp_path / "report.xlsx"
    folder = tmp_path / "job"
    folder.mkdir()
    log = folder / "Process.log"
    summary = folder / "summary.json"
    for path in (txt, report, log, summary):
        path.touch()

    class Reader:
        def __init__(self, path):
            pass

        def read(self):
            return object()

    class Engine:
        def run(self, **kwargs):
            return {
                "raw_log_count": 4,
                "paired_record_count": 2,
                "valid_record_count": 1,
                "anomaly_record_count": 1,
                "duplicate_removed_count": 0,
                "mdb_summary": [],
                "txt_result": {"generated_files": [{"file_path": str(txt)}]},
                "report_result": {"report_file": str(report)},
                "artifact_result": {
                    "artifact_folder": str(folder),
                    "process_log": str(log),
                    "summary_json": str(summary),
                },
            }

    events = []
    result = AttendanceAdapter(Reader, Engine).run(
        resolved(config, tmp_path),
        cancellation=AttendanceCancellationToken(),
        progress=events.append,
        log=events.append,
    )
    assert result.success and not result.cancelled
    assert {item.file_type for item in result.output_files} == {
        "HRIS_TXT",
        "EXCEL_REPORT",
        "OUTPUT_FOLDER",
        "PROCESS_LOG",
        "SUMMARY_JSON",
    }
    staged = tmp_path / "HRIS" / "HO" / txt.name
    assert staged.exists()
    assert txt.exists()
    assert events


def test_cancellation_before_engine_is_cooperative(tmp_path: Path) -> None:
    token = AttendanceCancellationToken()
    token.request()
    called = []

    class Engine:
        def run(self, **kwargs):
            called.append(True)

    result = AttendanceAdapter(engine_class=Engine).run(
        resolved(tmp_path / "config.xlsx", tmp_path),
        cancellation=token,
        progress=lambda event: None,
        log=lambda event: None,
    )
    assert result.cancelled and not called
