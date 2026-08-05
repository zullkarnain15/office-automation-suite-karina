from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from shared.database import SQLiteConnectionFactory, SchemaManager
from shared.database.repositories import GlobalSettingsRepository
from ui.adapters.att_data_repair_adapter import AttDataRepairAdapter
from ui.services.att_data_repair_service import AttDataRepairService
from ui.services.protocols import StorageStatusView
from ui.utilities_models import (
    AttDataRepairCancellationToken,
    AttDataRepairRunRequest,
)
from utilities.att_data_repair.constants import (
    INVALID_RECORDS_COLUMNS,
    VALID_RECORDS_COLUMNS,
)
from utilities.att_data_repair.engine import AttDataRepairEngine
from utilities.att_data_repair.txt_writer import AttDataRepairTxtWriter


class Storage:
    def __init__(self, root: Path, database: Path, valid: bool = True) -> None:
        self.root, self.database, self.valid = root, database, valid

    def resolve_status(self):
        return StorageStatusView(
            "READY" if self.valid else "INITIAL_SETUP_REQUIRED",
            self.root,
            self.database,
            self.database.exists(),
            self.valid,
            1 if self.valid else None,
            "test",
            "Fake Registry",
            self.root / "profiles",
            self.root / "backup",
            self.root / "output",
            self.root / "logs",
            self.root / "diagnostics",
        )


class FixedClock:
    def __call__(self) -> datetime:
        return datetime(2026, 8, 3, 9, 0, 0)


class CodeProvider:
    def __call__(self) -> int:
        return 4837


def make_database(tmp_path: Path) -> Path:
    database = tmp_path / "OAS-K.db"
    SchemaManager().initialize_database(database, "ui-s6-test")
    with SQLiteConnectionFactory().connect(database) as connection:
        GlobalSettingsRepository(connection).save_global_settings(
            output_root=str(tmp_path / "global-output"),
            period_start="2026-08-01",
            period_end="2026-08-31",
        )
        connection.execute(
            """
            UPDATE att_data_repair_settings
            SET use_global_output = 1,
                use_global_period = 1,
                generate_txt = 1,
                generate_excel_report = 1,
                txt_max_rows = 4321,
                weekday_default_in = '08:00',
                weekday_default_out = '16:30'
            WHERE att_data_repair_settings_id = 1
            """
        )
    return database


def test_resolves_global_values_and_per_job_flags_without_writing_settings(
    tmp_path: Path,
) -> None:
    database = make_database(tmp_path)
    source = tmp_path / "source.xlsx"
    _write_source(source)
    service = AttDataRepairService(Storage(tmp_path, database), AttDataRepairAdapter())

    resolved = service.resolve_request(
        AttDataRepairRunRequest(source, True, None, None, True, None, False, True)
    )

    assert resolved.period_start == "2026-08-01"
    assert resolved.period_end == "2026-08-31"
    assert resolved.output_root == tmp_path / "global-output"
    assert resolved.settings.txt_max_rows == 4321
    assert resolved.settings.weekday_default_in.hour == 8
    assert resolved.generate_txt is False
    assert resolved.generate_excel_report is True
    with SQLiteConnectionFactory().connect(database, read_only=True) as connection:
        row = connection.execute(
            "SELECT generate_txt, generate_excel_report FROM att_data_repair_settings"
        ).fetchone()
    assert tuple(row) == (1, 1)


def test_folder_source_discovers_single_valid_report(tmp_path: Path) -> None:
    database = make_database(tmp_path)
    source_folder = tmp_path / "source-reports"
    valid_folder = source_folder / "2026-08" / "HO"
    valid_folder.mkdir(parents=True)
    source = valid_folder / "source.xlsx"
    invalid = source_folder / "not-a-report.xlsx"
    _write_source(source)
    workbook = Workbook()
    workbook.active.title = "Wrong"
    workbook.save(invalid)
    workbook.close()
    service = AttDataRepairService(Storage(tmp_path, database), AttDataRepairAdapter())

    resolved, validation = service.preflight(
        AttDataRepairRunRequest(
            None,
            True,
            None,
            None,
            True,
            None,
            True,
            True,
            source_folder,
            True,
        ),
        cancellation=AttDataRepairCancellationToken(),
    )

    assert resolved.source_report == source
    assert resolved.source_report_folder == source_folder
    assert validation.valid
    assert validation.discovery_files_scanned == 2
    assert validation.discovered_valid_reports == 1
    assert validation.discovered_invalid_reports == 1
    assert any("Folder scan" in item for item in validation.warnings)


def test_folder_source_rejects_multiple_valid_reports(tmp_path: Path) -> None:
    database = make_database(tmp_path)
    source_folder = tmp_path / "source-reports"
    source_folder.mkdir()
    _write_source(source_folder / "source-a.xlsx")
    _write_source(source_folder / "source-b.xlsx")
    service = AttDataRepairService(Storage(tmp_path, database), AttDataRepairAdapter())

    with pytest.raises(ValueError, match="lebih dari satu Excel Report valid"):
        service.resolve_request(
            AttDataRepairRunRequest(
                None,
                True,
                None,
                None,
                True,
                None,
                True,
                True,
                source_folder,
                True,
            )
        )


def test_local_period_output_and_generate_validation(tmp_path: Path) -> None:
    database = make_database(tmp_path)
    source = tmp_path / "source.xlsx"
    _write_source(source)
    output = tmp_path / "manual-output"
    service = AttDataRepairService(Storage(tmp_path, database), AttDataRepairAdapter())

    resolved = service.resolve_request(
        AttDataRepairRunRequest(
            source,
            False,
            "2026-08-02",
            "2026-08-03",
            False,
            output,
            True,
            False,
        )
    )

    assert resolved.period_start == "2026-08-02"
    assert resolved.output_root == output
    assert not output.exists()
    with pytest.raises(ValueError, match="Tanggal mulai"):
        service.resolve_request(
            AttDataRepairRunRequest(
                source,
                False,
                "2026-08-04",
                "2026-08-03",
                False,
                output,
                True,
                True,
            )
        )
    with pytest.raises(ValueError, match="minimal salah satu"):
        service.resolve_request(
            AttDataRepairRunRequest(
                source,
                False,
                "2026-08-02",
                "2026-08-03",
                False,
                output,
                False,
                False,
            )
        )


def test_user_facing_validation_blocks_bad_inputs(tmp_path: Path) -> None:
    database = make_database(tmp_path)
    service = AttDataRepairService(Storage(tmp_path, database), AttDataRepairAdapter())

    with pytest.raises(ValueError, match="Pilih Excel Report"):
        service.resolve_request(
            AttDataRepairRunRequest(None, True, None, None, True, None, True, True)
        )
    with pytest.raises(ValueError, match="\\.xlsx"):
        service.resolve_request(
            AttDataRepairRunRequest(
                tmp_path / "source.xls",
                True,
                None,
                None,
                True,
                None,
                True,
                True,
            )
        )
    with pytest.raises(ValueError, match="tidak ditemukan"):
        service.resolve_request(
            AttDataRepairRunRequest(
                tmp_path / "missing.xlsx",
                True,
                None,
                None,
                True,
                None,
                True,
                True,
            )
        )
    with SQLiteConnectionFactory().connect(database) as connection:
        connection.execute(
            "UPDATE att_data_repair_settings SET enabled = 0 "
            "WHERE att_data_repair_settings_id = 1"
        )
    with pytest.raises(ValueError, match="dinonaktifkan"):
        service.resolve_request(
            AttDataRepairRunRequest(
                tmp_path / "missing.xlsx",
                True,
                None,
                None,
                True,
                None,
                True,
                True,
            )
        )


def test_global_missing_values_are_blocked(tmp_path: Path) -> None:
    database = make_database(tmp_path)
    source = tmp_path / "source.xlsx"
    _write_source(source)
    with SQLiteConnectionFactory().connect(database) as connection:
        connection.execute(
            "UPDATE global_settings SET period_start = NULL, period_end = NULL, "
            "output_root = '' WHERE global_settings_id = 1"
        )
    service = AttDataRepairService(Storage(tmp_path, database), AttDataRepairAdapter())

    with pytest.raises(ValueError, match="Periode pada General Settings"):
        service.resolve_request(
            AttDataRepairRunRequest(source, True, None, None, False, tmp_path, True, True)
        )
    with pytest.raises(ValueError, match="Output Root pada General Settings"):
        service.resolve_request(
            AttDataRepairRunRequest(
                source,
                False,
                "2026-08-01",
                "2026-08-31",
                True,
                None,
                True,
                True,
            )
        )


def test_preflight_reports_missing_sheet_and_missing_header(tmp_path: Path) -> None:
    database = make_database(tmp_path)
    missing_sheet = tmp_path / "missing-sheet.xlsx"
    missing_header = tmp_path / "missing-header.xlsx"
    _write_source(missing_sheet, include_invalid_sheet=False)
    _write_source(missing_header, omit_required_header=True)
    service = AttDataRepairService(Storage(tmp_path, database), AttDataRepairAdapter())

    for source, message in (
        (missing_sheet, "Sheet Valid_Records atau Invalid_Records"),
        (missing_header, "kolom wajib"),
    ):
        resolved, validation = service.preflight(
            AttDataRepairRunRequest(source, True, None, None, True, None, True, True),
            cancellation=AttDataRepairCancellationToken(),
        )
        assert resolved.source_report == source
        assert not validation.valid
        assert any(message in item for item in validation.errors)
        assert not (tmp_path / "global-output").exists()


def test_run_records_history_and_recovery_artifacts(tmp_path: Path) -> None:
    database = make_database(tmp_path)
    source = tmp_path / "source.xlsx"
    _write_source(source, invalid_nik=True)
    service = AttDataRepairService(
        Storage(tmp_path, database),
        AttDataRepairAdapter(
            AttDataRepairEngine(
                txt_writer=AttDataRepairTxtWriter(CodeProvider()),
                clock=FixedClock(),
            )
        ),
    )
    resolved, validation = service.preflight(
        AttDataRepairRunRequest(
            source,
            False,
            "2026-08-01",
            "2026-08-31",
            False,
            tmp_path / "manual-output",
            True,
            True,
        ),
        cancellation=AttDataRepairCancellationToken(),
    )
    logs = []
    result = service.run_job(
        resolved,
        validation,
        cancellation=AttDataRepairCancellationToken(),
        progress=lambda event: None,
        log=logs.append,
    )

    assert result.success
    assert result.warning_count >= 1
    assert result.output_folder and result.output_folder.exists()
    assert result.report_generated is True
    assert result.txt_file_count == 1
    roles = {str(item.role) for item in result.outputs}
    assert {"OUTPUT_FOLDER", "HRIS_TXT", "EXCEL_REPORT", "PROCESS_LOG", "SUMMARY_JSON"} <= roles
    assert any(item.stage == "HISTORY" for item in logs)
    with SQLiteConnectionFactory().connect(database, read_only=True) as connection:
        job = connection.execute("SELECT * FROM job_history").fetchone()
        assert job["feature_code"] == "Att Data Repair"
        assert job["unified_status"] == "COMPLETED_WITH_WARNING"
        assert connection.execute("SELECT COUNT(*) FROM job_files").fetchone()[0] >= 4
        assert connection.execute("SELECT COUNT(*) FROM job_status_events").fetchone()[0] >= 5
        assert (
            connection.execute(
                """
                SELECT COUNT(*)
                FROM job_files f
                LEFT JOIN job_history h ON h.job_pk = f.job_pk
                WHERE h.job_pk IS NULL
                """
            ).fetchone()[0]
            == 0
        )


def _write_source(
    path: Path,
    *,
    include_invalid_sheet: bool = True,
    omit_required_header: bool = False,
    invalid_nik: bool = False,
) -> None:
    workbook = Workbook()
    workbook.active.title = "Valid_Records"
    valid = workbook["Valid_Records"]
    invalid = workbook.create_sheet("Invalid_Records")
    headers = list(VALID_RECORDS_COLUMNS)
    if omit_required_header:
        headers.remove("NIK")
    valid.append(headers)
    invalid.append(INVALID_RECORDS_COLUMNS)
    valid.append(
        [
            1,
            "revision.xlsx",
            "nested/revision.xlsx",
            2,
            "HO",
            "000001234",
            "08/03/2026",
            "09:30",
            "08/03/2026",
            "17:00",
            "Outlook_Revisi_HO_001.txt",
            "VALID",
        ]
    )
    if invalid_nik:
        valid.append(
            [
                2,
                "revision.xlsx",
                "nested/revision.xlsx",
                3,
                "HO",
                "#REF!",
                "08/03/2026",
                "09:30",
                "08/03/2026",
                "17:00",
                "Outlook_Revisi_HO_001.txt",
                "VALID",
            ]
        )
    if not include_invalid_sheet:
        del workbook["Invalid_Records"]
    workbook.save(path)
    workbook.close()
