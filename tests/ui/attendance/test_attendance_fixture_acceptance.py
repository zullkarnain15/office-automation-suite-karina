from __future__ import annotations

import logging
import time
from pathlib import Path

from openpyxl import Workbook

from config.app_config import APP_VERSION, PROJECT_ROOT
from shared.database import REQUIRED_TABLES, SchemaManager, SQLiteConnectionFactory
from shared.database.repositories import GlobalSettingsRepository
from shared.storage import get_database_path
from shared.storage.registry import FakeRegistryBackend, StorageRegistryService
from ui.adapters.attendance_adapter import AttendanceAdapter
from ui.app import OASKUnifiedApp
from ui.context import AppContext
from ui.services.attendance_service import AttendanceService
from ui.services.service_container import build_default_app_services


def _configuration(path: Path, output: Path) -> tuple[Path, Path, Path]:
    ho = path.parent / "fixture-ho.mdb"
    branch = path.parent / "fixture-branch.mdb"
    ho.touch()
    branch.touch()
    book = Workbook()
    general = book.active
    general.title = "General"
    general.append(["Parameter", "Value", "Description"])
    general.append(["Split_TXT_Rows", 10_000, "fixture"])
    output_sheet = book.create_sheet("Output")
    output_sheet.append(["Parameter", "Value", "Description"])
    output_sheet.append(["Output_Root", str(output), "fixture"])
    for title, code, source in (
        ("MDB_HO", "HO-FIX", ho),
        ("MDB_Branch", "BR-FIX", branch),
    ):
        sheet = book.create_sheet(title)
        sheet.append(["Active", "Code", "Description", "MDB_Path"])
        sheet.append(["Y", code, code, str(source)])
    book.save(path)
    return path, ho, branch


class _FixtureEngine:
    mode = "success"

    def run(self, **values):
        time.sleep(0.15)
        if self.mode == "failed":
            raise RuntimeError("fixture engine failure")
        folder = (
            Path(values["output_root"])
            / f"fixture-{values['workflow'].lower()}-{time.time_ns()}"
        )
        folder.mkdir(parents=True)
        txt = folder / "attendance.txt"
        report = folder / "attendance.xlsx"
        process_log = folder / "Process.log"
        summary = folder / "summary.json"
        txt.touch()
        report.touch()
        process_log.touch()
        summary.touch()
        return {
            "raw_log_count": 2,
            "paired_record_count": 1,
            "valid_record_count": 1,
            "anomaly_record_count": 0,
            "duplicate_removed_count": 0,
            "mdb_summary": [],
            "txt_result": {"generated_files": [{"file_path": str(txt)}]},
            "report_result": {"report_file": str(report)},
            "artifact_result": {
                "artifact_folder": str(folder),
                "process_log": str(process_log),
                "summary_json": str(summary),
            },
        }


def _wait(root, predicate, timeout: float = 5.0) -> int:
    deadline = time.monotonic() + timeout
    ticks = 0
    while time.monotonic() < deadline:
        root.update()
        ticks += 1
        if predicate():
            return ticks
        time.sleep(0.01)
    raise AssertionError("Timed out while pumping the Tk event loop.")


def test_fixture_attendance_manual_acceptance(tk_root, tmp_path: Path) -> None:
    data_root = tmp_path / "Data"
    database = get_database_path(data_root)
    database.parent.mkdir(parents=True)
    SchemaManager().initialize_database(database, "ui4-acceptance")
    output = tmp_path / "Output"
    with SQLiteConnectionFactory().connect(database) as connection:
        GlobalSettingsRepository(connection).save_global_settings(
            output_root=str(output),
            period_start="2026-07-01",
            period_end="2026-07-31",
        )
        connection.execute(
            "INSERT INTO attendance_settings ("
            "attendance_settings_id, use_global_output, use_global_period, "
            "split_txt_rows, generate_report_default, default_workflow, updated_at"
            ") VALUES (1, 1, 1, 10000, 1, 'HO', '2026-07-01')"
        )
    configuration, _, _ = _configuration(tmp_path / "Attendance-Fixture.xlsx", output)
    backend = FakeRegistryBackend()
    registry = StorageRegistryService(backend)
    registry.write_storage_pointer(data_root, database)
    services = build_default_app_services(
        tk_root,
        application_version=APP_VERSION,
        project_root=PROJECT_ROOT,
        registry=registry,
        default_data_root=data_root,
    )
    services.attendance_service = AttendanceService(
        services.storage_service,
        AttendanceAdapter(engine_class=_FixtureEngine),
    )
    confirmations = []
    opened = []
    services.dialog_service = type(
        "Dialog",
        (),
        {
            "select_file": lambda self, **values: configuration,
            "select_folder": lambda self, **values: output,
            "confirm": lambda self, title, message: (
                confirmations.append(message) or True
            ),
            "warning": lambda *args: None,
            "error": lambda *args: None,
        },
    )()
    services.file_system_service = type(
        "FileSystem",
        (),
        {"open_folder": lambda self, path: opened.append(Path(path)) or True},
    )()
    context = AppContext(
        PROJECT_ROOT,
        PROJECT_ROOT / "assets",
        APP_VERSION,
        logging.getLogger("ui4.fixture-acceptance"),
        app_services=services,
    )
    app = OASKUnifiedApp(tk_root, context=context)
    assert app.navigate("attendance")
    page = app.navigation._cache["attendance"]
    _wait(tk_root, lambda: page._defaults_loaded and not page._busy)
    assert page.global_output_var.get() and page.global_period_var.get()

    page.browse_configuration()
    page.global_output_var.set(False)
    page.global_period_var.set(False)
    page._apply_global_state()
    page.output_var.set(str(output))
    page.start_var.set("07/02/2026")
    page.end_var.set("07/03/2026")
    for workflow in ("HO", "BRANCH"):
        page.workflow_var.set(workflow)
        page.validate()
        _wait(tk_root, lambda: not page._busy)
        assert f"Workflow: {workflow}" in page.validation_summary.label.cget("text")

    page.workflow_var.set("HO")
    page.run_attendance()
    _wait(tk_root, lambda: page._running)
    responsive_ticks = _wait(tk_root, lambda: not page._running)
    assert page._last_result.success
    assert responsive_ticks > 2
    page.open_output()
    page.open_process_log()

    page.workflow_var.set("BRANCH")
    page.run_attendance()
    _wait(tk_root, lambda: page._running)
    _wait(
        tk_root,
        lambda: (
            "Running Attendance engine for BRANCH" in page.log_text.get("1.0", "end")
        ),
    )
    page.cancel()
    _wait(tk_root, lambda: not page._running)
    assert page._last_result.cancelled

    _FixtureEngine.mode = "failed"
    page.workflow_var.set("HO")
    page.run_attendance()
    _wait(tk_root, lambda: page._running)
    _wait(tk_root, lambda: not page._running)
    assert not page._last_result.success
    assert not page._last_result.cancelled
    assert "FAILED" in page.result_summary.label.cget("text")
    _FixtureEngine.mode = "success"

    assert app.navigate("history")
    _wait(
        tk_root,
        lambda: (
            app.navigation.active_page_id == "history"
            and not app.navigation._cache["history"]._busy
        ),
    )
    assert app.navigate("dashboard")
    _wait(
        tk_root,
        lambda: (
            app.navigation.active_page_id == "dashboard"
            and not app.navigation._cache["dashboard"]._busy
        ),
    )
    with SQLiteConnectionFactory().connect(database, read_only=True) as connection:
        statuses = [
            row[0]
            for row in connection.execute(
                "SELECT unified_status FROM job_history "
                "WHERE module_code='ATTENDANCE' ORDER BY job_pk"
            )
        ]
        file_count = connection.execute("SELECT COUNT(*) FROM job_files").fetchone()[0]
        table_count = connection.execute(
            "SELECT COUNT(*) FROM sqlite_master "
            "WHERE type='table' AND name NOT LIKE 'sqlite_%'"
        ).fetchone()[0]
    assert statuses == ["COMPLETED", "CANCELLED", "FAILED"]
    assert file_count == 10
    assert table_count == len(REQUIRED_TABLES)
    assert len(confirmations) == 3
    assert len(opened) == 2
    assert backend.write_count == 1  # Explicit test pointer setup only.
    assert backend.delete_count == 0
    assert app.close()
    services.task_runner.shutdown()
