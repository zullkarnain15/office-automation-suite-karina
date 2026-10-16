"""Sprint 4 tests for Device Master and standalone Excel integration."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from shared.finger_monitoring import (
    DeviceMasterService, FingerMonitoringRepository, FingerMonitoringStorageService,
)
from shared.finger_monitoring.device_master_service import HEADERS, SHEET_NAME


@dataclass
class _Status:
    data_root: Path


class _Storage:
    def __init__(self, root): self.root = root
    def resolve_status(self): return _Status(self.root)


def _service(tmp_path):
    status = FingerMonitoringStorageService(_Storage(tmp_path)).initialize()
    repository = FingerMonitoringRepository(status.database_path)
    return repository, DeviceMasterService(repository)


def _workbook(path: Path, rows, headers=HEADERS):
    workbook = Workbook(); sheet = workbook.active; sheet.title = SHEET_NAME
    sheet.append(headers)
    for row in rows:
        if headers == HEADERS and len(row) == 4:
            branch, ip, model, active = row
            row = ("", "", branch, ip, model, active, "")
        sheet.append(row)
    workbook.save(path); workbook.close()


def test_crud_soft_delete_reactivate_search_and_global_ip(tmp_path):
    repository, service = _service(tmp_path)
    first = service.add_device(
        "Jakarta", "10.10.0.1", unit="CF", three_lc="MDU", note="Primary"
    )
    with pytest.raises(Exception):
        service.add_device("Bandung", "10.10.0.1")
    service.edit_device(first, "Jakarta Selatan", "10.10.0.2", device_model="X1800")
    device = service.search(search="selatan")[0]
    assert device["device_model"] == "X1800"
    assert (device["unit"], device["three_lc"], device["note"]) == ("CF", "MDU", "Primary")
    service.soft_delete(first)
    assert service.search() == ()
    assert service.search(active=False)[0]["device_id"] == first
    service.reactivate(first)
    assert service.search(device_model="X1800")[0]["is_active"] == 1
    with pytest.raises(ValueError): service.add_device("X", "not-an-ip")
    assert repository.list_devices()[0]["device_id"] == first


def test_device_master_accepts_other_device_models(tmp_path):
    _repository, service = _service(tmp_path)

    device_id = service.add_device(
        "Jakarta", "10.10.0.3", device_model="SOLUTION UTIME 5000",
    )

    assert service.search()[0]["device_id"] == device_id
    assert service.search()[0]["device_model"] == "SOLUTION UTIME 5000"


def test_import_accepts_other_device_models(tmp_path):
    _repository, service = _service(tmp_path)
    source = tmp_path / "other-model.xlsx"
    _workbook(source, [("Other", "10.10.0.6", "Solution UTime 5000", "YES")])

    preview = service.preview_import(source)

    assert preview.can_commit
    assert preview.items[0].device_model == "SOLUTION UTIME 5000"
    service.commit_import(preview)
    assert service.search()[0]["device_model"] == "SOLUTION UTIME 5000"


def test_v4_to_v5_migration_preserves_monitoring_history_and_allows_new_model(tmp_path):
    repository, service = _service(tmp_path)
    device = service.add_device("History", "10.10.0.4", device_model="X1500")
    cycle = repository.create_cycle("CHECK_ALL", 1)
    repository.persist_check_outcome(cycle, device, "OFFLINE", observed_date="2026-10-12")
    with repository.factory.connect(repository.database_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute("UPDATE monitoring_metadata SET schema_version=4")
        connection.execute("PRAGMA user_version=4")

    assert repository.initialize() == 6
    assert repository.list_devices()[0]["device_id"] == device
    with repository.factory.connect(repository.database_path, read_only=True) as connection:
        assert connection.execute(
            "SELECT device_id FROM check_result WHERE cycle_id=?", (cycle,)
        ).fetchone()[0] == device
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        table_sql = connection.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='device_master'"
        ).fetchone()[0]
        assert "IN ('X1500', 'X1800')" not in table_sql
    new_device = service.add_device("Other", "10.10.0.5", device_model="X9900")
    assert new_device > device


@pytest.mark.parametrize("count", (10, 200, 1000))
def test_preview_and_commit_large_imports(tmp_path, count):
    repository, service = _service(tmp_path)
    source = tmp_path / f"devices-{count}.xlsx"
    rows = [(f"Branch {number}", f"10.{number // 254}.{number % 254}.1", "X1800" if number % 2 else "", "YES") for number in range(1, count + 1)]
    _workbook(source, rows)
    preview = service.preview_import(source)
    assert preview.insert_count == count and preview.update_count == preview.rejected_count == 0
    assert service.commit_import(preview)["inserted"] == count
    assert len(repository.list_devices()) == count


def test_preview_update_rejection_and_invalid_import_changes_nothing(tmp_path):
    repository, service = _service(tmp_path)
    device = service.add_device("Old", "10.20.0.1")
    source = tmp_path / "mixed.xlsx"
    _workbook(source, [
        ("New", "10.20.0.2", "X1500", "YES"),
        ("Updated", "10.20.0.1", "X1800", "YES"),
        ("Broken", "nope", "X1500", "YES"),
        ("Duplicate", "10.20.0.2", "X1500", "YES"),
        ("", "", "", ""),
    ])
    preview = service.preview_import(source)
    assert (preview.insert_count, preview.update_count, preview.rejected_count) == (0, 1, 3)
    with pytest.raises(ValueError): service.commit_import(preview)
    assert repository.list_devices()[0]["device_id"] == device
    assert repository.list_devices()[0]["branch_name"] == "Old"


def test_import_revalidates_revision_and_rolls_back_on_apply_failure(tmp_path, monkeypatch):
    repository, service = _service(tmp_path)
    source = tmp_path / "valid.xlsx"; _workbook(source, [("A", "10.30.0.1", "", "YES"), ("B", "10.30.0.2", "", "YES")])
    preview = service.preview_import(source)
    service.add_device("Other", "10.30.0.3")
    with pytest.raises(RuntimeError): service.commit_import(preview)
    assert len(repository.list_devices()) == 1
    preview = service.preview_import(source)
    calls = {"count": 0}
    original = service._apply_item
    def fail_second(connection, item):
        calls["count"] += 1
        original(connection, item)
        if calls["count"] == 2: raise RuntimeError("forced rollback")
    monkeypatch.setattr(service, "_apply_item", fail_second)
    with pytest.raises(RuntimeError): service.commit_import(preview)
    assert len(repository.list_devices()) == 1


def test_template_export_filters_and_copy_tsv(tmp_path):
    _repository, service = _service(tmp_path)
    first = service.add_device("Jakarta", "10.40.0.1", device_model="X1500", unit="CF", three_lc="MDU", note="First")
    second = service.add_device("Bandung", "10.40.0.2", device_model="X1800", unit="MF", three_lc="BAN", note="Second")
    template = service.generate_template(tmp_path)
    workbook = load_workbook(template, read_only=True)
    try:
        assert tuple(workbook[SHEET_NAME].iter_rows(values_only=True))[1:] == (
            ("CF", "MDU", "Example Branch", "192.168.1.10", "X1500", "YES", "OTO AMU + CEMAS + ARD"),
            ("MF", "BAN", "Example Branch", "192.168.1.11", "X1500", "YES", None),
        )
    finally:
        workbook.close()
    output = tmp_path / "filtered.xlsx"
    service.export(output, device_model="X1800")
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert rows[0] == HEADERS and rows[1][:3] == ("MF", "BAN", "Bandung") and rows[1][6] == "Second" and len(rows) == 2
    all_rows = service.search(active=None)
    assert service.copy_selected(all_rows, (first,)).split("\t")[2] == "Jakarta"
    assert service.copy_ip_addresses(all_rows) == "10.40.0.2\n10.40.0.1"
    assert service.copy_branch_names(all_rows) == "Bandung\nJakarta"
    assert second in {row["device_id"] for row in all_rows}


def test_import_persists_unit_3lc_and_note(tmp_path):
    _repository, service = _service(tmp_path)
    source = tmp_path / "device-master.xlsx"
    _workbook(source, [("CF", "MDU", "Example Branch", "10.41.0.1", "X1500", "YES", "OTO AMU + CEMAS + ARD")])

    preview = service.preview_import(source)
    assert preview.can_commit and preview.insert_count == 1
    service.commit_import(preview)
    row = service.search()[0]
    assert (row["unit"], row["three_lc"], row["note"]) == (
        "CF", "MDU", "OTO AMU + CEMAS + ARD"
    )


def test_history_foreign_keys_survive_edit_and_soft_delete(tmp_path):
    repository, service = _service(tmp_path)
    device = service.add_device("History", "10.50.0.1")
    cycle = repository.create_cycle("CHECK_ALL", 1)
    repository.persist_check_outcome(cycle, device, "OFFLINE", observed_date="2026-10-12")
    service.edit_device(device, "History Updated", "10.50.0.2")
    service.soft_delete(device)
    with repository.factory.connect(repository.database_path, read_only=True) as connection:
        assert connection.execute("SELECT COUNT(*) FROM check_result WHERE device_id=?", (device,)).fetchone()[0] == 1


def test_wrong_headers_are_rejected(tmp_path):
    _repository, service = _service(tmp_path)
    source = tmp_path / "wrong.xlsx"; _workbook(source, [("A", "10.1.1.1")], headers=("Branch", "IP"))
    with pytest.raises(ValueError, match="Header"):
        service.preview_import(source)


def test_v1_to_v2_migration_enforces_global_active_ip(tmp_path):
    repository, service = _service(tmp_path)
    with repository.factory.connect(repository.database_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute("DROP INDEX uq_device_master_active_ip")
        connection.execute("CREATE UNIQUE INDEX uq_device_master_active_branch_ip ON device_master(branch_name, ip_address) WHERE is_active=1")
        connection.execute("UPDATE monitoring_metadata SET schema_version=1")
        connection.execute("PRAGMA user_version=1")
    assert repository.initialize() == 6
    first = repository.create_device("A", "10.60.0.1")
    assert first > 0
    with pytest.raises(Exception):
        repository.create_device("B", "10.60.0.1")
