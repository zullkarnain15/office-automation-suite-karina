"""Configured schema-v1 database fixture created through the import API."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from shared.database import SchemaManager
from shared.database.exporting import build_configuration_template
from shared.database.importing import (
    ConfigImportCommitRequest,
    ConfigImportService,
    ImportMode,
)


@pytest.fixture
def configured_database(tmp_path: Path) -> Path:
    database = tmp_path / "configured.db"
    source = tmp_path / "source-config.xlsx"
    SchemaManager().initialize_database(database, "db2b-test")
    build_configuration_template(source)
    workbook = load_workbook(source)
    try:
        _set_setting(workbook["Global_Settings"], "output_root", r"C:\OAS-K\Output")
        _set_setting(workbook["Global_Settings"], "period_start", "2026-07-01")
        _set_setting(workbook["Global_Settings"], "period_end", "2026-07-31")
        _set_setting(workbook["Outlook_Settings"], "mailbox_smtp", "oas-k@example.com")
        _set_setting(workbook["Outlook_Settings"], "payroll_period", "07-2026")
        _set_setting(workbook["HRIS_Settings"], "hris_url", "https://hris.example.com")

        _append(
            workbook["Attendance_Sources"],
            ("HO", "HO-01", "Kantor Pusat", r"C:\Data\attendance.mdb", "TRUE", 1),
        )
        _append(
            workbook["Outlook_HO_Senders"],
            (
                "", "", "10001", "Operator", "operator@example.com",
                "20001", "Supervisor", "supervisor@example.com", "TRUE",
            ),
        )
        _append(
            workbook["Outlook_Subject_Rules"],
            ("HO", "REVISI ATTENDANCE", "TRUE"),
        )
        _append(
            workbook["Outlook_Attachment_Rules"],
            ("HO", ".xlsx", "TRUE"),
        )
        _append(
            workbook["Outlook_Validation_Rules"],
            ("SUBJECT_REQUIRED", "ALL", "TRUE", "TRUE"),
        )
        _append(
            workbook["Outlook_Reply_Templates"],
            (
                "SUCCESS", "SENDER", "SUCCESS", "Berhasil {PERIOD}",
                "Baris pertama\nBaris kedua {SENDER_NAME}", "TRUE",
            ),
        )
        _append(
            workbook["Outlook_Summary_Recipients"],
            ("TO", "hr@example.com", "TRUE", 1),
        )
        for row in (
            ("HO", 1, "001", "HO pertama", "TRUE"),
            ("HO", 2, "02", "HO kedua", "TRUE"),
            ("BRANCH", 1, "3", "Branch pertama", "TRUE"),
        ):
            _append(workbook["HRIS_Run_Controls"], row)
        _append(
            workbook["HRIS_Assisted_Steps"],
            (
                1, "Isi Run Control", "type", "RUN_CONTROL_ID",
                "assisted", "TRUE", 0.5, "Langkah assisted", "TRUE",
            ),
        )
        workbook.save(source)
    finally:
        workbook.close()

    service = ConfigImportService()
    preview = service.preview(database, [source])
    assert preview.can_commit, tuple(
        (issue.code, issue.message) for issue in preview.issues
    )
    result = service.commit(
        database,
        ConfigImportCommitRequest(
            preview=preview,
            modules=(
                "GLOBAL",
                "ATTENDANCE",
                "OUTLOOK_REVISI",
                "HRIS",
                "UTILITIES",
            ),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
            operator="db2b-test",
        ),
    )
    assert result.committed, result.module_results
    return database


def _set_setting(sheet: object, key: str, value: object) -> None:
    for row in range(4, sheet.max_row + 1):
        if sheet.cell(row, 1).value == key:
            sheet.cell(row, 2, value)
            return
    raise AssertionError(f"Missing setting key: {key}")


def _append(sheet: object, values: tuple[object, ...]) -> None:
    row = 4
    while any(sheet.cell(row, column).value is not None for column in range(1, len(values) + 1)):
        row += 1
    for column, value in enumerate(values, start=1):
        sheet.cell(row, column, value)
