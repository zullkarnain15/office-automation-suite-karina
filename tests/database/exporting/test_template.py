"""Official template structure, safety, and validation tests."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from shared.database.exporting import (
    build_configuration_template,
    validate_configuration_workbook,
)
from shared.database.exporting.constants import (
    DATA_HEADERS,
    SETTING_HEADERS,
    SHEET_ORDER,
)

PROJECT_ROOT = Path(__file__).resolve().parents[3]


def test_template_has_exact_sheet_order(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    result = build_configuration_template(output)

    assert result.sheet_names == SHEET_ORDER
    assert result.validation.is_valid


def test_all_unified_sheet_names_fit_excel_limit(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output, read_only=True)
    try:
        assert all(len(name) <= 31 for name in workbook.sheetnames)
        assert workbook.sheetnames[-1] == "Attachment_Consolidation"
        assert len(workbook.sheetnames[-1]) == 24
        assert (
            workbook["Attachment_Consolidation"]["A1"].value
            == "Attachment Consolidation Settings"
        )
    finally:
        workbook.close()


def test_official_template_artifact_is_present_and_valid() -> None:
    official = (
        PROJECT_ROOT
        / "config"
        / "templates"
        / "OAS-K_Configuration_Template.xlsx"
    )

    assert official.is_file()
    assert validate_configuration_workbook(official).is_valid


@pytest.mark.parametrize("sheet_name", SHEET_ORDER[1:])
def test_each_configuration_sheet_has_contract_headers(
    tmp_path: Path,
    sheet_name: str,
) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output, read_only=True)
    try:
        expected = DATA_HEADERS.get(sheet_name, SETTING_HEADERS)
        actual = tuple(
            workbook[sheet_name].cell(3, column).value
            for column in range(1, len(expected) + 1)
        )
        assert actual == expected
    finally:
        workbook.close()


def test_template_contains_no_formulas_macros_or_external_links(
    tmp_path: Path,
) -> None:
    output = tmp_path / "template.xlsx"
    result = build_configuration_template(output)

    assert result.validation.formula_count == 0
    assert result.validation.external_link_count == 0
    assert result.validation.has_macros is False


def test_template_has_freeze_panes_filters_and_wrapping(
    tmp_path: Path,
) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    try:
        for sheet_name in SHEET_ORDER[1:]:
            assert workbook[sheet_name].freeze_panes == "A4"
            assert workbook[sheet_name].auto_filter.ref
        assert (
            workbook["Outlook_Reply_Templates"]["E4"].alignment.wrap_text
            is True
        )
    finally:
        workbook.close()


def test_run_control_id_column_is_text_formatted(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    try:
        sheet = workbook["HRIS_Run_Controls"]
        assert sheet["C4"].number_format == "@"
        assert sheet["C500"].number_format == "@"
    finally:
        workbook.close()


def test_hris_template_uses_production_verification_defaults(
    tmp_path: Path,
) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        settings = {
            str(row[0] or ""): row[1]
            for row in workbook["HRIS_Settings"].iter_rows(
                min_row=4,
                values_only=True,
            )
        }
        assert settings["verification_wait_seconds"] == 2
        assert settings["verification_timeout_seconds"] == 30
        assert (
            settings["verification_success_texts"]
            == "Process Instance|Submitted|Queued"
        )
        assert (
            settings["verification_failure_texts"]
            == "Error|Invalid|Failed"
        )
    finally:
        workbook.close()


def test_boolean_workflow_and_active_validations_exist(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    try:
        attendance = workbook["Attendance_Sources"]
        formulas = {
            validation.formula1
            for validation in attendance.data_validations.dataValidation
        }
        assert '"HO,BRANCH"' in formulas
        assert '"TRUE,FALSE"' in formulas
        assert workbook["Attendance_Settings"].data_validations.count >= 3
    finally:
        workbook.close()


def test_guide_is_indonesian_and_warns_about_safe_defaults(
    tmp_path: Path,
) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output, read_only=True)
    try:
        values = "\n".join(
            str(cell.value or "")
            for row in workbook["Guide"].iter_rows()
            for cell in row
        )
        assert "Panduan Konfigurasi" in values
        assert "auto_reply_enabled=FALSE" in values
        assert "leading zero" in values
    finally:
        workbook.close()


def test_template_has_no_developer_path_default(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output, read_only=True)
    try:
        values = [
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert workbook["Global_Settings"]["B4"].value is None
        assert not any("Python Project" in value for value in values)
    finally:
        workbook.close()


def test_global_period_keys_are_not_duplicated_in_module_settings(
    tmp_path: Path,
) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output, read_only=True)
    try:
        for sheet_name in (
            "Attendance_Settings",
            "Outlook_Settings",
            "HRIS_Settings",
            "Comparison_Settings",
            "Attachment_Consolidation",
        ):
            keys = {
                workbook[sheet_name].cell(row, 1).value
                for row in range(4, workbook[sheet_name].max_row + 1)
            }
            assert "period_start" not in keys
            assert "period_end" not in keys
    finally:
        workbook.close()


def test_parent_creation_and_overwrite_are_explicit(tmp_path: Path) -> None:
    output = tmp_path / "missing" / "template.xlsx"
    with pytest.raises(FileNotFoundError):
        build_configuration_template(output)

    build_configuration_template(output, create_parent=True)
    with pytest.raises(FileExistsError):
        build_configuration_template(output)
    build_configuration_template(output, overwrite=True)


def test_validator_rejects_formula(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    workbook["Attendance_Sources"]["A4"] = "=1+1"
    workbook.save(output)
    workbook.close()

    result = validate_configuration_workbook(output)
    assert not result.is_valid
    assert result.formula_count == 1


def test_validator_rejects_missing_sheet(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    del workbook["Comparison_Settings"]
    workbook.save(output)
    workbook.close()

    result = validate_configuration_workbook(output)
    assert not result.is_valid
    assert any("Sheet order" in error for error in result.errors)


def test_validator_rejects_non_text_run_control_column(
    tmp_path: Path,
) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    workbook["HRIS_Run_Controls"]["C4"].number_format = "General"
    workbook.save(output)
    workbook.close()

    result = validate_configuration_workbook(output)
    assert not result.is_valid
    assert any("run_control_id" in error for error in result.errors)


def test_validator_rejects_shifted_workflow_validation(
    tmp_path: Path,
) -> None:
    output = tmp_path / "template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    validations = workbook[
        "Attendance_Sources"
    ].data_validations.dataValidation
    workflow = next(
        item for item in validations if item.formula1 == '"HO,BRANCH"'
    )
    workflow.sqref = "B4:B500"
    workbook.save(output)
    workbook.close()

    result = validate_configuration_workbook(output)
    assert not result.is_valid
    assert any("Attendance_Sources" in error for error in result.errors)
