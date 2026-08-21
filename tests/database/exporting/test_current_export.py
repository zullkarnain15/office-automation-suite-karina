"""Current configuration export and roundtrip acceptance tests."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from shared.database.exporting import export_current_configuration
from shared.database.importing import ChangeOperation, ConfigImportService
from shared.database.importing.models import WorkbookIdentity
from shared.database.importing.workbook_detector import detect_workbook


def test_current_export_is_unified_and_valid(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "current.xlsx"
    result = export_current_configuration(configured_database, output)

    assert result.validation.is_valid
    assert detect_workbook(output).identity is WorkbookIdentity.OAS_K_UNIFIED


def test_import_export_preview_roundtrip_is_all_unchanged(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "current.xlsx"
    export_current_configuration(configured_database, output)

    preview = ConfigImportService().preview(configured_database, [output])
    assert preview.can_commit
    assert preview.changes
    assert {change.operation for change in preview.changes} == {
        ChangeOperation.UNCHANGED
    }


def test_attachment_consolidation_settings_remain_mapped(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "current.xlsx"
    export_current_configuration(configured_database, output)

    preview = ConfigImportService().preview(configured_database, [output])
    utilities = next(
        mapped
        for mapped in preview.mapped_modules
        if mapped.module == "UTILITIES"
    )
    attachment = utilities.tables[
        "attachment_consolidation_settings"
    ][0]
    assert attachment["use_global_output"] == 1
    assert attachment["txt_max_lines"] == 10000


def test_export_preserves_leading_zero_run_control_ids(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "current.xlsx"
    export_current_configuration(configured_database, output)
    workbook = load_workbook(output)
    try:
        values = [
            workbook["HRIS_Run_Controls"].cell(row, 3).value
            for row in range(4, 7)
        ]
        formats = [
            workbook["HRIS_Run_Controls"].cell(row, 3).number_format
            for row in range(4, 7)
        ]
        assert values == ["3", "001", "02"]
        assert formats == ["@", "@", "@"]
    finally:
        workbook.close()


def test_export_preserves_multiline_reply_body(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "current.xlsx"
    export_current_configuration(configured_database, output)
    workbook = load_workbook(output)
    try:
        cell = workbook["Outlook_Reply_Templates"]["E4"]
        assert cell.value == "Baris pertama\nBaris kedua {SENDER_NAME}"
        assert cell.alignment.wrap_text is True
    finally:
        workbook.close()


def test_export_includes_non_secret_metadata(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "current.xlsx"
    export_current_configuration(configured_database, output)
    workbook = load_workbook(output, read_only=True)
    try:
        guide_text = "\n".join(
            str(cell.value or "")
            for row in workbook["Guide"].iter_rows()
            for cell in row
        )
        assert "database_uuid" in guide_text
        assert "schema_version" in guide_text
        assert "password" not in guide_text.casefold()
        metadata_keys = {
            str(workbook["Guide"].cell(row, 1).value or "").casefold()
            for row in range(1, workbook["Guide"].max_row + 1)
        }
        assert not metadata_keys & {
            "password",
            "secret",
            "token",
            "credential",
            "api_key",
        }
    finally:
        workbook.close()


def test_current_export_requires_valid_database(tmp_path: Path) -> None:
    with pytest.raises(Exception, match="Database"):
        export_current_configuration(
            tmp_path / "missing.db",
            tmp_path / "current.xlsx",
        )


def test_current_export_does_not_overwrite_by_default(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "current.xlsx"
    export_current_configuration(configured_database, output)
    with pytest.raises(FileExistsError):
        export_current_configuration(configured_database, output)
    export_current_configuration(
        configured_database,
        output,
        overwrite=True,
    )


def test_current_export_does_not_modify_database_bytes(
    configured_database: Path,
    tmp_path: Path,
) -> None:
    before = hashlib.sha256(configured_database.read_bytes()).hexdigest()
    export_current_configuration(
        configured_database,
        tmp_path / "current.xlsx",
    )
    after = hashlib.sha256(configured_database.read_bytes()).hexdigest()

    assert after == before
