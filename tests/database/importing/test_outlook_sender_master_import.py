"""Outlook Sender HO Master import regression tests."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from shared.database import SQLiteConnectionFactory, SchemaManager
from shared.database.exporting import build_configuration_template
from shared.database.exporting.current_config_exporter import (
    export_current_configuration,
)
from shared.database.importing import (
    ChangeOperation,
    ConfigImportCommitRequest,
    ConfigImportService,
    ImportMode,
)
from shared.database.importing.models import ConfigImportChange


def test_sender_ho_master_no_change_is_unchanged(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    workbook = _workbook(tmp_path, _sender_rows())

    changes = _sender_changes(outlook_sender_database, workbook)

    assert _counts(changes) == {
        ChangeOperation.INSERT: 0,
        ChangeOperation.UPDATE: 0,
        ChangeOperation.DELETE: 0,
    }


def test_sender_ho_master_modify_existing_by_email_key(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    rows = _sender_rows()
    rows[0]["sender_name"] = "Updated Sender 01"
    workbook = _workbook(tmp_path, rows)

    changes = _sender_changes(outlook_sender_database, workbook)

    assert _counts(changes)[ChangeOperation.UPDATE] == 1


def test_sender_ho_master_add_sender_dynamic_row(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    workbook = _workbook(tmp_path, [*_sender_rows(), _sender(13)])

    changes = _sender_changes(outlook_sender_database, workbook)

    assert _counts(changes)[ChangeOperation.INSERT] == 1
    assert _preview(outlook_sender_database, workbook).error_count == 0


def test_sender_ho_master_imports_exactly_1323_rows(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    workbook = _workbook(
        tmp_path,
        [_sender(index) for index in range(1, 1324)],
    )
    service = ConfigImportService()
    preview = service.preview(outlook_sender_database, [workbook])

    sender_changes = tuple(
        change
        for change in preview.changes
        if change.setting_scope == "outlook_sender_master"
    )
    assert preview.error_count == 0
    assert preview.can_commit is True
    assert _counts(sender_changes)[ChangeOperation.INSERT] == 1311

    result = service.commit(
        outlook_sender_database,
        ConfigImportCommitRequest(
            preview=preview,
            modules=("OUTLOOK_REVISI",),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )

    with SQLiteConnectionFactory().connect(
        outlook_sender_database,
        read_only=True,
    ) as connection:
        sender_count = connection.execute(
            "SELECT COUNT(*) FROM outlook_sender_master WHERE workflow = 'HO'"
        ).fetchone()[0]

    assert result.committed is True
    assert sender_count == 1323


def test_sender_ho_master_remove_sender_dynamic_row(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    workbook = _workbook(tmp_path, _sender_rows()[1:])

    changes = _sender_changes(outlook_sender_database, workbook)

    assert _counts(changes)[ChangeOperation.DELETE] == 1
    assert _preview(outlook_sender_database, workbook).error_count == 0


def test_sender_ho_master_add_and_remove(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    workbook = _workbook(tmp_path, [*_sender_rows()[1:], _sender(99)])

    changes = _sender_changes(outlook_sender_database, workbook)

    assert _counts(changes)[ChangeOperation.INSERT] == 1
    assert _counts(changes)[ChangeOperation.DELETE] == 1


def test_sender_ho_master_modify_add_and_remove(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    rows = _sender_rows()[1:]
    rows[0]["required_cc_email"] = "new.supervisor@example.com"
    workbook = _workbook(tmp_path, [*rows, _sender(99)])

    changes = _sender_changes(outlook_sender_database, workbook)
    counts = _counts(changes)

    assert counts[ChangeOperation.UPDATE] == 1
    assert counts[ChangeOperation.INSERT] == 1
    assert counts[ChangeOperation.DELETE] == 1


def test_sender_ho_master_reorder_rows_does_not_change_diff(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    workbook = _workbook(tmp_path, list(reversed(_sender_rows())))

    changes = _sender_changes(outlook_sender_database, workbook)

    assert _counts(changes) == {
        ChangeOperation.INSERT: 0,
        ChangeOperation.UPDATE: 0,
        ChangeOperation.DELETE: 0,
    }


def test_sender_ho_master_duplicate_business_key_blocks_false_delete(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    rows = [_sender(1), _sender(1, name="Duplicate Sender")]
    workbook = _workbook(tmp_path, rows)

    preview = _preview(outlook_sender_database, workbook)
    changes = _sender_changes(outlook_sender_database, workbook)

    assert any(issue.code == "OUTLOOK_SENDER_DUPLICATE" for issue in preview.issues)
    assert _counts(changes)[ChangeOperation.DELETE] == 0


def test_sender_ho_master_identical_duplicate_is_ignored_with_warning(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    rows = _sender_rows()
    workbook = _workbook(tmp_path, [rows[0], rows[0], *rows[1:]])

    preview = _preview(outlook_sender_database, workbook)
    changes = _sender_changes(outlook_sender_database, workbook)

    assert preview.can_commit is True
    assert preview.error_count == 0
    assert any(
        issue.code == "OUTLOOK_SENDER_DUPLICATE_IDENTICAL_IGNORED"
        and issue.severity.value == "WARNING"
        and issue.sheet == "Outlook_HO_Senders"
        and issue.row_number == 5
        for issue in preview.issues
    )
    assert _counts(changes) == {
        ChangeOperation.INSERT: 0,
        ChangeOperation.UPDATE: 0,
        ChangeOperation.DELETE: 0,
    }


def test_sender_ho_master_blank_trailing_rows_are_ignored(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    workbook = _workbook(tmp_path, _sender_rows(), blank_trailing_rows=5)

    changes = _sender_changes(outlook_sender_database, workbook)

    assert _counts(changes) == {
        ChangeOperation.INSERT: 0,
        ChangeOperation.UPDATE: 0,
        ChangeOperation.DELETE: 0,
    }


def test_sender_ho_master_invalid_dataset_blocks_false_delete(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    rows = [_sender(index) for index in range(1, 13)]
    for row in rows:
        row["sender_email"] = ""
    workbook = _workbook(tmp_path, rows)

    preview = _preview(outlook_sender_database, workbook)
    changes = _sender_changes(outlook_sender_database, workbook)

    assert any(issue.code == "ACTIVE_SENDER_EMAIL_MISSING" for issue in preview.issues)
    assert _counts(changes)[ChangeOperation.DELETE] == 0


def test_sender_ho_master_invalid_added_value_reports_excel_location(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    rows = [*_sender_rows(), *(_sender(index) for index in range(13, 25))]
    workbook = _workbook(tmp_path, rows)
    excel_row = 27
    excel_column = 9
    excel = load_workbook(workbook)
    try:
        excel["Outlook_HO_Senders"].cell(
            excel_row,
            excel_column,
            "AKTIF",
        )
        excel.save(workbook)
    finally:
        excel.close()

    preview = _preview(outlook_sender_database, workbook)
    issue = next(
        item
        for item in preview.issues
        if item.code == "OUTLOOK_SENDER_ROW_INVALID"
    )

    assert preview.can_commit is False
    assert issue.module == "OUTLOOK_REVISI"
    assert issue.sheet == "Outlook_HO_Senders"
    assert issue.row_number == excel_row
    assert issue.field == "is_active"


def test_sender_ho_master_apply_uses_preview_operations(
    outlook_sender_database: Path,
    tmp_path: Path,
) -> None:
    rows = _sender_rows()[1:]
    rows[0]["sender_name"] = "Updated Sender 02"
    rows.extend([_sender(13), _sender(14)])
    workbook = _workbook(tmp_path, rows)
    service = ConfigImportService()
    preview = service.preview(outlook_sender_database, [workbook])

    result = service.commit(
        outlook_sender_database,
        ConfigImportCommitRequest(
            preview=preview,
            modules=("OUTLOOK_REVISI",),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )

    with SQLiteConnectionFactory().connect(outlook_sender_database, read_only=True) as connection:
        senders = connection.execute(
            """
            SELECT sender_name, sender_email
            FROM outlook_sender_master
            WHERE workflow = 'HO'
            ORDER BY sender_email
            """
        ).fetchall()

    assert result.committed is True
    assert len(senders) == 13
    assert "sender01@example.com" not in {row["sender_email"] for row in senders}
    assert any(row["sender_name"] == "Updated Sender 02" for row in senders)
    assert {"sender13@example.com", "sender14@example.com"} <= {
        row["sender_email"] for row in senders
    }


def test_sender_ho_master_apply_rolls_back_on_failure(
    outlook_sender_database: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = _workbook(tmp_path, [*_sender_rows(), _sender(13)])
    service = ConfigImportService()
    preview = service.preview(outlook_sender_database, [workbook])
    original = service.transaction_service._insert_rows

    def fail_sender_insert(connection, table, rows):
        original(connection, table, rows)
        if table == "outlook_sender_master":
            raise RuntimeError("forced sender failure")

    monkeypatch.setattr(service.transaction_service, "_insert_rows", fail_sender_insert)
    result = service.commit(
        outlook_sender_database,
        ConfigImportCommitRequest(
            preview=preview,
            modules=("OUTLOOK_REVISI",),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )

    with SQLiteConnectionFactory().connect(outlook_sender_database, read_only=True) as connection:
        sender_count = connection.execute(
            "SELECT COUNT(*) FROM outlook_sender_master"
        ).fetchone()[0]
        failed_batches = connection.execute(
            "SELECT COUNT(*) FROM config_import_batches WHERE status = 'FAILED'"
        ).fetchone()[0]

    assert result.committed is False
    assert sender_count == 12
    assert failed_batches == 1


def test_sender_ho_master_export_edit_add_import(tmp_path: Path) -> None:
    database = _database(tmp_path)
    exported = tmp_path / "export-add.xlsx"
    export_current_configuration(database, exported, overwrite=True)
    _edit_sender_sheet(exported, lambda rows: [*rows, _sender(13)])

    changes = _sender_changes(database, exported)

    assert _counts(changes)[ChangeOperation.INSERT] == 1


def test_sender_ho_master_export_edit_remove_import(tmp_path: Path) -> None:
    database = _database(tmp_path)
    exported = tmp_path / "export-remove.xlsx"
    export_current_configuration(database, exported, overwrite=True)
    _edit_sender_sheet(exported, lambda rows: rows[1:])

    changes = _sender_changes(database, exported)

    assert _counts(changes)[ChangeOperation.DELETE] == 1


def test_sender_ho_master_export_edit_row_count_import(tmp_path: Path) -> None:
    database = _database(tmp_path)
    exported = tmp_path / "export-row-count.xlsx"
    export_current_configuration(database, exported, overwrite=True)
    _edit_sender_sheet(exported, lambda rows: [*rows[2:], _sender(13), _sender(14)])

    changes = _sender_changes(database, exported)
    counts = _counts(changes)

    assert counts[ChangeOperation.INSERT] == 2
    assert counts[ChangeOperation.DELETE] == 2


@pytest.fixture
def outlook_sender_database(tmp_path: Path) -> Path:
    return _database(tmp_path)


def _database(tmp_path: Path) -> Path:
    database = tmp_path / "sender-master.db"
    SchemaManager().initialize_database(database, "sender-master-test")
    timestamp = "2026-08-13T00:00:00"
    with SQLiteConnectionFactory().connect(database) as connection:
        connection.execute(
            """
            INSERT INTO outlook_settings (
                outlook_settings_id,
                use_global_output,
                use_global_period,
                integration_method,
                mailbox_smtp,
                source_folder,
                send_transport,
                smtp_port,
                smtp_timeout_seconds,
                save_smtp_copy_to_sent,
                processed_folder,
                auto_reply_enabled,
                send_mode,
                txt_max_lines,
                module_display_name,
                payroll_period,
                updated_at
            ) VALUES (1, 1, 1, 'OOM_COM', 'mailbox@example.com',
                'Inbox', 'OUTLOOK', 25, 30, 1, 'Deleted Items', 0,
                'DRAFT', 10000, 'Outlook Revisi', '08-2026', ?)
            """,
            (timestamp,),
        )
        for row in _sender_rows():
            connection.execute(
                """
                INSERT INTO outlook_sender_master (
                    workflow,
                    company_code,
                    branch_code,
                    sender_nik,
                    sender_name,
                    sender_email,
                    supervisor_nik,
                    supervisor_name,
                    required_cc_email,
                    is_active,
                    created_at,
                    updated_at
                ) VALUES ('HO', '', '', ?, ?, ?, ?, ?, ?, 1, ?, ?)
                """,
                (
                    row["sender_nik"],
                    row["sender_name"],
                    row["sender_email"],
                    row["supervisor_nik"],
                    row["supervisor_name"],
                    row["required_cc_email"],
                    timestamp,
                    timestamp,
                ),
            )
    return database


def _workbook(
    tmp_path: Path,
    senders: list[dict[str, Any]],
    *,
    blank_trailing_rows: int = 0,
) -> Path:
    path = tmp_path / "sender-master.xlsx"
    build_configuration_template(path)
    _edit_settings(path)
    _edit_sender_sheet(
        path,
        lambda _rows: senders,
        blank_trailing_rows=blank_trailing_rows,
    )
    return path


def _edit_settings(path: Path) -> None:
    workbook = load_workbook(path)
    try:
        for sheet_name, values in {
            "Global_Settings": {
                "output_root": r"C:\OAS-K\Output",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
            },
            "Outlook_Settings": {
                "mailbox_smtp": "mailbox@example.com",
                "payroll_period": "08-2026",
            },
            "HRIS_Settings": {
                "hris_url": "https://hris.example.com",
            },
        }.items():
            sheet = workbook[sheet_name]
            for row in range(4, sheet.max_row + 1):
                key = sheet.cell(row, 1).value
                if key in values:
                    sheet.cell(row, 2, values[key])
        workbook.save(path)
    finally:
        workbook.close()


def _edit_sender_sheet(
    path: Path,
    edit,
    *,
    blank_trailing_rows: int = 0,
) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook["Outlook_HO_Senders"]
        headers = [sheet.cell(3, column).value for column in range(1, 10)]
        rows = []
        for row_number in range(4, sheet.max_row + 1):
            values = [sheet.cell(row_number, column).value for column in range(1, 10)]
            if any(value is not None for value in values):
                rows.append(dict(zip(headers, values, strict=True)))
        updated_rows = edit(rows)
        for row_number in range(4, sheet.max_row + blank_trailing_rows + 10):
            for column in range(1, 10):
                sheet.cell(row_number, column).value = None
        for row_number, row in enumerate(updated_rows, start=4):
            sheet.cell(row_number, 1, row.get("company_code", ""))
            sheet.cell(row_number, 2, row.get("branch_code", ""))
            sheet.cell(row_number, 3, row.get("sender_nik"))
            sheet.cell(row_number, 4, row.get("sender_name"))
            sheet.cell(row_number, 5, row.get("sender_email"))
            sheet.cell(row_number, 6, row.get("supervisor_nik"))
            sheet.cell(row_number, 7, row.get("supervisor_name"))
            sheet.cell(row_number, 8, row.get("required_cc_email"))
            sheet.cell(row_number, 9, "TRUE" if row.get("is_active", 1) else "FALSE")
        if blank_trailing_rows:
            start = 4 + len(updated_rows)
            for row_number in range(start, start + blank_trailing_rows):
                sheet.cell(row_number, 1).value = None
                sheet.cell(row_number, 9).value = None
        workbook.save(path)
    finally:
        workbook.close()


def _sender_rows() -> list[dict[str, Any]]:
    return [_sender(index) for index in range(1, 13)]


def _sender(index: int, *, name: str | None = None) -> dict[str, Any]:
    return {
        "company_code": "",
        "branch_code": "",
        "sender_nik": f"{index:05d}",
        "sender_name": name or f"Sender {index:02d}",
        "sender_email": f"sender{index:02d}@example.com",
        "supervisor_nik": f"9{index:04d}",
        "supervisor_name": f"Supervisor {index:02d}",
        "required_cc_email": f"supervisor{index:02d}@example.com",
        "is_active": 1,
    }


def _preview(database: Path, workbook: Path):
    return ConfigImportService().preview(database, [workbook])


def _sender_changes(database: Path, workbook: Path) -> tuple[ConfigImportChange, ...]:
    preview = _preview(database, workbook)
    return tuple(
        change
        for change in preview.changes
        if change.setting_scope == "outlook_sender_master"
    )


def _counts(
    changes: tuple[ConfigImportChange, ...],
) -> dict[ChangeOperation, int]:
    return {
        operation: sum(change.operation == operation for change in changes)
        for operation in (
            ChangeOperation.INSERT,
            ChangeOperation.UPDATE,
            ChangeOperation.DELETE,
        )
    }
