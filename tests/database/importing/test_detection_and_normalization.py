"""Workbook detection, reader safety, and normalization tests."""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from shared.database.importing import WorkbookIdentity, detect_workbook
from shared.database.importing.legacy import map_hris, map_outlook
from shared.database.importing.normalizer import (
    normalize_boolean,
    normalize_date,
    normalize_workflow,
)
from shared.database.importing.workbook_reader import read_workbook
from tests.database.importing.conftest import (
    ATTENDANCE_WORKBOOK,
    HRIS_WORKBOOK,
    OUTLOOK_WORKBOOK,
)


@pytest.mark.parametrize(
    ("path", "identity"),
    [
        (ATTENDANCE_WORKBOOK, WorkbookIdentity.ATTENDANCE_LEGACY),
        (OUTLOOK_WORKBOOK, WorkbookIdentity.OUTLOOK_REVISI_LEGACY),
        (HRIS_WORKBOOK, WorkbookIdentity.HRIS_LEGACY),
    ],
)
def test_detects_actual_legacy_workbooks(
    path: Path,
    identity: WorkbookIdentity,
) -> None:
    assert detect_workbook(path).identity == identity


def test_unknown_and_ambiguous_workbooks_are_detected(
    tmp_path: Path,
) -> None:
    unknown = tmp_path / "unknown.xlsx"
    workbook = Workbook()
    workbook.active.title = "Random"
    workbook.save(unknown)
    workbook.close()

    ambiguous = tmp_path / "ambiguous.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in (
        "General",
        "MDB_HO",
        "MDB_Branch",
        "Output",
        "Reference",
        "Run_Control",
        "Browser",
        "Upload",
    ):
        workbook.create_sheet(name)
    workbook.save(ambiguous)
    workbook.close()

    assert detect_workbook(unknown).identity == WorkbookIdentity.UNKNOWN
    assert detect_workbook(ambiguous).identity == WorkbookIdentity.AMBIGUOUS


def test_actual_workbook_is_read_without_changing_hash() -> None:
    before = _sha256(ATTENDANCE_WORKBOOK)
    data, issues = read_workbook(ATTENDANCE_WORKBOOK)
    after = _sha256(ATTENDANCE_WORKBOOK)

    assert data.detection.identity == WorkbookIdentity.ATTENDANCE_LEGACY
    assert issues == ()
    assert before == after


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("TRUE", 1),
        ("YES", 1),
        ("Y", 1),
        (1, 1),
        ("FALSE", 0),
        ("NO", 0),
        ("N", 0),
        (0, 0),
    ],
)
def test_boolean_normalization(value: object, expected: int) -> None:
    assert normalize_boolean(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [("HO", "HO"), ("Branch", "BRANCH"), ("branch", "BRANCH")],
)
def test_workflow_normalization(value: str, expected: str) -> None:
    assert normalize_workflow(value) == expected


def test_date_normalization_and_ambiguous_rejection() -> None:
    assert normalize_date(date(2026, 7, 20)) == "2026-07-20"
    assert normalize_date("03/30/2026", legacy_order="MDY") == "2026-03-30"
    with pytest.raises(ValueError, match="Ambiguous"):
        normalize_date("03/04/2026")


def test_hris_leading_zero_and_outlook_multiline_are_preserved() -> None:
    hris_data, _ = read_workbook(HRIS_WORKBOOK)
    hris = map_hris(hris_data)
    identifiers = [
        row["run_control_id"] for row in hris.tables["hris_run_controls"]
    ]

    outlook_data, _ = read_workbook(OUTLOOK_WORKBOOK)
    outlook = map_outlook(outlook_data)
    bodies = [
        row["body_template"]
        for row in outlook.tables["outlook_reply_templates"]
    ]

    assert identifiers[:3] == ["001", "02", "3"]
    assert any("\n\n" in body for body in bodies)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
