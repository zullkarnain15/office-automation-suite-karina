from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from shared.database import SQLiteConnectionFactory
from shared.database.exporting import build_configuration_template, export_current_configuration
from shared.database.exporting.constants import SHEET_ORDER
from shared.database.importing import (
    ConfigImportCommitRequest,
    ConfigImportService,
    ImportMode,
)
from shared.database.importing.models import WorkbookIdentity
from shared.database.importing.workbook_detector import detect_workbook


def test_template_contains_att_data_repair_sheet_and_defaults(tmp_path: Path) -> None:
    output = tmp_path / "template.xlsx"
    result = build_configuration_template(output)

    workbook = load_workbook(output, data_only=True)
    try:
        assert result.sheet_names == SHEET_ORDER
        assert workbook.sheetnames[-1] == "Att_Data_Repair"
        settings = _settings(workbook["Att_Data_Repair"])
        assert settings["enabled"] == "TRUE"
        assert settings["minimum_duration_minutes"] == 61
        assert settings["weekday_default_in"] == "09:30"
        assert settings["saturday_missing_out_default"] == "11:00"
        assert settings["midnight_time_out_default"] == "23:59"
        assert settings["txt_max_rows"] == 10000
        assert settings["generate_txt"] == "TRUE"
        assert settings["generate_excel_report"] == "TRUE"
    finally:
        workbook.close()


def test_blank_template_preview_reports_required_fields_without_crashing(
    database_path: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "blank-template.xlsx"
    build_configuration_template(output)

    preview = ConfigImportService().preview(database_path, [output])

    assert not preview.can_commit
    codes = {issue.code for issue in preview.issues}
    assert "OUTLOOK_MAILBOX_INVALID" in codes
    assert "OUTLOOK_PAYROLL_PERIOD_INVALID" in codes
    assert "HRIS_URL_INVALID" in codes


def test_v3_att_data_repair_sheet_imports_missing_v4_keys_as_defaults(
    database_path: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "schema-v3-template.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    _fill_required(workbook)
    _delete_setting(workbook["Att_Data_Repair"], "saturday_missing_out_default")
    _delete_setting(workbook["Att_Data_Repair"], "midnight_time_out_default")
    workbook.save(output)
    workbook.close()

    service = ConfigImportService()
    preview = service.preview(database_path, [output])

    assert preview.can_commit
    result = service.commit(
        database_path,
        ConfigImportCommitRequest(
            preview=preview,
            modules=("UTILITIES",),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )
    assert result.committed
    with SQLiteConnectionFactory().connect(database_path, read_only=True) as connection:
        row = connection.execute(
            "SELECT saturday_missing_out_default, midnight_time_out_default "
            "FROM att_data_repair_settings WHERE att_data_repair_settings_id = 1"
        ).fetchone()
    assert tuple(row) == ("11:00", "23:59")


def test_old_unified_workbook_without_att_data_repair_imports_defaults(
    database_path: Path,
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "old.xlsx"
    build_configuration_template(workbook_path)
    workbook = load_workbook(workbook_path)
    del workbook["Att_Data_Repair"]
    _fill_required(workbook)
    workbook.save(workbook_path)
    workbook.close()

    assert detect_workbook(workbook_path).identity is WorkbookIdentity.OAS_K_UNIFIED
    service = ConfigImportService()
    preview = service.preview(database_path, [workbook_path])

    assert preview.can_commit
    assert any(issue.code == "ATT_DATA_REPAIR_DEFAULTS_APPLIED" for issue in preview.issues)
    result = service.commit(
        database_path,
        ConfigImportCommitRequest(
            preview=preview,
            modules=("UTILITIES",),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )
    assert result.committed
    with SQLiteConnectionFactory().connect(database_path, read_only=True) as connection:
        row = connection.execute(
            "SELECT minimum_duration_minutes, generate_txt, generate_excel_report "
            "FROM att_data_repair_settings"
        ).fetchone()
    assert tuple(row) == (61, 1, 1)


def test_new_workbook_import_validation_and_failed_import_preserves_db(
    database_path: Path,
    tmp_path: Path,
) -> None:
    valid = tmp_path / "valid.xlsx"
    invalid = tmp_path / "invalid.xlsx"
    build_configuration_template(valid)
    workbook = load_workbook(valid)
    _fill_required(workbook)
    _set_setting(workbook["Att_Data_Repair"], "minimum_duration_minutes", 75)
    _set_setting(workbook["Att_Data_Repair"], "weekday_default_in", "08:00")
    _set_setting(workbook["Att_Data_Repair"], "weekday_default_out", "16:30")
    _set_setting(workbook["Att_Data_Repair"], "generate_txt", "FALSE")
    workbook.save(valid)
    workbook.close()

    service = ConfigImportService()
    preview = service.preview(database_path, [valid])
    assert preview.can_commit
    service.commit(
        database_path,
        ConfigImportCommitRequest(
            preview=preview,
            modules=("UTILITIES",),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )
    with SQLiteConnectionFactory().connect(database_path, read_only=True) as connection:
        before = dict(connection.execute("SELECT * FROM att_data_repair_settings").fetchone())

    workbook = load_workbook(valid)
    _set_setting(workbook["Att_Data_Repair"], "weekday_default_in", "9;30")
    workbook.save(invalid)
    workbook.close()

    bad_preview = service.preview(database_path, [invalid])
    assert not bad_preview.can_commit
    assert any(issue.code == "WORKBOOK_MAPPING_FAILED" for issue in bad_preview.issues)
    with SQLiteConnectionFactory().connect(database_path, read_only=True) as connection:
        after = dict(connection.execute("SELECT * FROM att_data_repair_settings").fetchone())
    assert after == before


def test_export_import_roundtrip_preserves_att_data_repair_settings(
    database_path: Path,
    tmp_path: Path,
) -> None:
    seed = tmp_path / "seed.xlsx"
    build_configuration_template(seed)
    workbook = load_workbook(seed)
    _fill_required(workbook)
    workbook.save(seed)
    workbook.close()

    service = ConfigImportService()
    seed_preview = service.preview(database_path, [seed])
    assert seed_preview.can_commit
    service.commit(
        database_path,
        ConfigImportCommitRequest(
            preview=seed_preview,
            modules=("GLOBAL", "OUTLOOK_REVISI", "HRIS", "UTILITIES"),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )

    with SQLiteConnectionFactory().connect(database_path) as connection:
        connection.execute(
            """
            UPDATE att_data_repair_settings
            SET minimum_duration_minutes = 88,
                weekday_default_in = '08:15',
                weekday_default_out = '16:45',
                saturday_missing_out_default = '11:10',
                midnight_time_out_default = '23:58',
                txt_max_rows = 1234,
                generate_txt = 0,
                generate_excel_report = 1
            WHERE att_data_repair_settings_id = 1
            """
        )
    output = tmp_path / "current.xlsx"
    export_current_configuration(database_path, output)
    workbook = load_workbook(output, data_only=True)
    try:
        exported = _settings(workbook["Att_Data_Repair"])
    finally:
        workbook.close()
    assert exported["minimum_duration_minutes"] == 88
    assert exported["weekday_default_in"] == "08:15"
    assert exported["saturday_missing_out_default"] == "11:10"
    assert exported["midnight_time_out_default"] == "23:58"
    assert exported["generate_txt"] == "FALSE"

    new_db = tmp_path / "roundtrip.db"
    from shared.database import SchemaManager

    SchemaManager().initialize_database(new_db, "roundtrip")
    preview = service.preview(new_db, [output])
    assert preview.can_commit
    service.commit(
        new_db,
        ConfigImportCommitRequest(
            preview=preview,
            modules=("UTILITIES",),
            mode=ImportMode.REPLACE_MODULE_CONFIGURATION,
            confirmed=True,
        ),
    )
    with SQLiteConnectionFactory().connect(new_db, read_only=True) as connection:
        row = connection.execute(
            "SELECT minimum_duration_minutes, weekday_default_in, "
            "saturday_missing_out_default, midnight_time_out_default, txt_max_rows, "
            "generate_txt, generate_excel_report FROM att_data_repair_settings"
        ).fetchone()
    assert tuple(row) == (88, "08:15", "11:10", "23:58", 1234, 0, 1)


def test_import_preview_reports_generate_flags_both_false(
    database_path: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "invalid-flags.xlsx"
    build_configuration_template(output)
    workbook = load_workbook(output)
    _fill_required(workbook)
    _set_setting(workbook["Att_Data_Repair"], "generate_txt", "FALSE")
    _set_setting(workbook["Att_Data_Repair"], "generate_excel_report", "FALSE")
    workbook.save(output)
    workbook.close()

    preview = ConfigImportService().preview(database_path, [output])

    assert not preview.can_commit
    assert any(issue.code == "ATT_DATA_REPAIR_OUTPUT_DISABLED" for issue in preview.issues)


def _fill_required(workbook: Workbook) -> None:
    _set_setting(workbook["Global_Settings"], "output_root", r"C:\OAS-K\Output")
    _set_setting(workbook["Global_Settings"], "period_start", "2026-08-01")
    _set_setting(workbook["Global_Settings"], "period_end", "2026-08-31")
    _set_setting(workbook["Outlook_Settings"], "mailbox_smtp", "oas-k@example.com")
    _set_setting(workbook["Outlook_Settings"], "payroll_period", "08-2026")
    _set_setting(workbook["HRIS_Settings"], "hris_url", "https://hris.example.com")


def _set_setting(sheet, key: str, value: object) -> None:
    for row in range(4, sheet.max_row + 1):
        if sheet.cell(row, 1).value == key:
            sheet.cell(row, 2, value)
            return
    raise AssertionError(f"Missing setting key: {key}")


def _delete_setting(sheet, key: str) -> None:
    for row in range(4, sheet.max_row + 1):
        if sheet.cell(row, 1).value == key:
            sheet.delete_rows(row)
            return
    raise AssertionError(f"Missing setting key: {key}")


def _settings(sheet) -> dict[str, object]:
    return {
        str(row[0]): row[1]
        for row in sheet.iter_rows(min_row=4, values_only=True)
        if row[0]
    }
