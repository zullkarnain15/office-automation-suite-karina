from __future__ import annotations

import json
from datetime import date, datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from utilities.att_data_repair.artifacts import (
    AttDataRepairJobRequest,
    JobPaths,
    TxtArtifact,
)
from utilities.att_data_repair.constants import (
    INVALID_RECORDS_COLUMNS,
    REPORT_SHEET_ORDER,
    VALID_RECORDS_COLUMNS,
)
from utilities.att_data_repair.engine import AttDataRepairEngine
from utilities.att_data_repair.models import (
    AttDataRepairAnalysisResult,
    AttDataRepairRequest,
    FinalRecord,
    ReportAlreadyExistsError,
    RepairChange,
    SourceRecord,
)
from utilities.att_data_repair.report_writer import AttDataRepairReportWriter
from utilities.att_data_repair.statuses import ChangeCode, FinalStatus, JobStatus
from utilities.att_data_repair.txt_writer import AttDataRepairTxtWriter


class FixedClock:
    def __call__(self) -> datetime:
        return datetime(2026, 8, 3, 9, 0, 0)


class CodeProvider:
    def __call__(self) -> int:
        return 4837


def _default_valid(**overrides):
    row = {
        "No": 1,
        "Source_File": "revision.xlsx",
        "Relative_Path": "nested/revision.xlsx",
        "Source_Row": 2,
        "Workflow": "HO",
        "NIK": "000001234",
        "Date_In": "08/03/2026",
        "Time_In": "09:30",
        "Date_Out": "08/03/2026",
        "Time_Out": "17:00",
        "Output_TXT": "Outlook_Revisi_HO_001.txt",
        "Status": "VALID",
    }
    row.update(overrides)
    return row


def _default_invalid(**overrides):
    row = {
        "No": 7,
        "Source_File": "invalid.xlsx",
        "Relative_Path": "bad/invalid.xlsx",
        "Source_Row": 9,
        "Workflow": "Branch",
        "NIK": "000000888",
        "Date_In": "08/09/2026",
        "Time_In": "bad",
        "Date_Out": "08/09/2026",
        "Time_Out": "#REF!",
        "Status_Code": "TIME_FORMAT",
        "Reason": "Invalid time from source.",
        "Raw_Value": "raw invalid row",
    }
    row.update(overrides)
    return row


def _write_source(path: Path, *, valid_rows=(), invalid_rows=()) -> None:
    workbook = Workbook()
    workbook.active.title = "Valid_Records"
    valid = workbook["Valid_Records"]
    invalid = workbook.create_sheet("Invalid_Records")
    valid.append(VALID_RECORDS_COLUMNS)
    invalid.append(INVALID_RECORDS_COLUMNS)
    for item in valid_rows:
        valid.append([item.get(column) for column in VALID_RECORDS_COLUMNS])
    for item in invalid_rows:
        invalid.append([item.get(column) for column in INVALID_RECORDS_COLUMNS])
    workbook.save(path)
    workbook.close()


def _request(path: Path) -> AttDataRepairRequest:
    return AttDataRepairRequest(path, date(2026, 8, 1), date(2026, 8, 31))


def _fixture_source(path: Path) -> None:
    _write_source(
        path,
        valid_rows=(
            _default_valid(No=1, Workflow="HO", NIK="000001234"),
            _default_valid(
                No=2,
                Workflow="HO",
                NIK="000001234",
                Time_In="09:00",
                Time_Out="09:30",
                Date_Out="08/04/2026",
            ),
            _default_valid(
                No=3,
                Workflow="Branch",
                NIK="000000777",
                Time_In="7;20",
                Time_Out="7.40",
            ),
            _default_valid(No=4, Workflow="HO", NIK="#REF!"),
            _default_valid(No=5, Workflow="HO", NIK="000001234", Date_In="bad date"),
            _default_valid(No=6, Workflow="BAD", NIK="000009999"),
        ),
        invalid_rows=(
            _default_invalid(),
        ),
    )


def _run_fixture(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    _fixture_source(source)
    before = source.read_bytes()
    engine = AttDataRepairEngine(
        txt_writer=AttDataRepairTxtWriter(CodeProvider()),
        clock=FixedClock(),
    )
    result = engine.run_job(AttDataRepairJobRequest(_request(source), tmp_path))
    assert source.read_bytes() == before
    return source, result


def _rows(sheet):
    return list(sheet.iter_rows(values_only=True))


def _row_by_value(sheet, value):
    for row in _rows(sheet):
        if row and row[0] == value:
            return row
    raise AssertionError(f"Row not found: {value}")


def _as_date(value):
    return value.date() if isinstance(value, datetime) else value


def test_report_workbook_contract_and_guide_status(tmp_path: Path) -> None:
    source, result = _run_fixture(tmp_path)

    assert result.status == JobStatus.PARTIAL_SUCCESS
    assert result.report_artifact is not None
    assert result.report_artifact.file_name == "Att_Data_Repair_Report_2026-08-03_01.xlsx"
    assert result.report_artifact.file_path == result.paths.report_folder / result.report_artifact.file_name
    assert "4837" not in result.report_artifact.file_name
    assert not list(result.paths.report_folder.glob("*.tmp.xlsx"))
    assert not (result.paths.job_folder / "Original").exists()

    workbook = load_workbook(result.report_artifact.file_path, data_only=True)
    try:
        assert workbook.sheetnames == list(REPORT_SHEET_ORDER)
        assert workbook.active.title == "Guide_Status"
        assert workbook.properties.title == "Att Data Repair Report"
        assert workbook.properties.subject == "OAS-K Attendance Data Repair Audit Report"
        assert workbook.properties.creator == "Office Automation Suite - Karina"
        guide_text = "\n".join(str(value) for row in _rows(workbook["Guide_Status"]) for value in row if value)
        for expected in (
            "Job ID",
            "2026-08-03_01",
            "08/01/2026",
            str(source),
            "NIK wajib tersedia.",
            "Leading zero dipertahankan.",
            "Date_In menjadi tanggal utama.",
            "Senin-Jumat: 09:30-17:00.",
            "Sabtu: pasangan default tetap 09:30-12:05.",
            "Minggu hanya diproses jika merupakan hari terakhir bulan.",
            "Maksimum 10.000 baris per TXT.",
            "Att_Data_Repair_[Workflow]-[NNN]_[RRRR].txt.",
            "Workbook sumber tidak pernah diubah.",
            "Tidak ada salinan sumber pada folder output.",
            "Tidak ada folder Original",
            "VALID_UNCHANGED",
            "REPAIRED",
            "TIME_FORMAT_NORMALIZED",
            "DUPLICATE_RECORD",
            "SUNDAY_NOT_MONTH_END",
            "INVALID_NIK",
            "PARTIAL_SUCCESS",
        ):
            assert expected in guide_text
    finally:
        workbook.close()


def test_process_summary_and_summary_json_report_metadata(tmp_path: Path) -> None:
    _source, result = _run_fixture(tmp_path)

    workbook = load_workbook(result.report_artifact.file_path, data_only=True)
    try:
        summary = workbook["Process_Summary"]
        assert _row_by_value(summary, "Job Status")[1] == "PARTIAL_SUCCESS"
        assert _row_by_value(summary, "Source Records Total")[1] == 7
        assert _row_by_value(summary, "Source Valid_Records")[1] == 6
        assert _row_by_value(summary, "Source Invalid_Records")[1] == 1
        assert _row_by_value(summary, "Final Records")[1] == 3
        assert _row_by_value(summary, "Changed Records")[1] == 2
        assert _row_by_value(summary, "Anomaly Records")[1] == 4
        assert _row_by_value(summary, "TXT Files")[1] == 2
        assert _row_by_value(summary, "TXT Rows")[1] == 3
        assert any(
            row[:2] == ("INVALID_NIK", 1)
            and "panjang/polanya tidak valid" in row[2]
            for row in _rows(summary)
        )
        assert any(row[:2] == ("DATE_OUT_ALIGNED_TO_DATE_IN", 1) for row in _rows(summary))
        assert _row_by_value(summary, "saturday_missing_out_default")[1] == "11:00"
        assert _row_by_value(summary, "midnight_time_out_default")[1] == "23:59"
        assert any(row[:4] == ("TXT", "HO", "Att_Data_Repair_HO-001_4837.txt", 2) for row in _rows(summary))
        assert any(row[:3] == ("EXCEL_REPORT", None, result.report_artifact.file_name) for row in _rows(summary))
    finally:
        workbook.close()

    payload = json.loads(result.paths.summary_json.read_text(encoding="utf-8"))
    assert payload["report"]["generated"] is True
    assert payload["report"]["file_name"] == result.report_artifact.file_name
    assert payload["report"]["file_size_bytes"] == result.report_artifact.file_size_bytes
    assert payload["report"]["sheet_count"] == 8
    assert payload["duration_seconds"] == 0.0
    assert payload["duration_minutes"] == 0.0
    assert [item["name"] for item in payload["report"]["sheets"]] == list(REPORT_SHEET_ORDER)

    log = result.paths.process_log.read_text(encoding="utf-8")
    assert "Report generation started." in log
    assert f"Report generated: {result.report_artifact.file_name}" in log
    assert "Report size:" in log
    assert "Job duration: 0.00 seconds (0.00 minutes)" in log


def test_report_data_sheets_values_formats_and_styles(tmp_path: Path) -> None:
    _source, result = _run_fixture(tmp_path)

    workbook = load_workbook(result.report_artifact.file_path, data_only=True)
    try:
        summary_employee = workbook["Summary_Per_Karyawan"]
        employee_rows = _rows(summary_employee)[1:]
        assert employee_rows[0][:11] == (
            1,
            "HO",
            "000001234",
            3,
            2,
            1,
            1,
            1,
            employee_rows[0][8],
            employee_rows[0][9],
            "PARTIAL",
        )
        assert _as_date(employee_rows[0][8]) == date(2026, 8, 3)
        assert _as_date(employee_rows[0][9]) == date(2026, 8, 3)
        invalid_employee = next(row for row in employee_rows if row[1:3] == ("Invalid", "000009999"))
        assert invalid_employee[3:8] == (1, 0, 0, 0, 1)
        assert invalid_employee[10] == "EXCLUDED"
        assert all(row[2] != "#REF!" for row in employee_rows)
        assert summary_employee["C2"].number_format == "@"

        final = workbook["Final_Records"]
        assert [row[1] for row in _rows(final)[1:]] == [
            "ADR-000001",
            "ADR-000002",
            "ADR-000003",
        ]
        assert final["D2"].value == "000001234"
        assert final["D2"].number_format == "@"
        assert _as_date(final["E2"].value) == date(2026, 8, 3)
        assert final["E2"].number_format == "mm/dd/yyyy"
        assert final["F2"].value == time(9, 30)
        assert final["F2"].number_format == "hh:mm"
        assert final["Q2"].value == "Att_Data_Repair_HO-001_4837.txt"
        assert final["A1"].border.left.style == "thin"
        assert final["A2"].border.left.style is None
        assert final["P2"].fill.fill_type == "solid"

        changed = workbook["Changed_Records"]
        changed_rows = _rows(changed)[1:]
        assert _rows(changed)[0][3:5] == ("Original_NIK", "Final_NIK")
        assert [row[1] for row in changed_rows] == [
            "ADR-000002",
            "ADR-000003",
        ]
        branch_row = next(row for row in changed_rows if row[1] == "ADR-000003")
        assert branch_row[3:5] == ("000000777", "000000777")
        assert branch_row[11] == "7;20"
        assert branch_row[15] == "7.40"
        assert "TIME_FORMAT_NORMALIZED" in branch_row[18]
        assert branch_row[17] >= 2
        assert changed["A2"].border.left.style is None

        anomaly = workbook["Anomaly"]
        anomaly_rows = _rows(anomaly)[1:]
        assert [row[1] for row in anomaly_rows] == [
            "ADR-000004",
            "ADR-000005",
            "ADR-000006",
            "ADR-000007",
        ]
        assert "TXT_File_Name" not in _rows(anomaly)[0]
        assert any(row[16] == "INVALID_WORKFLOW" and "Workflow" in row[18] for row in anomaly_rows)

        change_log = workbook["Change_Log"]
        change_rows = _rows(change_log)[1:]
        assert all(row[1] != "ADR-000001" for row in change_rows)
        assert [row[11] for row in change_rows if row[1] == "ADR-000002"][0] == "DATE_OUT_ALIGNED_TO_DATE_IN"

        inventory = workbook["Source_Inventory"]
        inventory_rows = _rows(inventory)[1:]
        assert [row[4] for row in inventory_rows] == ["Valid_Records", "Invalid_Records"]
        assert [row[5] for row in inventory_rows] == [6, 1]
        assert inventory_rows[0][2] == Path(result.analysis_result.source_records[0].source_workbook).stat().st_size
        assert isinstance(inventory_rows[0][3], datetime)

        for sheet_name in ("Summary_Per_Karyawan", "Final_Records", "Changed_Records", "Anomaly", "Change_Log", "Source_Inventory"):
            sheet = workbook[sheet_name]
            assert sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref
            assert sheet["A1"].font.bold
    finally:
        workbook.close()


def test_report_audits_nik_repair_without_styling_body_borders(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    _write_source(
        source,
        valid_rows=(_default_valid(NIK="0000031815"),),
    )
    result = AttDataRepairEngine(
        txt_writer=AttDataRepairTxtWriter(CodeProvider()),
        clock=FixedClock(),
    ).run_job(AttDataRepairJobRequest(_request(source), tmp_path))

    workbook = load_workbook(result.report_artifact.file_path, data_only=True)
    try:
        changed = workbook["Changed_Records"]
        row = _rows(changed)[1]
        assert row[3:5] == ("0000031815", "20000031815")
        assert changed["D2"].number_format == "@"
        assert changed["E2"].number_format == "@"
        assert changed["D2"].border.left.style is None
        assert changed["E2"].border.left.style is None

        nik_change = next(
            row for row in _rows(workbook["Change_Log"])[1:] if row[8] == "NIK"
        )
        assert nik_change[9:12] == (
            "0000031815",
            "20000031815",
            "NIK_LEADING_TWO_RESTORED",
        )
    finally:
        workbook.close()


def test_no_valid_records_still_generates_audit_report(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _write_source(
        source,
        valid_rows=(
            _default_valid(NIK="#N/A"),
            _default_valid(No=2, Workflow="BAD", NIK="000009999"),
        ),
    )
    engine = AttDataRepairEngine(
        txt_writer=AttDataRepairTxtWriter(CodeProvider()),
        clock=FixedClock(),
    )

    result = engine.run_job(AttDataRepairJobRequest(_request(source), tmp_path))

    assert result.status == JobStatus.NO_VALID_RECORDS
    assert result.report_artifact is not None
    assert not list(result.paths.txt_folder.glob("*.txt"))
    workbook = load_workbook(result.report_artifact.file_path, data_only=True)
    try:
        assert workbook["Final_Records"].max_row == 1
        assert workbook["Changed_Records"].max_row == 1
        assert workbook["Anomaly"].max_row == 3
        assert workbook["Guide_Status"].max_row > 1
    finally:
        workbook.close()
    payload = json.loads(result.paths.summary_json.read_text(encoding="utf-8"))
    assert payload["status"] == "NO_VALID_RECORDS"
    assert payload["report"]["generated"] is True


def test_report_existing_file_is_not_overwritten(tmp_path: Path) -> None:
    source, result = _run_fixture(tmp_path)
    existing = result.report_artifact.file_path
    before = existing.read_bytes()
    writer = AttDataRepairReportWriter()

    with pytest.raises(ReportAlreadyExistsError):
        writer.write(
            report_folder=result.paths.report_folder,
            job_id=result.job_id,
            status=result.status,
            request=AttDataRepairJobRequest(_request(source), tmp_path),
            paths=result.paths,
            analysis=result.analysis_result,
            txt_artifacts=result.txt_artifacts,
            record_txt_assignment=result.record_txt_assignment,
            generated_at=FixedClock()(),
        )

    assert existing.read_bytes() == before


def test_report_save_failure_cleans_temporary_file(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "source.xlsx"
    _write_source(source, valid_rows=(_default_valid(),))
    analysis = AttDataRepairEngine(clock=FixedClock()).analyze(_request(source))
    paths = _paths(tmp_path)
    original_save = Workbook.save

    def fail_save(self, filename):
        original_save(self, filename)
        raise OSError("simulated save failure")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(Exception):
        AttDataRepairReportWriter().write(
            report_folder=paths.report_folder,
            job_id=paths.job_id,
            status="SUCCESS",
            request=AttDataRepairJobRequest(_request(source), tmp_path),
            paths=paths,
            analysis=analysis,
            txt_artifacts=(),
            record_txt_assignment={},
            generated_at=FixedClock()(),
        )
    monkeypatch.setattr(Workbook, "save", original_save)

    assert not list(paths.report_folder.glob("*.tmp.xlsx"))
    assert not list(paths.report_folder.glob("Att_Data_Repair_Report_*.xlsx"))


def test_failed_report_marks_job_failed_but_retains_txt(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _write_source(source, valid_rows=(_default_valid(),))

    class FailingReportWriter:
        def write(self, **_kwargs):
            raise RuntimeError("report boom")

    engine = AttDataRepairEngine(
        txt_writer=AttDataRepairTxtWriter(CodeProvider()),
        report_writer=FailingReportWriter(),
        clock=FixedClock(),
    )

    result = engine.run_job(AttDataRepairJobRequest(_request(source), tmp_path))

    assert result.status == JobStatus.FAILED
    assert result.txt_artifacts
    assert len(list(result.paths.txt_folder.glob("*.txt"))) == 1
    payload = json.loads(result.paths.summary_json.read_text(encoding="utf-8"))
    assert payload["txt"]["generated"] is True
    assert payload["report"]["generated"] is False
    assert "report boom" in payload["report"]["reason"]
    assert "report boom" in result.paths.process_log.read_text(encoding="utf-8")


def test_moderate_large_report_row_counts(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _write_source(source)
    records = tuple(
        FinalRecord(
            record_id=f"ADR-{index:06d}",
            workflow="HO",
            nik=f"{index:09d}",
            date_in=date(2026, 8, 3),
            time_in=time(9, 30),
            date_out=date(2026, 8, 3),
            time_out=time(17, 0),
            duration_minutes=450,
            source_sheet="Valid_Records",
            source_file="source.xlsx",
            source_row=index + 1,
            final_status=str(FinalStatus.VALID_UNCHANGED),
        )
        for index in range(1, 2501)
    )
    sources = tuple(
        SourceRecord(
            record_id=record.record_id,
            source_workbook=source,
            source_sheet="Valid_Records",
            source_record_no=index,
            source_file="source.xlsx",
            relative_path="source.xlsx",
            source_row=index + 1,
            workflow_raw="HO",
            nik_raw=record.nik,
            date_in_raw="08/03/2026",
            time_in_raw="09:30",
            date_out_raw="08/03/2026",
            time_out_raw="17:00",
            source_status="VALID",
            source_reason="",
            raw_value="",
        )
        for index, record in enumerate(records, 1)
    )
    changes = (
        RepairChange(
            record_id="ADR-000001",
            field_name="Time_Out",
            original_value="09:30",
            final_value="10:01",
            change_code=str(ChangeCode.MINIMUM_DURATION_APPLIED),
            reason="Durasi kurang dari minimum; Time_Out disesuaikan.",
        ),
    )
    analysis = AttDataRepairAnalysisResult(
        final_records=records,
        changed_records=(records[0],),
        change_log=changes,
        source_counts={"Valid_Records": len(records), "Invalid_Records": 0, "total": len(records)},
        source_records=sources,
    )
    paths = _paths(tmp_path)
    artifact = AttDataRepairReportWriter().write(
        report_folder=paths.report_folder,
        job_id=paths.job_id,
        status="SUCCESS",
        request=AttDataRepairJobRequest(_request(source), tmp_path),
        paths=paths,
        analysis=analysis,
        txt_artifacts=(
            TxtArtifact(
                workflow="HO",
                sequence=1,
                unique_code=4837,
                file_name="Att_Data_Repair_HO-001_4837.txt",
                file_path=paths.txt_folder / "Att_Data_Repair_HO-001_4837.txt",
                row_count=len(records),
                record_ids=tuple(record.record_id for record in records),
            ),
        ),
        record_txt_assignment={
            record.record_id: "Att_Data_Repair_HO-001_4837.txt"
            for record in records
        },
        generated_at=FixedClock()(),
    )

    workbook = load_workbook(artifact.file_path, read_only=True, data_only=True)
    try:
        assert sum(1 for _row in workbook["Final_Records"].iter_rows()) == 2501
        assert artifact.sheet_row_counts["Final_Records"] == 2501
    finally:
        workbook.close()


def _paths(root: Path) -> JobPaths:
    job_folder = (
        root
        / "Utilities"
        / "Att_Data_Repair"
        / "2026-08"
        / "2026-08-03_01"
    )
    txt_folder = job_folder / "TXT"
    report_folder = job_folder / "Report"
    txt_folder.mkdir(parents=True, exist_ok=True)
    report_folder.mkdir(parents=True, exist_ok=True)
    return JobPaths(
        output_root=root,
        utilities_root=root / "Utilities",
        feature_root=root / "Utilities" / "Att_Data_Repair",
        job_folder=job_folder,
        txt_folder=txt_folder,
        report_folder=report_folder,
        process_log=job_folder / "Process.log",
        summary_json=job_folder / "summary.json",
        job_id="2026-08-03_01",
    )
