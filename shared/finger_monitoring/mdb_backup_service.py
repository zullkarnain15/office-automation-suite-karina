"""Manual backup-MDB monitoring, export, and Outlook reporting."""

from __future__ import annotations

import fnmatch
import html
import os
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from shared.database.connection_factory import SQLiteConnectionFactory
from shared.finger_monitoring.device_master_service import DeviceMasterService
from shared.finger_monitoring.email_service import OutlookClassicSender, _com_initialized


STATUS_CURRENT = "Memenuhi"
STATUS_STALE = "Kedaluwarsa"
STATUS_MISSING = "File wajib hilang"
STATUS_ACCESS_ERROR = "Gagal diakses"
STATUS_PARTIAL = "Sebagian diperbarui"
STATUS_PROBLEM = "Bermasalah"


@dataclass(frozen=True, slots=True)
class MDBRequirement:
    unit: str
    branch: str
    file_name: str
    source: str = "Manual"

    @property
    def key(self) -> tuple[str, str, str]:
        return tuple(
            value.strip().casefold()
            for value in (self.unit, self.branch, self.file_name)
        )


@dataclass(frozen=True, slots=True)
class MDBBackupFileResult:
    unit: str
    branch: str
    file_name: str
    path: Path
    modified_at: datetime | None
    size_bytes: int | None
    status: str
    detail: str

    @property
    def is_problem(self) -> bool:
        return self.status != STATUS_CURRENT


@dataclass(frozen=True, slots=True)
class MDBBackupBranchSummary:
    unit: str
    branch: str
    total_files: int
    current_files: int
    stale_files: int
    missing_files: int
    access_errors: int
    status: str


@dataclass(frozen=True, slots=True)
class MDBBackupSnapshot:
    root_folder: Path
    checked_at: datetime
    reference_date: date
    cutoff_date: date
    files: tuple[MDBBackupFileResult, ...]
    branches: tuple[MDBBackupBranchSummary, ...]

    @property
    def problem_files(self) -> tuple[MDBBackupFileResult, ...]:
        return tuple(row for row in self.files if row.is_problem)

    @property
    def all_current(self) -> bool:
        return bool(self.files) and not self.problem_files


ProgressCallback = Callable[[int, int, str], None]


class MDBBackupScanner:
    """Recursively inspect each MDB using file metadata, never folder dates."""

    def __init__(self, *, stat_provider=None, scandir_provider=None, now=None) -> None:
        self._stat = stat_provider or (lambda path: path.stat())
        self._scandir = scandir_provider or os.scandir
        self._now = now or datetime.now

    def scan(
        self,
        root_folder: str | Path,
        reference_date: date,
        *,
        required: Iterable[MDBRequirement] = (),
        archive_exclusions: Iterable[str] = (),
        progress_callback: ProgressCallback | None = None,
    ) -> MDBBackupSnapshot:
        root = Path(root_folder)
        if not root.is_dir():
            raise ValueError("Folder sharing utama tidak tersedia atau bukan folder.")
        if not isinstance(reference_date, date):
            raise ValueError("Tanggal acuan tidak valid.")
        exclusions = tuple(
            value.strip().casefold()
            for value in archive_exclusions
            if value and value.strip()
        )
        cutoff = reference_date - timedelta(days=3)
        candidates: list[Path] = []
        access_rows: list[MDBBackupFileResult] = []

        if progress_callback:
            progress_callback(0, 0, "Memindai struktur folder...")
        self._walk(
            root, root, exclusions, candidates, access_rows, progress_callback
        )

        requirements = self._deduplicate_requirements(required)
        present_keys = set()
        for path in candidates:
            unit, branch = self._location(path, root)
            present_keys.add((unit.casefold(), branch.casefold(), path.name.casefold()))
        present_keys.update(
            (row.unit.casefold(), row.branch.casefold(), row.file_name.casefold())
            for row in access_rows
            if row.file_name and row.file_name != "(folder)"
        )
        missing_rows = []
        for item in requirements:
            expected = self._expected_path(root, item)
            if self._is_excluded(expected, root, exclusions) or item.key in present_keys:
                continue
            missing_rows.append(MDBBackupFileResult(
                item.unit,
                item.branch,
                item.file_name,
                expected,
                None,
                None,
                STATUS_MISSING,
                f"MDB wajib dari {item.source} tidak ditemukan.",
            ))

        total = len(candidates) + len(access_rows) + len(missing_rows)
        rows: list[MDBBackupFileResult] = []
        current = 0
        for path in candidates:
            unit, branch = self._location(path, root)
            try:
                metadata = self._stat(path)
                modified = datetime.fromtimestamp(metadata.st_mtime)
                status = STATUS_CURRENT if modified.date() >= cutoff else STATUS_STALE
                detail = (
                    f"Modified {modified:%Y-%m-%d %H:%M:%S} memenuhi batas H-3."
                    if status == STATUS_CURRENT
                    else f"Modified lebih lama dari batas {cutoff.isoformat()}."
                )
                row = MDBBackupFileResult(
                    unit, branch, path.name, path, modified, int(metadata.st_size),
                    status, detail,
                )
            except (OSError, PermissionError) as exc:
                row = MDBBackupFileResult(
                    unit, branch, path.name, path, None, None,
                    STATUS_ACCESS_ERROR, str(exc),
                )
            rows.append(row)
            current += 1
            if progress_callback:
                progress_callback(current, total, f"Memeriksa {path.name}")

        for row in (*access_rows, *missing_rows):
            rows.append(row)
            current += 1
            if progress_callback:
                progress_callback(current, total, row.detail)

        rows.sort(key=lambda row: (
            row.unit.casefold(), row.branch.casefold(), row.file_name.casefold(),
            str(row.path).casefold(),
        ))
        branches = self._summarize(rows)
        if progress_callback:
            progress_callback(total, total, f"Selesai: {len(rows)} hasil MDB")
        return MDBBackupSnapshot(
            root.resolve(), self._now(), reference_date, cutoff,
            tuple(rows), branches,
        )

    def _walk(
        self,
        directory: Path,
        root: Path,
        exclusions: tuple[str, ...],
        candidates: list[Path],
        access_rows: list[MDBBackupFileResult],
        progress_callback: ProgressCallback | None,
    ) -> None:
        if directory != root and self._is_excluded(directory, root, exclusions):
            return
        if progress_callback:
            try:
                label = str(directory.relative_to(root)) or "."
            except ValueError:
                label = str(directory)
            progress_callback(0, 0, f"Memindai folder {label}")
        try:
            with self._scandir(directory) as entries:
                listed = list(entries)
        except (OSError, PermissionError) as exc:
            unit, branch = self._location(directory / "(folder)", root)
            access_rows.append(MDBBackupFileResult(
                unit, branch, "(folder)", directory, None, None,
                STATUS_ACCESS_ERROR, f"Folder gagal diakses: {exc}",
            ))
            return
        for entry in listed:
            path = Path(entry.path)
            if self._is_excluded(path, root, exclusions):
                continue
            try:
                if entry.is_dir(follow_symlinks=False):
                    self._walk(
                        path, root, exclusions, candidates, access_rows,
                        progress_callback,
                    )
                elif entry.is_file(follow_symlinks=False) and path.suffix.casefold() == ".mdb":
                    candidates.append(path)
            except (OSError, PermissionError) as exc:
                unit, branch = self._location(path, root)
                access_rows.append(MDBBackupFileResult(
                    unit, branch, path.name, path, None, None,
                    STATUS_ACCESS_ERROR, str(exc),
                ))

    @staticmethod
    def _location(path: Path, root: Path) -> tuple[str, str]:
        try:
            parts = list(path.relative_to(root).parts[:-1])
        except ValueError:
            parts = list(path.parts[:-1])
        lowered = [part.casefold() for part in parts]
        if "mdb" in lowered:
            parts = parts[len(parts) - 1 - lowered[::-1].index("mdb") + 1:]
        if len(parts) >= 2:
            return parts[-2], parts[-1]
        if len(parts) == 1:
            return "", parts[0]
        return "", ""

    @staticmethod
    def _expected_path(root: Path, requirement: MDBRequirement) -> Path:
        base = root if root.name.casefold() == "mdb" else root / "MDB"
        return base / requirement.unit / requirement.branch / requirement.file_name

    @staticmethod
    def _is_excluded(path: Path, root: Path, exclusions: tuple[str, ...]) -> bool:
        if not exclusions:
            return False
        try:
            relative = path.relative_to(root)
        except ValueError:
            relative = path
        parts = tuple(part.casefold() for part in relative.parts)
        joined = "/".join(parts)
        return any(
            fnmatch.fnmatchcase(joined, pattern)
            or any(fnmatch.fnmatchcase(part, pattern) for part in parts)
            for pattern in exclusions
        )

    @staticmethod
    def _deduplicate_requirements(
        required: Iterable[MDBRequirement],
    ) -> tuple[MDBRequirement, ...]:
        result = {}
        for item in required:
            result.setdefault(item.key, item)
        return tuple(result.values())

    @staticmethod
    def _summarize(
        rows: Iterable[MDBBackupFileResult],
    ) -> tuple[MDBBackupBranchSummary, ...]:
        grouped: dict[tuple[str, str], list[MDBBackupFileResult]] = defaultdict(list)
        for row in rows:
            grouped[(row.unit, row.branch)].append(row)
        result = []
        for (unit, branch), items in sorted(
            grouped.items(), key=lambda item: (item[0][0].casefold(), item[0][1].casefold())
        ):
            current = sum(item.status == STATUS_CURRENT for item in items)
            stale = sum(item.status == STATUS_STALE for item in items)
            missing = sum(item.status == STATUS_MISSING for item in items)
            inaccessible = sum(item.status == STATUS_ACCESS_ERROR for item in items)
            if current == len(items):
                status = STATUS_CURRENT
            elif current:
                status = STATUS_PARTIAL
            elif inaccessible and not stale and not missing:
                status = STATUS_ACCESS_ERROR
            elif stale and not inaccessible and not missing:
                status = STATUS_STALE
            elif missing and not inaccessible and not stale:
                status = STATUS_MISSING
            else:
                status = STATUS_PROBLEM
            result.append(MDBBackupBranchSummary(
                unit, branch, len(items), current, stale, missing, inaccessible, status,
            ))
        return tuple(result)


class MDBBackupMonitoringService:
    """One snapshot powers the table, Excel export, and Outlook email."""

    def __init__(
        self,
        core_storage_service=None,
        *,
        sender=None,
        scanner: MDBBackupScanner | None = None,
        now=None,
        factory: SQLiteConnectionFactory | None = None,
    ) -> None:
        self.core_storage_service = core_storage_service
        self.sender = sender or OutlookClassicSender()
        self.now = now or datetime.now
        self.scanner = scanner or MDBBackupScanner(now=self.now)
        self.factory = factory or SQLiteConnectionFactory()

    def scan(
        self,
        root_folder: str | Path,
        reference_date: date,
        *,
        required: Iterable[MDBRequirement] = (),
        archive_exclusions: Iterable[str] = (),
        progress_callback: ProgressCallback | None = None,
        include_attendance_required: bool = True,
    ) -> MDBBackupSnapshot:
        combined = list(required)
        if include_attendance_required:
            combined.extend(self.load_attendance_requirements())
        return self.scanner.scan(
            root_folder,
            reference_date,
            required=combined,
            archive_exclusions=archive_exclusions,
            progress_callback=progress_callback,
        )

    def load_attendance_requirements(self) -> tuple[MDBRequirement, ...]:
        if self.core_storage_service is None:
            return ()
        try:
            status = self.core_storage_service.resolve_status()
        except (OSError, RuntimeError, ValueError):
            return ()
        database = getattr(status, "database_path", None)
        if not getattr(status, "database_valid", False) or database is None:
            return ()
        try:
            with self.factory.connect(database, read_only=True) as connection:
                paths = tuple(
                    Path(row[0])
                    for row in connection.execute(
                        """SELECT mdb_path FROM attendance_sources
                           WHERE is_active=1 AND mdb_path IS NOT NULL
                           AND trim(mdb_path)<>'' ORDER BY workflow,sort_order,source_code"""
                    )
                )
        except (OSError, ValueError, sqlite3.Error):
            return ()
        return self.requirements_from_attendance_paths(paths)

    @staticmethod
    def requirements_from_attendance_paths(
        paths: Iterable[Path],
    ) -> tuple[MDBRequirement, ...]:
        result = []
        for path in paths:
            parts = list(Path(path).parts)
            lowered = [part.casefold() for part in parts]
            if "mdb" in lowered:
                index = len(parts) - 1 - lowered[::-1].index("mdb")
                tail = parts[index + 1:]
            else:
                tail = parts[-3:]
            if len(tail) < 3 or Path(tail[-1]).suffix.casefold() != ".mdb":
                continue
            result.append(MDBRequirement(tail[-3], tail[-2], tail[-1], "Attendance"))
        return MDBBackupScanner._deduplicate_requirements(result)

    @staticmethod
    def parse_required_entries(value: str) -> tuple[MDBRequirement, ...]:
        result = []
        for raw in re.split(r"[;\n,]+", value or ""):
            text = raw.strip().replace("\\", "/")
            if not text:
                continue
            parts = [part.strip() for part in text.split("/") if part.strip()]
            if len(parts) < 3 or Path(parts[-1]).suffix.casefold() != ".mdb":
                raise ValueError(
                    f"MDB wajib '{raw.strip()}' harus berformat Unit/Cabang/file.mdb."
                )
            result.append(MDBRequirement(parts[-3], parts[-2], parts[-1], "Manual"))
        return MDBBackupScanner._deduplicate_requirements(result)

    @staticmethod
    def parse_archive_exclusions(value: str) -> tuple[str, ...]:
        return tuple(
            item.strip() for item in re.split(r"[;\n,]+", value or "") if item.strip()
        )

    def export(self, snapshot: MDBBackupSnapshot, path: str | Path) -> Path:
        target = Path(path)
        if target.suffix.casefold() != ".xlsx":
            raise ValueError("Export Monitoring Backup MDB harus berformat .xlsx.")
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Ringkasan Cabang"
        detail = workbook.create_sheet("Detail File")
        for sheet in (summary, detail):
            sheet.append(("Waktu Pemeriksaan", snapshot.checked_at.isoformat(sep=" ")))
            sheet.append(("Tanggal Acuan", snapshot.reference_date.isoformat()))
            sheet.append(("Batas H-3", snapshot.cutoff_date.isoformat()))
            sheet.append(("Folder Sharing", str(snapshot.root_folder)))
            sheet.append(())

        summary_headers = (
            "Unit", "Cabang", "Total MDB", "Memenuhi", "Kedaluwarsa",
            "Wajib Hilang", "Gagal Diakses", "Status",
        )
        detail_headers = (
            "Unit", "Cabang", "File", "Modified", "Ukuran (bytes)",
            "Status", "Keterangan", "Path",
        )
        summary.append(summary_headers)
        detail.append(detail_headers)
        for row in snapshot.branches:
            summary.append((
                row.unit, row.branch, row.total_files, row.current_files,
                row.stale_files, row.missing_files, row.access_errors, row.status,
            ))
        for row in snapshot.files:
            detail.append((
                row.unit, row.branch, row.file_name,
                row.modified_at.isoformat(sep=" ") if row.modified_at else "",
                row.size_bytes, row.status, row.detail, str(row.path),
            ))
        self._style_sheet(summary, (14, 18, 13, 12, 15, 14, 15, 24))
        self._style_sheet(detail, (14, 18, 28, 21, 16, 22, 48, 65))
        DeviceMasterService._atomic_save(workbook, target)
        return target

    def send_email(self, snapshot: MDBBackupSnapshot, recipient: str) -> bool:
        problems = snapshot.problem_files
        if not problems:
            return False
        address = recipient.strip()
        if not address:
            raise ValueError("Recipient Email Finger Monitoring belum diisi.")
        body_rows = "".join(
            "<tr>"
            f"<td>{html.escape(row.unit)}</td>"
            f"<td>{html.escape(row.branch)}</td>"
            f"<td>{html.escape(row.file_name)}</td>"
            f"<td>{html.escape(row.status)}</td>"
            f"<td>{html.escape(row.detail)}</td>"
            "</tr>"
            for row in problems
        )
        body = (
            "<html><body><h3>Monitoring Backup MDB</h3>"
            f"<p>Waktu pemeriksaan: {html.escape(snapshot.checked_at.isoformat(sep=' '))}<br>"
            f"Tanggal acuan: {snapshot.reference_date.isoformat()}<br>"
            f"Batas H-3: {snapshot.cutoff_date.isoformat()}</p>"
            "<table border='1' cellpadding='4'><tr><th>Unit</th><th>Cabang</th>"
            "<th>File</th><th>Status</th><th>Keterangan</th></tr>"
            + body_rows
            + "</table></body></html>"
        )
        with _com_initialized():
            self.sender.send_html(
                address,
                f"[OAS-K] Monitoring Backup MDB - {len(problems)} file bermasalah",
                body,
            )
        return True

    @staticmethod
    def _style_sheet(sheet, widths: tuple[int, ...]) -> None:
        for cell in sheet[6]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A7"
        sheet.auto_filter.ref = f"A6:{get_column_letter(len(widths))}{sheet.max_row}"
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
