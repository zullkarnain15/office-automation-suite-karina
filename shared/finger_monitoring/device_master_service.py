"""Device Master and standalone Excel operations; deliberately GUI-free."""

from __future__ import annotations

import ipaddress
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from shared.database.time_utils import current_timestamp

TEMPLATE_FILENAME = "OASK_Finger_Device_Master_Template.xlsx"
SHEET_NAME = "Device_Master"
HEADERS = (
    "Unit",
    "3LC",
    "Branch_Name",
    "IP_Address",
    "Device_Model",
    "Active",
    "Note",
)
LEGACY_HEADERS = ("Branch_Name", "IP_Address", "Device_Model", "Active")
ACTIVE_VALUES = {"YES": True, "Y": True, "TRUE": True, "1": True, "NO": False, "N": False, "FALSE": False, "0": False}


@dataclass(frozen=True, slots=True)
class DeviceImportItem:
    row_number: int
    action: str
    branch_name: str
    ip_address: str
    device_model: str
    active: bool
    unit: str = ""
    three_lc: str = ""
    note: str = ""
    device_id: int | None = None
    reason: str = ""


@dataclass(frozen=True, slots=True)
class DeviceImportPreview:
    source_path: Path
    revision: tuple[int, int, str | None]
    items: tuple[DeviceImportItem, ...]

    @property
    def insert_count(self) -> int:
        return sum(item.action == "INSERT" for item in self.items)

    @property
    def update_count(self) -> int:
        return sum(item.action == "UPDATE" for item in self.items)

    @property
    def rejected_count(self) -> int:
        return sum(item.action == "REJECTED" for item in self.items)

    @property
    def can_commit(self) -> bool:
        return bool(self.items) and self.rejected_count == 0


class DeviceMasterService:
    """Service layer for master maintenance, Excel, and future GUI copy actions."""

    def __init__(self, repository, *, activity_gate=None) -> None:
        self.repository = repository
        self.activity_gate = activity_gate

    def add_device(
        self,
        branch_name: str,
        ip_address: str,
        *,
        device_model: str = "X1500",
        unit: str | None = None,
        three_lc: str | None = None,
        note: str | None = None,
    ) -> int:
        return self.repository.create_device(
            branch_name,
            ip_address,
            device_model=device_model,
            unit=unit,
            three_lc=three_lc,
            note=note,
        )

    def edit_device(
        self,
        device_id: int,
        branch_name: str,
        ip_address: str,
        *,
        device_model: str = "X1500",
        unit: str | None = None,
        three_lc: str | None = None,
        note: str | None = None,
    ) -> None:
        self.repository.update_device(
            device_id,
            branch_name,
            ip_address,
            device_model=device_model,
            unit=unit,
            three_lc=three_lc,
            note=note,
        )

    def soft_delete(self, device_id: int) -> None:
        self.repository.soft_delete_device(device_id)

    def reactivate(self, device_id: int) -> None:
        self.repository.reactivate_device(device_id)

    def search(self, *, search: str = "", device_model: str | None = None, active: bool | None = True):
        return self.repository.list_devices(
            include_deleted=active is None,
            search=search,
            device_model=device_model,
            active=active,
        )

    @staticmethod
    def generate_template(path: str | Path) -> Path:
        path = Path(path)
        if path.name != TEMPLATE_FILENAME:
            path = path / TEMPLATE_FILENAME if path.suffix.lower() != ".xlsx" else path
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        DeviceMasterService._style_header(sheet, HEADERS)
        sheet.append(("CF", "MDU", "Example Branch", "192.168.1.10", "X1500", "YES", "OTO AMU + CEMAS + ARD"))
        sheet.append(("MF", "BAN", "Example Branch", "192.168.1.11", "X1500", "YES", ""))
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for index, width in enumerate((14, 14, 30, 20, 18, 12, 38), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        DeviceMasterService._atomic_save(workbook, path)
        return path

    def preview_import(self, path: str | Path) -> DeviceImportPreview:
        if self.activity_gate is not None:
            with self.activity_gate.activity("IMPORT"):
                return self._preview_import(path)
        return self._preview_import(path)

    def _preview_import(self, path: str | Path) -> DeviceImportPreview:
        source = Path(path)
        if source.suffix.lower() != ".xlsx":
            raise ValueError("Import Device Master harus file .xlsx.")
        workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise ValueError("Sheet Device_Master tidak ditemukan.")
            sheet = workbook[SHEET_NAME]
            headers = tuple(self._text(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            if headers not in {HEADERS, LEGACY_HEADERS}:
                raise ValueError("Header harus tepat: " + ", ".join(HEADERS))
            raw_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        finally:
            workbook.close()

        prepared: list[DeviceImportItem] = []
        ip_rows: dict[str, list[int]] = {}
        for index, row in enumerate(raw_rows, 2):
            values = tuple(self._text(value) for value in row[: len(headers)])
            if not any(values):
                continue
            if headers == LEGACY_HEADERS:
                values = ("", "", *values, "")
            item = self._parse_import_row(index, values)
            prepared.append(item)
            if item.action != "REJECTED":
                ip_rows.setdefault(item.ip_address, []).append(len(prepared) - 1)
        for indices in ip_rows.values():
            if len(indices) > 1:
                for item_index in indices:
                    item = prepared[item_index]
                    prepared[item_index] = DeviceImportItem(
                        item.row_number, "REJECTED", item.branch_name, item.ip_address,
                    item.device_model, item.active, item.unit, item.three_lc, item.note,
                    reason="Duplicate IP_Address di file import.",
                    )

        existing = self.repository.list_devices(include_deleted=True, active=None)
        by_ip: dict[str, dict] = {}
        for row in sorted(existing, key=lambda value: (value["is_active"], value["device_id"]), reverse=True):
            by_ip.setdefault(row["ip_address"], row)
        planned = []
        for item in prepared:
            if item.action == "REJECTED":
                planned.append(item)
                continue
            previous = by_ip.get(item.ip_address)
            action = "UPDATE" if previous is not None else "INSERT"
            planned.append(DeviceImportItem(
                item.row_number, action, item.branch_name, item.ip_address,
                item.device_model, item.active, item.unit, item.three_lc, item.note,
                int(previous["device_id"]) if previous else None,
            ))
        return DeviceImportPreview(source, self.repository.device_revision(), tuple(planned))

    def commit_import(self, preview: DeviceImportPreview) -> dict[str, int]:
        if not preview.can_commit:
            raise ValueError("Preview memiliki rejection atau tidak ada data untuk diimport.")
        if self.activity_gate is not None:
            with self.activity_gate.activity("IMPORT"):
                return self._commit_import(preview)
        return self._commit_import(preview)

    def _commit_import(self, preview: DeviceImportPreview) -> dict[str, int]:
        factory = self.repository.factory
        with factory.connect(self.repository.database_path) as connection:
            connection.execute("BEGIN IMMEDIATE")
            if self.repository._device_revision(connection) != preview.revision:
                raise RuntimeError("Device Master berubah setelah preview. Buat preview baru.")
            for item in preview.items:
                self._apply_item(connection, item)
        return {"inserted": preview.insert_count, "updated": preview.update_count, "rejected": 0}

    @staticmethod
    def _apply_item(connection, item: DeviceImportItem) -> None:
        now = current_timestamp()
        if item.action == "INSERT":
            connection.execute(
                "INSERT INTO device_master(unit,three_lc,branch_name,ip_address,device_model,is_active,note,created_at,updated_at,deleted_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (item.unit, item.three_lc, item.branch_name, item.ip_address, item.device_model, int(item.active), item.note, now, now, None if item.active else now),
            )
        elif item.action == "UPDATE" and item.device_id is not None:
            changed = connection.execute(
                "UPDATE device_master SET unit=?,three_lc=?,branch_name=?,ip_address=?,device_model=?,is_active=?,note=?,updated_at=?,deleted_at=? WHERE device_id=?",
                (item.unit, item.three_lc, item.branch_name, item.ip_address, item.device_model, int(item.active), item.note, now, None if item.active else now, item.device_id),
            ).rowcount
            if changed != 1:
                raise RuntimeError("Perangkat import tidak lagi tersedia.")
        else:
            raise ValueError("Item import tidak valid.")

    def export(self, path: str | Path, *, search: str = "", device_model: str | None = None, active: bool | None = True) -> Path:
        target = Path(path)
        if target.suffix.lower() != ".xlsx":
            raise ValueError("Export harus berformat .xlsx.")
        rows = self.search(search=search, device_model=device_model, active=active)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        self._style_header(sheet, HEADERS)
        for row in rows:
            sheet.append((row["unit"], row["three_lc"], row["branch_name"], row["ip_address"], row["device_model"], "YES" if row["is_active"] else "NO", row["note"]))
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for index, width in enumerate((14, 14, 30, 20, 18, 12, 38), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        self._atomic_save(workbook, target)
        return target

    @staticmethod
    def copy_rows(rows: Iterable[dict], *, fields: tuple[str, ...] = ("unit", "three_lc", "branch_name", "ip_address", "device_model", "is_active", "note"), include_header: bool = False) -> str:
        lines = []
        if include_header:
            lines.append("\t".join(fields))
        for row in rows:
            values = []
            for field in fields:
                value = row.get(field, "")
                if field == "is_active":
                    value = "YES" if value else "NO"
                values.append(str(value))
            lines.append("\t".join(values))
        return "\n".join(lines)

    @classmethod
    def copy_selected(cls, all_rows: Iterable[dict], selected_ids: Iterable[int]) -> str:
        selected = {int(value) for value in selected_ids}
        return cls.copy_rows(row for row in all_rows if int(row["device_id"]) in selected)

    @classmethod
    def copy_ip_addresses(cls, rows: Iterable[dict]) -> str:
        return cls.copy_rows(rows, fields=("ip_address",))

    @classmethod
    def copy_branch_names(cls, rows: Iterable[dict]) -> str:
        return cls.copy_rows(rows, fields=("branch_name",))

    @staticmethod
    def _parse_import_row(row_number: int, values: tuple[str, ...]) -> DeviceImportItem:
        unit, three_lc, branch, address, model, active, note = values
        if not branch or not address:
            return DeviceImportItem(row_number, "REJECTED", branch, address, model, True, unit, three_lc, note, reason="Branch_Name dan IP_Address wajib diisi.")
        try:
            parsed = ipaddress.ip_address(address)
            if parsed.version != 4:
                raise ValueError
        except ValueError:
            return DeviceImportItem(row_number, "REJECTED", branch, address, model, True, unit, three_lc, note, reason="IP_Address harus IPv4 valid.")
        model = (model or "X1500").strip().upper()
        if not model or len(model) > 100:
            return DeviceImportItem(row_number, "REJECTED", branch, str(parsed), model, True, unit, three_lc, note, reason="Device_Model wajib diisi dan maksimal 100 karakter.")
        active_value = ACTIVE_VALUES.get((active or "YES").upper())
        if active_value is None:
            return DeviceImportItem(row_number, "REJECTED", branch, str(parsed), model, True, unit, three_lc, note, reason="Active harus YES atau NO.")
        return DeviceImportItem(row_number, "PENDING", branch, str(parsed), model, active_value, unit, three_lc, note)

    @staticmethod
    def _text(value) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _style_header(sheet, headers) -> None:
        for index, value in enumerate(headers, 1):
            cell = sheet.cell(1, index, value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center")

    @staticmethod
    def _atomic_save(workbook, target: Path) -> None:
        target.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary = tempfile.mkstemp(prefix=".finger_device_", suffix=".xlsx", dir=target.parent)
        os.close(descriptor)
        temp_path = Path(temporary)
        try:
            workbook.save(temp_path)
            os.replace(temp_path, target)
        finally:
            workbook.close()
            if temp_path.exists():
                temp_path.unlink()
