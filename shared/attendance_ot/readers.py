"""Low-memory, anchor-based XLSX/XLS row readers."""

from __future__ import annotations

import re
from collections.abc import Iterator
from contextlib import suppress
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from shared.attendance_ot.models import ParsedSourceRow, ScheduleSubtype, SourceType

_WHITESPACE = re.compile(r"\s+")

ATTENDANCE_ANCHORS = frozenset(
    {
        "emplid",
        "date in",
        "time in",
        "date out",
        "time out",
        "overtime hour",
        "overtime minute",
        "ot amount",
        "meal ot amount",
        "location descr",
        "name",
        "paylink",
    }
)
EMPLOYEE_ANCHORS = frozenset(
    {
        "emplid",
        "name",
        "pay group",
        "status",
        "location descr",
        "regional",
        "dept desc",
        "business desc",
        "jobcd desc",
    }
)
SCHEDULE_ANCHORS = frozenset({"emplid", "schedule id", "schedule descr", "effdt"})
SCHEDULE_RULE_ANCHORS = frozenset(
    {
        "keterangan",
        "rule",
        "schedule rule",
        "working hours",
        "jam kerja",
        "scheduled in",
        "scheduled out",
        "time in",
        "time out",
        "weekday",
        "workday",
    }
)
HEADER_ALIASES = {
    "employee id": "emplid",
    "employee name": "name",
    "paygroup": "pay group",
    "department": "dept desc",
    "department desc": "dept desc",
    "bus. desc.": "business desc",
    "bus. desc": "business desc",
    "bus desc": "business desc",
    "business description": "business desc",
    "jobcd desc.": "jobcd desc",
    "job code desc": "jobcd desc",
    "job description": "jobcd desc",
    "effective date": "effdt",
    "end effective date": "end effdt",
    "snapshot date": "snapshot date",
    "as of date": "snapshot date",
    "as-of date": "snapshot date",
    "employee period": "snapshot date",
    "period": "snapshot date",
    "periode": "snapshot date",
    "data period": "snapshot date",
    "payroll period": "snapshot date",
}


class SourceFormatError(ValueError):
    """Raised when no compatible anchor-based table is found."""


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return _WHITESPACE.sub(" ", str(value).strip()).casefold()


def header_map(values: tuple[object, ...]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, value in enumerate(values):
        normalized = HEADER_ALIASES.get(
            normalize_header(value), normalize_header(value)
        )
        if normalized and normalized not in result:
            result[normalized] = index
    return result


def detect_source_type(headers: dict[str, int]) -> SourceType | None:
    detected = detect_source_contract(headers)
    return detected[0] if detected is not None else None


def detect_source_contract(
    headers: dict[str, int],
) -> tuple[SourceType, ScheduleSubtype | None] | None:
    names = set(headers)
    scores: list[tuple[int, SourceType]] = []
    attendance_matches = len(names & ATTENDANCE_ANCHORS)
    if {"emplid", "date in", "time in"} <= names and attendance_matches >= 6:
        scores.append((attendance_matches, SourceType.ATTENDANCE_OT, None))
    employee_matches = len(names & EMPLOYEE_ANCHORS)
    if {"emplid", "name"} <= names and employee_matches >= 4:
        scores.append((employee_matches, SourceType.EMPLOYEE, None))
    if SCHEDULE_ANCHORS <= names:
        scores.append(
            (len(SCHEDULE_ANCHORS) + 2, SourceType.SCHEDULE, ScheduleSubtype.ASSIGNMENT)
        )
    rule_matches = len(names & SCHEDULE_RULE_ANCHORS)
    if {"schedule id", "schedule descr"} <= names and rule_matches:
        scores.append((2 + rule_matches, SourceType.SCHEDULE, ScheduleSubtype.RULE))
    if not scores:
        return None
    scores.sort(key=lambda item: item[0], reverse=True)
    return scores[0][1], scores[0][2]


def iter_source_rows(
    path: Path,
    expected_type: SourceType,
) -> Iterator[ParsedSourceRow]:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        yield from _iter_xlsx(path, expected_type)
        return
    if suffix == ".xls":
        yield from _iter_xls(path, expected_type)
        return
    raise SourceFormatError(f"Format file tidak didukung: {path.suffix}")


def _iter_xlsx(path: Path, expected_type: SourceType) -> Iterator[ParsedSourceRow]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    matched = False
    mismatches: set[str] = set()
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            found = _find_header(rows)
            if found is None:
                continue
            header_row, headers, detected, subtype = found
            if detected is not expected_type:
                mismatches.add(detected.value)
                continue
            matched = True
            for row_number, values in enumerate(rows, start=header_row + 1):
                row = tuple(values)
                if _has_identifier(row, headers, subtype):
                    yield ParsedSourceRow(
                        expected_type,
                        worksheet.title,
                        row_number,
                        headers,
                        row,
                        subtype,
                    )
    finally:
        workbook.close()
    if not matched:
        raise _format_error(path, expected_type, mismatches)


def _iter_xls(path: Path, expected_type: SourceType) -> Iterator[ParsedSourceRow]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(
            "Reader XLS belum tersedia. Install dependency xlrd==2.0.2."
        ) from exc

    workbook = xlrd.open_workbook(path, on_demand=True)
    matched = False
    mismatches: set[str] = set()
    try:
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            rows = (
                _xls_row(sheet, index, workbook.datemode)
                for index in range(sheet.nrows)
            )
            found = _find_header(rows)
            if found is None:
                workbook.unload_sheet(sheet_name)
                continue
            header_row, headers, detected, subtype = found
            if detected is not expected_type:
                mismatches.add(detected.value)
                workbook.unload_sheet(sheet_name)
                continue
            matched = True
            for zero_based in range(header_row, sheet.nrows):
                row = _xls_row(sheet, zero_based, workbook.datemode)
                if _has_identifier(row, headers, subtype):
                    yield ParsedSourceRow(
                        expected_type, sheet_name, zero_based + 1, headers, row, subtype
                    )
            workbook.unload_sheet(sheet_name)
    finally:
        with suppress(Exception):
            workbook.release_resources()
    if not matched:
        raise _format_error(path, expected_type, mismatches)


def _find_header(rows: Iterator[tuple[object, ...]]):
    for row_number, values in enumerate(rows, start=1):
        row = tuple(values)
        headers = header_map(row)
        detected = detect_source_contract(headers)
        if detected is not None:
            return row_number, headers, detected[0], detected[1]
    return None


def _has_identifier(
    values: tuple[object, ...],
    headers: dict[str, int],
    subtype: ScheduleSubtype | None,
) -> bool:
    if subtype is ScheduleSubtype.RULE:
        return any(
            index < len(values) and values[index] not in (None, "")
            for name, index in headers.items()
            if name in ({"schedule id", "schedule descr"} | SCHEDULE_RULE_ANCHORS)
        )
    index = headers.get("emplid")
    return index is not None and index < len(values) and values[index] not in (None, "")


def _xls_row(sheet, row_index: int, datemode: int) -> tuple[object, ...]:
    import xlrd

    values: list[object] = []
    for cell in sheet.row(row_index):
        value = cell.value
        if cell.ctype == xlrd.XL_CELL_DATE:
            parts = xlrd.xldate_as_tuple(value, datemode)
            if parts[:3] == (0, 0, 0):
                value = time(parts[3], parts[4], parts[5])
            else:
                value = datetime(*parts)
        values.append(value)
    return tuple(values)


def _format_error(
    path: Path,
    expected_type: SourceType,
    mismatches: set[str],
) -> SourceFormatError:
    if mismatches:
        detected = ", ".join(sorted(mismatches))
        return SourceFormatError(
            f"Anchor file menunjukkan source type {detected}, bukan {expected_type.value}."
        )
    return SourceFormatError(
        f"Header {expected_type.value} tidak ditemukan pada workbook {path.name}."
    )


def identifier_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def text_value(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def date_text(value: object) -> str | None:
    parsed = parse_date_value(value)
    return parsed.isoformat() if parsed is not None else None


def parse_date_value(value: object) -> date | None:
    """Parse Excel dates, ISO dates, and production MM/DD/YYYY text."""

    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        raise ValueError(f"Tanggal tidak valid: {value!r}.")
    if isinstance(value, (int, float)):
        number = float(value)
        if number.is_integer() and 10_000_101 <= number <= 99_991_231:
            text = str(int(number))
            try:
                return date(int(text[:4]), int(text[4:6]), int(text[6:8]))
            except ValueError as exc:
                raise ValueError(f"Tanggal tidak valid: {value!r}.") from exc
        try:
            converted = from_excel(number)
        except (OverflowError, ValueError) as exc:
            raise ValueError(f"Serial tanggal Excel tidak valid: {value!r}.") from exc
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
        raise ValueError(f"Serial tanggal Excel tidak valid: {value!r}.")

    text = _WHITESPACE.sub(" ", str(value).replace("\u00a0", " ").strip())
    iso = re.fullmatch(
        r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:[ T].*)?", text
    )
    production = re.fullmatch(
        r"(\d{1,2})/(\d{1,2})/(\d{4})(?:[ T].*)?", text
    )
    match = iso or production
    if match is None:
        raise ValueError(
            f"Format tanggal tidak didukung: {value!r}. Gunakan YYYY-MM-DD atau MM/DD/YYYY."
        )
    if iso:
        year, month, day = (int(part) for part in match.groups())
    else:
        month, day, year = (int(part) for part in match.groups())
    try:
        return date(year, month, day)
    except ValueError as exc:
        raise ValueError(f"Tanggal tidak valid: {value!r}.") from exc


def period_date_text(value: object) -> str | None:
    raw = text_value(value)
    if raw is None:
        return None
    match = re.fullmatch(r"(\d{4})[-/](\d{1,2})", raw)
    if match:
        year, month = int(match.group(1)), int(match.group(2))
        try:
            return date(year, month, 1).isoformat()
        except ValueError as exc:
            raise ValueError(f"Periode tidak valid: {value!r}.") from exc
    match = re.fullmatch(r"(\d{1,2})[-/](\d{4})", raw)
    if match:
        month, year = int(match.group(1)), int(match.group(2))
        try:
            return date(year, month, 1).isoformat()
        except ValueError as exc:
            raise ValueError(f"Periode tidak valid: {value!r}.") from exc
    return date_text(value)


def time_text(value: object) -> str | None:
    parsed = parse_time_value(value)
    return parsed.replace(microsecond=0).isoformat() if parsed is not None else None


def parse_time_value(value: object) -> time | None:
    """Parse native Excel time, 24-hour text, and 12-hour AM/PM text."""

    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, bool):
        raise ValueError(f"Waktu tidak valid: {value!r}.")
    if isinstance(value, (int, float)):
        number = float(value)
        if not 0 <= number < 1:
            raise ValueError(f"Serial waktu Excel tidak valid: {value!r}.")
        seconds = round(number * 86400) % 86400
        return (datetime.min + timedelta(seconds=seconds)).time()

    text = _WHITESPACE.sub(" ", str(value).replace("\u00a0", " ").strip())
    text = re.sub(r"(?<=\d)[.](?=\d)", ":", text)
    text = re.sub(r"(?i)\b([AP])\s*\.?\s*M\.?\b", r"\1M", text).upper()
    if "T" in text or re.match(r"^\d{4}-\d{1,2}-\d{1,2} ", text):
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).time().replace(
                microsecond=0
            )
        except ValueError:
            pass
    for pattern in ("%I:%M:%S %p", "%I:%M %p", "%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, pattern).time()
        except ValueError:
            continue
    try:
        number = float(text)
    except ValueError:
        number = -1
    if 0 <= number < 1:
        seconds = round(number * 86400) % 86400
        return (datetime.min + timedelta(seconds=seconds)).time()
    raise ValueError(
        f"Format waktu tidak didukung: {value!r}. Gunakan HH:MM:SS atau h:mm:ss AM/PM."
    )


def numeric_value(value: object) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    if "," in text and "." in text:
        text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None
