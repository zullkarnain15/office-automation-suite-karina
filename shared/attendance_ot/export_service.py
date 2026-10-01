"""Streaming Excel reporting using a consistent read snapshot."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from shared.attendance_ot.reporting import (
    AnalyticsFilters,
    DETAIL_LABELS,
    MONTHS,
    SOURCE_COLUMNS,
)


class AttendanceOTExportService:
    def __init__(self, queries, *, sheet_row_limit=1_048_576):
        if not 2 <= sheet_row_limit <= 1_048_576:
            raise ValueError("Excel sheet row limit tidak valid.")
        self.queries = queries
        self.sheet_row_limit = sheet_row_limit

    def export(self, path, filters=AnalyticsFilters(), report=None):
        path = Path(path).resolve()
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Output harus berformat .xlsx.")
        report = report or (lambda _text: None)
        workbook = Workbook(write_only=True)
        temp_path = None
        try:
            with self.queries.snapshot() as session:
                for row in session.connection.execute(
                    "SELECT source_path FROM source_registry"
                ):
                    if Path(row[0]).resolve() == path:
                        raise ValueError("Output tidak boleh menimpa raw source file.")
                report("Export Overview")
                dashboard = session.dashboard(filters)
                overview = self._sheet(workbook, "Overview", ("Metric", "Value"))
                self._append(overview, ("Active filters", filters.describe()))
                for key, value in dashboard["kpi"].items():
                    self._append(overview, (key, value))
                self._append(
                    overview,
                    (
                        "Classification",
                        "HEAD OFFICE Location Descr = HO; otherwise BRANCH (Sprint 3)",
                    ),
                )
                self._append(
                    overview,
                    ("Amounts", "OT Amount and Meal Amount from imported source"),
                )
                relation, args = session.relation(filters)
                months = ",".join(
                    f"SUM(CASE WHEN substr(year_month,6,2)='{i:02d}' THEN late_count ELSE 0 END) {name}"
                    for i, name in enumerate(MONTHS, 1)
                )
                late_sql = f"""WITH filtered AS ({relation}) SELECT substr(year_month,1,4) year,
                    classification, CASE WHEN classification='HO' THEN department ELSE branch END group_name,
                    emplid, MAX(employee_name) name, {months}, SUM(late_count) Total
                    FROM filtered GROUP BY year, classification, group_name, emplid
                    HAVING SUM(late_count)>0 ORDER BY year, classification, group_name, emplid"""
                self._write_stream(
                    workbook,
                    "Late_Summary",
                    (
                        "Year",
                        "Classification",
                        "Department / Branch",
                        "EMPLID",
                        "Name",
                        *MONTHS,
                        "Total",
                    ),
                    session.stream(late_sql, args),
                    report,
                )
                detail_sql, detail_args = session.detail_sql(filters)
                where, count_args = filters.sql()
                detail_count = session.connection.execute(
                    f"SELECT COUNT(*) FROM enriched_attendance WHERE {where}",
                    count_args,
                ).fetchone()[0]
                self._write_stream(
                    workbook,
                    "Detail_Data",
                    DETAIL_LABELS,
                    session.stream(detail_sql, detail_args),
                    report,
                    numbered=detail_count > self.sheet_row_limit - 1,
                )
                employee_sql = f"""WITH filtered AS ({relation}) SELECT emplid, MAX(employee_name) name,
                    classification, branch, department, division, pay_group, SUM(total_records) records,
                    SUM(late_count) late_records, SUM(total_ot_minutes)/60.0 ot_hours,
                    SUM(ot_amount) ot_amount, SUM(meal_ot_amount) meal_amount
                    FROM filtered GROUP BY emplid, classification, branch, department, division, pay_group
                    ORDER BY emplid, classification, branch, department, division, pay_group"""
                self._write_stream(
                    workbook,
                    "Employee_Summary",
                    (
                        "EMPLID",
                        "Name",
                        "Classification",
                        "Branch",
                        "Department",
                        "Division",
                        "Pay Group",
                        "Records",
                        "Late Records",
                        "OT Hours",
                        "OT Amount",
                        "Meal Amount",
                    ),
                    session.stream(employee_sql, args),
                    report,
                )
                # Registry is audit context for files contributing attendance in this scope.
                if filters == AnalyticsFilters():
                    source_sql, source_args = (
                        f"SELECT {','.join(SOURCE_COLUMNS)} FROM source_registry ORDER BY source_type,filename",
                        (),
                    )
                else:
                    w, source_args = filters.sql(alias="e")
                    source_sql = f"""SELECT {",".join(SOURCE_COLUMNS)} FROM source_registry WHERE source_id IN (
                        SELECT DISTINCT e.source_id FROM enriched_attendance e WHERE {w}
                    ) ORDER BY source_type,filename"""
                self._write_stream(
                    workbook,
                    "Source_Files",
                    (
                        "File",
                        "Source Type",
                        "Rows",
                        "Status",
                        "Last Import",
                        "Path",
                        "Error",
                    ),
                    session.stream(source_sql, source_args),
                    report,
                )
                descriptor, filename = tempfile.mkstemp(
                    prefix=".attendance_ot_", suffix=".xlsx", dir=path.parent
                )
                os.close(descriptor)
                temp_path = Path(filename)
                workbook.save(temp_path)
            os.replace(temp_path, path)
            report(f"Export completed: {detail_count:,} detail rows")
            return path
        finally:
            workbook.close()
            if temp_path is not None and temp_path.exists():
                temp_path.unlink()

    def _write_stream(self, workbook, name, headers, rows, report, *, numbered=False):
        sheet_no, count = 1, 0
        sheet = self._sheet(workbook, f"{name}_01" if numbered else name, headers)
        report(f"Export {name}")
        for row in rows:
            if count and count % (self.sheet_row_limit - 1) == 0:
                if sheet_no == 1 and not numbered:
                    sheet.title = f"{name}_01"
                sheet_no += 1
                sheet = self._sheet(workbook, f"{name}_{sheet_no:02d}", headers)
            self._append(sheet, tuple(row))
            count += 1
            if count % 10000 == 0:
                report(f"Export {name}: {count:,} rows")

    @staticmethod
    def _sheet(workbook, name, headers):
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = "C2"
        sheet.sheet_view.showGridLines = False
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                42, max(16, len(header) + 3)
            )
        cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=header)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cells.append(cell)
        sheet.append(cells)
        return sheet

    @staticmethod
    def _append(sheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                # Literal text also protects source strings beginning with '='.
                cell.data_type = "s"
            elif isinstance(value, float):
                cell.number_format = "#,##0.00"
            cell.font = Font(name="Arial", size=10)
            cells.append(cell)
        sheet.append(cells)
