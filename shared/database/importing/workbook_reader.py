"""Read workbook cells without saving or following external links."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from shared.database.importing.exceptions import WorkbookReadError
from shared.database.importing.models import (
    CellData,
    ConfigImportIssue,
    IssueSeverity,
    SheetData,
    WorkbookData,
)
from shared.database.importing.workbook_detector import detect_workbook


def read_workbook(path: str | Path) -> tuple[WorkbookData, tuple[ConfigImportIssue, ...]]:
    """Read all non-empty cells in safe read-only mode."""

    detection = detect_workbook(path)
    issues: list[ConfigImportIssue] = []
    try:
        workbook = load_workbook(
            detection.path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise WorkbookReadError(
            f"Unable to read workbook {detection.path}: {exc}"
        ) from exc

    sheets: dict[str, SheetData] = {}
    try:
        for sheet in workbook.worksheets:
            rows: list[tuple[CellData, ...]] = []
            for row_number, raw_row in enumerate(
                sheet.iter_rows(),
                start=1,
            ):
                cells = tuple(
                    CellData(
                        value=cell.value,
                        row_number=getattr(cell, "row", row_number),
                        column_number=getattr(cell, "column", column_number),
                        data_type=getattr(cell, "data_type", "n"),
                        number_format=getattr(
                            cell,
                            "number_format",
                            "General",
                        ),
                    )
                    for column_number, cell in enumerate(raw_row, start=1)
                )
                if any(cell.value is not None for cell in cells):
                    rows.append(cells)
                    for cell in cells:
                        if cell.data_type == "f":
                            issues.append(
                                ConfigImportIssue(
                                    code="FORMULA_CELL_DETECTED",
                                    severity=IssueSeverity.WARNING,
                                    module="UNKNOWN",
                                    sheet=sheet.title,
                                    row_number=cell.row_number,
                                    field=str(cell.column_number),
                                    message=(
                                        "Formula is preserved as text and is "
                                        "not evaluated during import preview."
                                    ),
                                )
                            )
            sheets[sheet.title] = SheetData(
                name=sheet.title,
                rows=tuple(rows),
            )
    finally:
        workbook.close()

    return WorkbookData(detection=detection, sheets=sheets), tuple(issues)
