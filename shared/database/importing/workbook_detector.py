"""Workbook identity detection based on structure, not filename."""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from shared.database.importing.constants import WORKBOOK_SIGNATURES
from shared.database.importing.exceptions import WorkbookReadError
from shared.database.importing.models import (
    WorkbookDetectionResult,
    WorkbookIdentity,
)


def detect_workbook(path: str | Path) -> WorkbookDetectionResult:
    """Detect one workbook from required sheet signatures."""

    workbook_path = Path(path)
    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise WorkbookReadError(
            f"Unable to read workbook {workbook_path}: {exc}"
        ) from exc

    try:
        sheets = frozenset(workbook.sheetnames)
        matches = [
            identity
            for identity, signature in WORKBOOK_SIGNATURES.items()
            if signature.issubset(sheets)
        ]
        if len(matches) == 1:
            identity = matches[0]
        elif not matches:
            identity = WorkbookIdentity.UNKNOWN
        else:
            identity = WorkbookIdentity.AMBIGUOUS

        detection_warnings: list[str] = []
        if (
            identity == WorkbookIdentity.HRIS_LEGACY
            and "Assisted_Steps" not in sheets
        ):
            detection_warnings.append("HRIS_ASSISTED_STEPS_MISSING")

        normalized = _normalized_workbook_payload(workbook)
    finally:
        workbook.close()

    return WorkbookDetectionResult(
        path=workbook_path,
        identity=identity,
        sheet_names=tuple(sorted(sheets)),
        sha256=_sha256(workbook_path),
        normalized_content_hash=hashlib.sha256(
            normalized.encode("utf-8")
        ).hexdigest(),
        warnings=tuple(detection_warnings),
    )


def _normalized_workbook_payload(workbook: Any) -> str:
    payload: list[Any] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            values = [_json_value(cell.value) for cell in row]
            while values and values[-1] is None:
                values.pop()
            if values:
                rows.append(values)
        payload.append((sheet.title, rows))
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file_handle:
        for chunk in iter(lambda: file_handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
