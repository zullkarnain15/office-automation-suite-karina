"""Reusable light professional styling for OAS-K workbooks."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "1F4E78"
BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
LIGHT_REQUIRED = "FFF2CC"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="C9D1D9")


def style_title(sheet: Worksheet, title: str, last_column: int) -> None:
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column,
    )
    cell = sheet.cell(1, 1, title)
    cell.font = Font(name="Segoe UI", size=14, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24


def style_note(sheet: Worksheet, note: str, last_column: int) -> None:
    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=last_column,
    )
    cell = sheet.cell(2, 1, note)
    cell.font = Font(name="Segoe UI", size=9, italic=True, color="44546A")
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 30


def style_headers(sheet: Worksheet, headers: tuple[str, ...], row: int = 3) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row, column, header)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )
    sheet.row_dimensions[row].height = 30


def style_data_area(
    sheet: Worksheet,
    *,
    first_row: int,
    last_row: int,
    last_column: int,
) -> None:
    for row in sheet.iter_rows(
        min_row=first_row,
        max_row=last_row,
        min_col=1,
        max_col=last_column,
    ):
        for cell in row:
            cell.font = Font(name="Segoe UI", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )


def mark_required(cell: object) -> None:
    cell.fill = PatternFill("solid", fgColor=LIGHT_REQUIRED)


def set_column_widths(
    sheet: Worksheet,
    headers: tuple[str, ...],
) -> None:
    for index, header in enumerate(headers, start=1):
        width = min(max(len(header) + 3, 14), 35)
        if any(token in header for token in ("description", "template", "path")):
            width = 34
        elif "email" in header:
            width = 28
        sheet.column_dimensions[get_column_letter(index)].width = width
