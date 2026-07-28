"""Build the official editable OAS-K configuration template."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from shared.database.exporting.constants import (
    BOOLEAN_COLUMNS,
    DATA_HEADERS,
    DATA_START_ROW,
    HEADER_ROW,
    MAX_TEMPLATE_ROWS,
    SETTING_DEFINITIONS,
    SETTING_HEADERS,
    SHEET_ORDER,
    TEMPLATE_VERSION,
    WORKFLOW_COLUMNS,
)
from shared.database.exporting.exceptions import (
    ConfigExportError,
    WorkbookValidationError,
)
from shared.database.exporting.models import ExportResult
from shared.database.exporting.style import (
    mark_required,
    set_column_widths,
    style_data_area,
    style_headers,
    style_note,
    style_title,
)
from shared.database.exporting.workbook_validator import (
    validate_configuration_workbook,
)
from shared.database.importing.models import WorkbookIdentity
from shared.database.constants import SCHEMA_VERSION
from shared.database.time_utils import current_timestamp


def build_configuration_template(
    output_path: Path,
    *,
    overwrite: bool = False,
    create_parent: bool = False,
) -> ExportResult:
    """Create and validate the deterministic DB2B unified template."""

    generated_at = current_timestamp()
    path = _prepare_output_path(
        output_path,
        overwrite=overwrite,
        create_parent=create_parent,
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "OAS-K Configuration Template"
    workbook.properties.subject = "Office Automation Suite - Karina"
    workbook.properties.creator = "OAS-K"
    workbook.properties.keywords = (
        f"OAS-K, configuration, template, version {TEMPLATE_VERSION}"
    )

    _build_guide(workbook.create_sheet("Guide"), generated_at)
    for sheet_name in SHEET_ORDER[1:]:
        sheet = workbook.create_sheet(sheet_name)
        if sheet_name in SETTING_DEFINITIONS:
            _build_settings_sheet(sheet, sheet_name)
        else:
            _build_data_sheet(sheet, sheet_name)

    try:
        workbook.save(path)
    except Exception as exc:
        raise ConfigExportError(f"Unable to save workbook {path}: {exc}") from exc
    finally:
        workbook.close()

    validation = validate_configuration_workbook(path)
    if not validation.is_valid:
        raise WorkbookValidationError(
            "Generated workbook is invalid: " + "; ".join(validation.errors)
        )
    return ExportResult(
        output_path=path,
        workbook_identity=WorkbookIdentity.OAS_K_UNIFIED,
        sheet_names=SHEET_ORDER,
        validation=validation,
        exported_at=generated_at,
    )


def _prepare_output_path(
    output_path: str | Path,
    *,
    overwrite: bool,
    create_parent: bool,
) -> Path:
    path = Path(output_path).expanduser().resolve()
    if path.suffix.casefold() != ".xlsx":
        raise ConfigExportError("Output file must use the .xlsx extension.")
    if path.exists() and not overwrite:
        raise FileExistsError(f"Output workbook already exists: {path}")
    if not path.parent.exists():
        if create_parent:
            path.parent.mkdir(parents=True, exist_ok=True)
        else:
            raise FileNotFoundError(
                f"Output parent directory does not exist: {path.parent}"
            )
    return path


def _build_guide(sheet: object, generated_at: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 78
    sheet.merge_cells("A1:B1")
    title = sheet["A1"]
    title.value = "Panduan Konfigurasi OAS-K"
    title.font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    rows = (
        ("Nama aplikasi", "Office Automation Suite – Karina (OAS-K)"),
        ("Nama template", "OAS-K Configuration Template"),
        ("Versi template", TEMPLATE_VERSION),
        ("Versi schema", SCHEMA_VERSION),
        ("Generated timestamp", generated_at),
        ("Identitas", "OAS_K_UNIFIED"),
        ("Tujuan", "Template konfigurasi untuk Attendance, Outlook Revisi, HRIS, Comparison, dan Attachment Consolidation Settings."),
        ("Cara mengisi", "Isi sel data mulai baris 4. Jangan mengubah nama, urutan, atau header sheet."),
        ("Nilai wajib", "Sel berwarna kuning wajib diperiksa/diisi sebelum import."),
        ("Boolean", "Gunakan hanya TRUE atau FALSE."),
        ("Workflow", "Gunakan HO atau BRANCH; ALL hanya pada rule validasi Outlook bila relevan."),
        ("Periode", "Periode tanggal bersama dikelola di Global_Settings. Outlook memakai payroll_period khusus dengan format MM-YYYY."),
        ("Path", "Gunakan path operasional. Jangan menyalin path developer, sample, atau mock."),
        ("Outlook", "Default aman: auto_reply_enabled=FALSE dan send_mode=DRAFT. Gmail test tidak disediakan."),
        ("HRIS", "run_control_id wajib diperlakukan sebagai TEXT agar leading zero seperti 001 tetap utuh."),
        ("Multiline", "Baris baru dan placeholder pada reply template harus dipertahankan apa adanya."),
        ("Keamanan", "Workbook ini tidak mengandung formula, macro, external link, atau secret."),
        ("Global", "Output dan periode bersama dikelola hanya melalui Global_Settings."),
        ("Media Excel", "Excel adalah media import/export; SQLite menjadi konfigurasi aktif setelah preview dan commit eksplisit."),
        ("Status engine", "Engine lama belum otomatis membaca SQLite sampai migrasi/aktivasi berikutnya disetujui."),
        ("Import", "Lakukan preview terlebih dahulu; commit konfigurasi harus eksplisit."),
        ("Peringatan", "Template kosong dapat memerlukan nilai operasional wajib sebelum lolos validasi import."),
    )
    for row_number, (label, value) in enumerate(rows, start=3):
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, value)
        for column in (1, 2):
            cell = sheet.cell(row_number, column)
            cell.font = Font(name="Segoe UI", size=10, bold=column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row_number % 2:
            for column in (1, 2):
                sheet.cell(row_number, column).fill = PatternFill(
                    "solid", fgColor="F2F2F2"
                )
        sheet.row_dimensions[row_number].height = 32
    sheet.freeze_panes = "A3"


def _build_settings_sheet(sheet: object, sheet_name: str) -> None:
    definitions = SETTING_DEFINITIONS[sheet_name]
    title = (
        "Attachment Consolidation Settings"
        if sheet_name == "Attachment_Consolidation"
        else sheet_name.replace("_", " ")
    )
    style_title(sheet, title, len(SETTING_HEADERS))
    style_note(
        sheet,
        "Struktur key/value. Kolom required adalah informasi dan tidak diimpor.",
        len(SETTING_HEADERS),
    )
    style_headers(sheet, SETTING_HEADERS)
    last_row = DATA_START_ROW + len(definitions) - 1
    style_data_area(
        sheet,
        first_row=DATA_START_ROW,
        last_row=last_row,
        last_column=len(SETTING_HEADERS),
    )
    for row_number, definition in enumerate(definitions, start=DATA_START_ROW):
        sheet.cell(row_number, 1, definition.key)
        sheet.cell(row_number, 2, _excel_value(definition.default))
        sheet.cell(row_number, 3, "TRUE" if definition.required else "FALSE")
        sheet.cell(row_number, 4, definition.description)
        if definition.required:
            mark_required(sheet.cell(row_number, 2))
        if definition.value_type == "DATE":
            sheet.cell(row_number, 2).number_format = "yyyy-mm-dd"
        if definition.value_type == "TEXT":
            sheet.cell(row_number, 2).number_format = "@"
        if definition.value_type == "BOOLEAN":
            _add_list_validation(sheet, f"B{row_number}", ("TRUE", "FALSE"))
        if definition.value_type == "WORKFLOW":
            _add_list_validation(sheet, f"B{row_number}", ("HO", "BRANCH"))
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:D{last_row}"
    set_column_widths(sheet, SETTING_HEADERS)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 70


def _build_data_sheet(sheet: object, sheet_name: str) -> None:
    headers = DATA_HEADERS[sheet_name]
    style_title(sheet, sheet_name.replace("_", " "), len(headers))
    style_note(
        sheet,
        "Isi data mulai baris 4. Jangan ubah header; baris kosong akan diabaikan.",
        len(headers),
    )
    style_headers(sheet, headers)
    style_data_area(
        sheet,
        first_row=DATA_START_ROW,
        last_row=DATA_START_ROW,
        last_column=len(headers),
    )
    set_column_widths(sheet, headers)
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = (
        f"A{HEADER_ROW}:{_column_letter(len(headers))}{HEADER_ROW}"
    )

    for index, header in enumerate(headers, start=1):
        column = _column_letter(index)
        target = f"{column}{DATA_START_ROW}:{column}{MAX_TEMPLATE_ROWS}"
        if header in BOOLEAN_COLUMNS:
            _add_list_validation(sheet, target, ("TRUE", "FALSE"))
        elif header in WORKFLOW_COLUMNS:
            values = (
                ("HO", "BRANCH", "ALL")
                if sheet_name == "Outlook_Validation_Rules"
                else ("HO", "BRANCH")
            )
            _add_list_validation(sheet, target, values)
        elif header == "recipient_type":
            values = (
                ("SENDER", "PIC_HR")
                if sheet_name == "Outlook_Reply_Templates"
                else ("TO", "CC")
            )
            _add_list_validation(sheet, target, values)
        elif header == "action":
            _add_list_validation(
                sheet,
                target,
                (
                    "click", "click_type", "type", "press", "attach_file",
                    "wait", "manual_continue",
                ),
            )
        elif header == "input_source":
            _add_list_validation(
                sheet,
                target,
                ("NONE", "RUN_CONTROL_ID", "START_DATE", "END_DATE", "TXT_FILE_PATH"),
            )
        elif header == "method":
            _add_list_validation(
                sheet, target, ("coordinate", "playwright", "manual", "assisted")
            )
        if header == "run_control_id":
            for row in range(DATA_START_ROW, MAX_TEMPLATE_ROWS + 1):
                sheet.cell(row, index).number_format = "@"
        if header in {"body_template", "description"}:
            for row in range(DATA_START_ROW, MAX_TEMPLATE_ROWS + 1):
                sheet.cell(row, index).alignment = Alignment(
                    vertical="top", wrap_text=True
                )

    active_columns = [
        index
        for index, header in enumerate(headers, start=1)
        if header == "is_active"
    ]
    for index in active_columns:
        column = _column_letter(index)
        sheet.conditional_formatting.add(
            f"{column}{DATA_START_ROW}:{column}{MAX_TEMPLATE_ROWS}",
            FormulaRule(
                formula=[f'{column}{DATA_START_ROW}="FALSE"'],
                fill=PatternFill("solid", fgColor="F2F2F2"),
            ),
        )


def _add_list_validation(
    sheet: object,
    target: str,
    values: tuple[str, ...],
) -> None:
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
    )
    validation.error = "Pilih nilai dari daftar yang tersedia."
    validation.errorTitle = "Nilai tidak valid"
    validation.prompt = "Gunakan salah satu nilai yang tersedia."
    validation.promptTitle = "OAS-K"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(target)


def _excel_value(value: object) -> object:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return value


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result
