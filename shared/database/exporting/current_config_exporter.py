"""Export current schema-v1 configuration to the unified DB2B workbook."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from shared.database.connection_factory import SQLiteConnectionFactory
from shared.database.database_validator import DatabaseValidator
from shared.database.exporting.constants import (
    BOOLEAN_COLUMNS,
    DATA_HEADERS,
    DATA_START_ROW,
    HEADER_ROW,
    SETTING_DEFINITIONS,
    SHEET_ORDER,
    SHEET_TABLES,
)

SENDER_WORKFLOW_BY_SHEET = {
    "Outlook_HO_Senders": "HO",
    "Outlook_Branch_Senders": "BRANCH",
}
from shared.database.exporting.exceptions import (
    ConfigExportError,
    WorkbookValidationError,
)
from shared.database.exporting.models import ExportResult
from shared.database.exporting.template_builder import (
    _prepare_output_path,
    build_configuration_template,
)
from shared.database.exporting.workbook_validator import (
    validate_configuration_workbook,
)
from shared.database.importing.models import WorkbookIdentity
from shared.database.time_utils import current_timestamp

_FORBIDDEN_COLUMN_TOKENS = (
    "password",
    "passwd",
    "secret",
    "token",
    "credential",
    "api_key",
)


def export_current_configuration(
    database_path: str | Path,
    output_path: str | Path,
    *,
    overwrite: bool = False,
    create_parent: bool = False,
) -> ExportResult:
    """Read the current database in read-only mode and export safely."""

    database = Path(database_path).expanduser().resolve()
    output = _prepare_output_path(
        output_path,
        overwrite=overwrite,
        create_parent=create_parent,
    )
    DatabaseValidator().validate_or_raise(database)
    exported_at = current_timestamp()

    temp_path: Path | None = None
    try:
        descriptor, temp_name = tempfile.mkstemp(
            prefix=".oas-k-config-",
            suffix=".xlsx",
            dir=output.parent,
        )
        os.close(descriptor)
        temp_path = Path(temp_name)
        build_configuration_template(temp_path, overwrite=True)
        configuration, metadata = _read_configuration(database)
        _populate_workbook(
            temp_path,
            configuration=configuration,
            metadata=metadata,
            exported_at=exported_at,
        )
        validation = validate_configuration_workbook(temp_path)
        if not validation.is_valid:
            raise WorkbookValidationError(
                "Exported workbook is invalid: "
                + "; ".join(validation.errors)
            )
        os.replace(temp_path, output)
        temp_path = None
        final_validation = validate_configuration_workbook(output)
        return ExportResult(
            output_path=output,
            workbook_identity=WorkbookIdentity.OAS_K_UNIFIED,
            sheet_names=SHEET_ORDER,
            validation=final_validation,
            exported_at=exported_at,
        )
    except (ConfigExportError, FileExistsError, FileNotFoundError):
        raise
    except Exception as exc:
        raise ConfigExportError(
            f"Unable to export current configuration: {exc}"
        ) from exc
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def _read_configuration(
    database_path: Path,
) -> tuple[dict[str, tuple[dict[str, Any], ...]], dict[str, Any]]:
    factory = SQLiteConnectionFactory()
    result: dict[str, tuple[dict[str, Any], ...]] = {}
    with factory.connect(database_path, read_only=True) as connection:
        metadata_row = connection.execute(
            """
            SELECT database_uuid, schema_version, application_version
            FROM database_metadata
            WHERE metadata_id = 1
            """
        ).fetchone()
        if metadata_row is None:
            raise ConfigExportError("Database metadata singleton is missing.")

        for sheet_name, table in SHEET_TABLES.items():
            columns = _export_columns(sheet_name)
            _assert_no_secret_columns(columns)
            sql_columns = ", ".join(f'"{column}"' for column in columns)
            order_clause = _order_clause(table, columns)
            rows = connection.execute(
                f'SELECT {sql_columns} FROM "{table}"{order_clause}'
            ).fetchall()
            result[table] = tuple(dict(row) for row in rows)
    return result, dict(metadata_row)


def _export_columns(sheet_name: str) -> tuple[str, ...]:
    if sheet_name in SETTING_DEFINITIONS:
        return tuple(item.key for item in SETTING_DEFINITIONS[sheet_name])
    if sheet_name in SENDER_WORKFLOW_BY_SHEET:
        return ("workflow", *DATA_HEADERS[sheet_name])
    return DATA_HEADERS[sheet_name]


def _assert_no_secret_columns(columns: tuple[str, ...]) -> None:
    unsafe = [
        column
        for column in columns
        if any(token in column.casefold() for token in _FORBIDDEN_COLUMN_TOKENS)
    ]
    if unsafe:
        raise ConfigExportError(
            "Secret-bearing columns cannot be exported: " + ", ".join(unsafe)
        )


def _order_clause(table: str, columns: tuple[str, ...]) -> str:
    preferred = {
        "attendance_sources": ("workflow", "sort_order", "source_code"),
        "outlook_sender_master": (
            "workflow", "company_code", "branch_code", "sender_email"
        ),
        "outlook_subject_rules": ("workflow", "subject_pattern"),
        "outlook_attachment_rules": ("workflow", "extension"),
        "outlook_validation_rules": ("rule_code", "workflow"),
        "outlook_reply_templates": ("reply_code",),
        "outlook_summary_recipients": (
            "recipient_type", "sort_order", "email_address"
        ),
        "hris_run_controls": ("workflow", "sequence"),
        "hris_assisted_steps": ("sequence",),
    }.get(table, ())
    usable = tuple(column for column in preferred if column in columns)
    if not usable:
        return ""
    return " ORDER BY " + ", ".join(f'"{column}"' for column in usable)


def _populate_workbook(
    path: Path,
    *,
    configuration: dict[str, tuple[dict[str, Any], ...]],
    metadata: dict[str, Any],
    exported_at: str,
) -> None:
    workbook = load_workbook(path, keep_links=False)
    try:
        guide = workbook["Guide"]
        next_row = guide.max_row + 2
        guide.cell(next_row, 1, "Metadata export")
        guide.cell(next_row, 2, "Nilai berikut bersifat informasional.")
        metadata_values = (
            ("exported_at", exported_at),
            ("schema_version", metadata["schema_version"]),
            ("application_version", metadata["application_version"]),
            ("database_uuid", metadata["database_uuid"]),
        )
        for offset, (key, value) in enumerate(metadata_values, start=1):
            guide.cell(next_row + offset, 1, key)
            guide.cell(next_row + offset, 2, value)

        for sheet_name, table in SHEET_TABLES.items():
            rows = configuration[table]
            sheet = workbook[sheet_name]
            if sheet_name in SETTING_DEFINITIONS:
                _populate_settings(
                    sheet,
                    SETTING_DEFINITIONS[sheet_name],
                    rows[0] if rows else None,
                )
            else:
                _populate_table(
                    sheet,
                    DATA_HEADERS[sheet_name],
                    _rows_for_sheet(sheet_name, rows),
                )
        workbook.save(path)
    finally:
        workbook.close()


def _populate_settings(
    sheet: object,
    definitions: tuple[object, ...],
    row: dict[str, Any] | None,
) -> None:
    for row_number, definition in enumerate(
        definitions,
        start=DATA_START_ROW,
    ):
        value = None if row is None else row.get(definition.key)
        cell = sheet.cell(row_number, 2, _to_excel(value, definition.key))
        if definition.value_type == "TEXT":
            cell.number_format = "@"


def _populate_table(
    sheet: object,
    headers: tuple[str, ...],
    rows: tuple[dict[str, Any], ...],
) -> None:
    for row_number, row in enumerate(rows, start=DATA_START_ROW):
        for column_number, header in enumerate(headers, start=1):
            cell = sheet.cell(
                row_number,
                column_number,
                _to_excel(row.get(header), header),
            )
            if header == "run_control_id":
                cell.number_format = "@"
            if header in {"body_template", "description"}:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    last_row = max(HEADER_ROW, DATA_START_ROW + len(rows) - 1)
    sheet.auto_filter.ref = (
        f"A{HEADER_ROW}:{get_column_letter(len(headers))}{last_row}"
    )


def _rows_for_sheet(
    sheet_name: str,
    rows: tuple[dict[str, Any], ...],
) -> tuple[dict[str, Any], ...]:
    workflow = SENDER_WORKFLOW_BY_SHEET.get(sheet_name)
    if workflow is None:
        return rows
    return tuple(
        row for row in rows if str(row.get("workflow", "")).upper() == workflow
    )


def _to_excel(value: Any, column: str) -> Any:
    if value is None:
        return None
    if column in BOOLEAN_COLUMNS or column in {
        "use_global_output",
        "use_global_period",
        "generate_report_default",
        "auto_reply_enabled",
        "save_smtp_copy_to_sent",
        "browser_headless",
        "stop_on_first_failure",
        "manual_recovery_enabled",
        "require_profile_match",
        "verification_enabled",
        "manual_verification_on_unknown",
        "manual_verification_on_error",
    }:
        return "TRUE" if bool(value) else "FALSE"
    if column in {"run_control_id", "payroll_period"}:
        return str(value)
    return value
