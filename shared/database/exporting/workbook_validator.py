"""Structural validator for generated OAS-K configuration workbooks."""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from shared.database.exporting.constants import (
    BOOLEAN_COLUMNS,
    DATA_HEADERS,
    DATA_START_ROW,
    HEADER_ROW,
    MAX_TEMPLATE_ROWS,
    SETTING_DEFINITIONS,
    SETTING_HEADERS,
    SHEET_ORDER,
    WORKFLOW_COLUMNS,
)
from shared.database.exporting.models import WorkbookValidationResult
from shared.database.importing.models import WorkbookIdentity
from shared.database.importing.workbook_detector import detect_workbook


def validate_configuration_workbook(
    path: str | Path,
) -> WorkbookValidationResult:
    """Validate exact sheets, safety properties, and editable structure."""

    workbook_path = Path(path).expanduser().resolve()
    errors: list[str] = []
    validation_warnings: list[str] = []
    formulas = 0
    external_links = 0
    has_macros = False

    if not workbook_path.is_file():
        return WorkbookValidationResult(
            path=workbook_path,
            is_valid=False,
            sheet_names=(),
            formula_count=0,
            external_link_count=0,
            has_macros=False,
            errors=(f"Workbook does not exist: {workbook_path}",),
        )

    try:
        workbook = load_workbook(
            workbook_path,
            data_only=False,
            read_only=False,
            keep_links=True,
            keep_vba=False,
        )
    except Exception as exc:
        return WorkbookValidationResult(
            path=workbook_path,
            is_valid=False,
            sheet_names=(),
            formula_count=0,
            external_link_count=0,
            has_macros=False,
            errors=(f"Workbook cannot be opened: {exc}",),
        )

    try:
        sheet_names = tuple(workbook.sheetnames)
        if sheet_names != SHEET_ORDER:
            errors.append(
                "Sheet order/names do not match the exact OAS-K contract."
            )
        overlong_sheet_names = tuple(
            name for name in sheet_names if len(name) > 31
        )
        if overlong_sheet_names:
            errors.append(
                "Worksheet names exceed Excel's 31-character limit: "
                + ", ".join(overlong_sheet_names)
            )

        try:
            with ZipFile(workbook_path) as archive:
                has_macros = any(
                    name.casefold().endswith("vbaproject.bin")
                    for name in archive.namelist()
                )
        except BadZipFile:
            has_macros = False
        if has_macros or workbook_path.suffix.casefold() == ".xlsm":
            errors.append("Macros are not permitted.")

        links = getattr(workbook, "_external_links", ())
        external_links = len(links)
        if external_links:
            errors.append("External workbook links are not permitted.")
        defined_names = tuple(workbook.defined_names)
        if defined_names:
            errors.append("Defined names are not permitted.")

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formulas += 1
            if sheet.title != "Guide":
                if sheet.freeze_panes != "A4":
                    errors.append(f"{sheet.title}: freeze panes must be A4.")
                if not sheet.auto_filter.ref:
                    errors.append(f"{sheet.title}: AutoFilter is missing.")

        if formulas:
            errors.append("Formulas are not permitted.")

        for sheet_name in SHEET_ORDER[1:]:
            if sheet_name not in workbook.sheetnames:
                continue
            expected = (
                DATA_HEADERS[sheet_name]
                if sheet_name in DATA_HEADERS
                else SETTING_HEADERS
            )
            actual = tuple(
                workbook[sheet_name].cell(HEADER_ROW, column).value
                for column in range(1, len(expected) + 1)
            )
            if actual != expected:
                errors.append(f"{sheet_name}: headers do not match contract.")
            _validate_data_validations(
                workbook[sheet_name],
                sheet_name,
                errors,
            )

        if "HRIS_Run_Controls" in workbook.sheetnames:
            sheet = workbook["HRIS_Run_Controls"]
            run_control_column = DATA_HEADERS[
                "HRIS_Run_Controls"
            ].index("run_control_id") + 1
            if sheet.cell(4, run_control_column).number_format != "@":
                errors.append(
                    "HRIS_Run_Controls: run_control_id must use TEXT format."
                )

        if "Outlook_Reply_Templates" in workbook.sheetnames:
            sheet = workbook["Outlook_Reply_Templates"]
            body_column = DATA_HEADERS[
                "Outlook_Reply_Templates"
            ].index("body_template") + 1
            if not sheet.cell(4, body_column).alignment.wrap_text:
                errors.append(
                    "Outlook_Reply_Templates: body_template must wrap text."
                )

        try:
            identity = detect_workbook(workbook_path).identity
            if identity is not WorkbookIdentity.OAS_K_UNIFIED:
                errors.append(
                    "Workbook detector does not identify this as OAS_K_UNIFIED."
                )
        except Exception as exc:
            errors.append(f"Workbook detection failed: {exc}")
    finally:
        workbook.close()

    return WorkbookValidationResult(
        path=workbook_path,
        is_valid=not errors,
        sheet_names=sheet_names,
        formula_count=formulas,
        external_link_count=external_links,
        has_macros=has_macros,
        errors=tuple(errors),
        warnings=tuple(validation_warnings),
    )


def _validate_data_validations(
    sheet: object,
    sheet_name: str,
    errors: list[str],
) -> None:
    validations = tuple(sheet.data_validations.dataValidation)
    if sheet_name in SETTING_DEFINITIONS:
        for row_number, definition in enumerate(
            SETTING_DEFINITIONS[sheet_name],
            start=DATA_START_ROW,
        ):
            expected_values = {
                "BOOLEAN": '"TRUE,FALSE"',
                "WORKFLOW": '"HO,BRANCH"',
            }.get(definition.value_type)
            if expected_values and not _has_validation(
                validations,
                f"B{row_number}",
                expected_values,
            ):
                errors.append(
                    f"{sheet_name}: validation missing at B{row_number}."
                )
        return

    for column_number, header in enumerate(
        DATA_HEADERS[sheet_name],
        start=1,
    ):
        formula = None
        if header in BOOLEAN_COLUMNS:
            formula = '"TRUE,FALSE"'
        elif header in WORKFLOW_COLUMNS:
            formula = (
                '"HO,BRANCH,ALL"'
                if sheet_name == "Outlook_Validation_Rules"
                else '"HO,BRANCH"'
            )
        if formula is None:
            continue
        letter = _column_letter(column_number)
        target = f"{letter}{DATA_START_ROW}:{letter}{MAX_TEMPLATE_ROWS}"
        if not _has_validation(validations, target, formula):
            errors.append(
                f"{sheet_name}: validation missing at {target}."
            )


def _has_validation(
    validations: tuple[object, ...],
    target: str,
    formula: str,
) -> bool:
    return any(
        str(validation.sqref) == target
        and str(validation.formula1) == formula
        for validation in validations
    )


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result
