"""Read Attachment Consolidation reports without modifying them."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from utilities.att_data_repair.constants import (
    INVALID_RECORDS_COLUMNS,
    INVALID_RECORDS_SHEET,
    REQUIRED_SHEETS,
    VALID_RECORDS_COLUMNS,
    VALID_RECORDS_SHEET,
)
from utilities.att_data_repair.models import (
    MissingRequiredColumnError,
    MissingRequiredSheetError,
    SourceRecord,
    SourceWorkbookError,
)


class AttDataRepairReportReader:
    """Read Valid_Records and Invalid_Records from an AC report."""

    def read(self, path: str | Path) -> tuple[SourceRecord, ...]:
        """Return source records in workbook order and close the workbook."""

        source_path = Path(path)
        if source_path.suffix.casefold() != ".xlsx" or not source_path.is_file():
            raise SourceWorkbookError(
                f"Source workbook Att Data Repair tidak valid: {source_path}"
            )

        try:
            workbook = load_workbook(
                source_path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception as exc:
            raise SourceWorkbookError(
                f"Source workbook tidak dapat dibaca: {source_path}: {exc}"
            ) from exc

        try:
            sheet_map = self._required_sheet_map(workbook.sheetnames)
            records: list[SourceRecord] = []
            for canonical_sheet in REQUIRED_SHEETS:
                actual_sheet = sheet_map[canonical_sheet]
                expected_columns = (
                    VALID_RECORDS_COLUMNS
                    if canonical_sheet == VALID_RECORDS_SHEET
                    else INVALID_RECORDS_COLUMNS
                )
                sheet = workbook[actual_sheet]
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if headers is None:
                    raise MissingRequiredColumnError(
                        f"{actual_sheet}: header wajib tidak ditemukan."
                    )
                columns = self._column_map(
                    actual_sheet,
                    headers,
                    expected_columns,
                )
                for row_number, row in enumerate(rows, start=2):
                    if self._empty_row(row):
                        continue
                    record_id = f"ADR-{len(records) + 1:06d}"
                    records.append(
                        self._source_record(
                            record_id,
                            source_path,
                            actual_sheet,
                            canonical_sheet,
                            row,
                            columns,
                            row_number,
                        )
                    )
            return tuple(records)
        finally:
            workbook.close()

    @classmethod
    def _required_sheet_map(cls, sheet_names: list[str]) -> dict[str, str]:
        normalized = {_normalize_name(name): name for name in sheet_names}
        result = {}
        for required in REQUIRED_SHEETS:
            actual = normalized.get(_normalize_name(required))
            if actual is None:
                raise MissingRequiredSheetError(
                    f"Sheet wajib tidak ditemukan: {required}"
                )
            result[required] = actual
        return result

    @classmethod
    def _column_map(
        cls,
        sheet_name: str,
        headers: tuple[Any, ...],
        expected_columns: tuple[str, ...],
    ) -> dict[str, int]:
        actual = {
            _normalize_name(value): index
            for index, value in enumerate(headers)
            if str(value or "").strip()
        }
        missing = [
            column
            for column in expected_columns
            if _normalize_name(column) not in actual
        ]
        if missing:
            raise MissingRequiredColumnError(
                f"{sheet_name}: kolom wajib tidak ditemukan: "
                + ", ".join(missing)
            )
        return {column: actual[_normalize_name(column)] for column in expected_columns}

    @staticmethod
    def _source_record(
        record_id: str,
        source_path: Path,
        actual_sheet: str,
        canonical_sheet: str,
        row: tuple[Any, ...],
        columns: dict[str, int],
        row_number: int,
    ) -> SourceRecord:
        source_status = (
            _text(_value(row, columns, "Status"))
            if canonical_sheet == VALID_RECORDS_SHEET
            else _text(_value(row, columns, "Status_Code"))
        )
        source_reason = (
            ""
            if canonical_sheet == VALID_RECORDS_SHEET
            else _text(_value(row, columns, "Reason"))
        )
        raw_value = (
            _text(_value(row, columns, "Raw_Value"))
            if canonical_sheet == INVALID_RECORDS_SHEET
            else _raw_value(row)
        )
        return SourceRecord(
            record_id=record_id,
            source_workbook=source_path,
            source_sheet=actual_sheet,
            source_record_no=_value(row, columns, "No"),
            source_file=_text(_value(row, columns, "Source_File")),
            relative_path=_text(_value(row, columns, "Relative_Path")),
            source_row=_value(row, columns, "Source_Row") or row_number,
            workflow_raw=_value(row, columns, "Workflow"),
            nik_raw=_value(row, columns, "NIK"),
            date_in_raw=_value(row, columns, "Date_In"),
            time_in_raw=_value(row, columns, "Time_In"),
            date_out_raw=_value(row, columns, "Date_Out"),
            time_out_raw=_value(row, columns, "Time_Out"),
            source_status=source_status,
            source_reason=source_reason,
            raw_value=raw_value,
        )

    @staticmethod
    def _empty_row(row: tuple[Any, ...]) -> bool:
        return not row or not any(_text(value) for value in row)


def _normalize_name(value: Any) -> str:
    return re.sub(r"[\s_]+", "", str(value or "")).casefold()


def _value(row: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns[name]
    return row[index] if index < len(row) else None


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _raw_value(values: tuple[Any, ...]) -> str:
    return " | ".join(_text(value) for value in values)[:1000]
