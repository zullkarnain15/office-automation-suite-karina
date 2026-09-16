"""Excel audit report writer for Att Data Repair."""

from __future__ import annotations

import uuid
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utilities.att_data_repair.artifacts import (
    AttDataRepairJobRequest,
    JobPaths,
    ReportArtifact,
    TxtArtifact,
)
from utilities.att_data_repair.constants import (
    ENGINE_VERSION,
    FEATURE_NAME,
    INVALID_RECORDS_SHEET,
    REPORT_FILENAME_PREFIX,
    REPORT_SHEET_ORDER,
    VALID_RECORDS_SHEET,
    WORKFLOW_BRANCH,
    WORKFLOW_HO,
)
from utilities.att_data_repair.models import (
    AttDataRepairAnalysisResult,
    ReportAlreadyExistsError,
    ReportWriteError,
)
from utilities.att_data_repair.normalizer import (
    format_date,
    format_time,
    normalize_date,
    normalize_nik,
    normalize_workflow,
)
from utilities.att_data_repair.statuses import AnomalyCode, ChangeCode, FinalStatus

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
THIN_BORDER = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)
STATUS_FILLS = {
    "VALID_UNCHANGED": PatternFill("solid", fgColor="C6EFCE"),
    "SUCCESS": PatternFill("solid", fgColor="C6EFCE"),
    "INCLUDED": PatternFill("solid", fgColor="C6EFCE"),
    "REPAIRED": PatternFill("solid", fgColor="FFEB9C"),
    "PARTIAL_SUCCESS": PatternFill("solid", fgColor="FFEB9C"),
    "PARTIAL": PatternFill("solid", fgColor="FFEB9C"),
    "DEFAULT_TIME_APPLIED": PatternFill("solid", fgColor="FCE4D6"),
    "MINIMUM_DURATION_APPLIED": PatternFill("solid", fgColor="FCE4D6"),
    "FAILED": PatternFill("solid", fgColor="FFC7CE"),
    "NO_VALID_RECORDS": PatternFill("solid", fgColor="FFC7CE"),
    "EXCLUDED": PatternFill("solid", fgColor="FFC7CE"),
}
class AttDataRepairReportWriter:
    """Write the final Att Data Repair Excel audit workbook."""

    def write(
        self,
        *,
        report_folder: str | Path,
        job_id: str,
        status: str,
        request: AttDataRepairJobRequest,
        paths: JobPaths,
        analysis: AttDataRepairAnalysisResult,
        txt_artifacts: tuple[TxtArtifact, ...],
        record_txt_assignment: dict[str, str],
        generated_at: datetime,
    ) -> ReportArtifact:
        """Create one collision-safe workbook report for a job."""

        folder = Path(report_folder)
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"{REPORT_FILENAME_PREFIX}_{job_id}.xlsx"
        if target.exists():
            raise ReportAlreadyExistsError(f"Report sudah ada: {target}")
        temporary = folder / f".{target.stem}.{uuid.uuid4().hex}.tmp.xlsx"
        # Stream rows directly to the XLSX archive. A 200k-row report can
        # contain millions of cells across the audit sheets.
        workbook = Workbook(write_only=True)
        try:
            self._prepare_workbook(workbook)
            self._write_all_sheets(
                workbook,
                target,
                job_id,
                status,
                request,
                paths,
                analysis,
                txt_artifacts,
                record_txt_assignment,
                generated_at,
            )
            row_counts = {
                name: int(getattr(workbook[name], "_oas_row_count", 0))
                for name in REPORT_SHEET_ORDER
            }
            workbook.save(temporary)
            if target.exists():
                raise ReportAlreadyExistsError(f"Report sudah ada: {target}")
            temporary.rename(target)
        except Exception as exc:
            temporary.unlink(missing_ok=True)
            if isinstance(exc, ReportWriteError):
                raise
            raise ReportWriteError(f"Gagal menulis report {target}: {exc}") from exc
        finally:
            workbook.close()
        return ReportArtifact(
            file_name=target.name,
            file_path=target,
            sheet_names=REPORT_SHEET_ORDER,
            sheet_row_counts=row_counts,
            file_size_bytes=target.stat().st_size,
        )

    @staticmethod
    def _prepare_workbook(workbook: Workbook) -> None:
        for name in REPORT_SHEET_ORDER:
            sheet = workbook.create_sheet(name)
            sheet._oas_row_count = 0
        properties = workbook.properties
        properties.title = "Att Data Repair Report"
        properties.subject = "OAS-K Attendance Data Repair Audit Report"
        properties.creator = "Office Automation Suite - Karina"
        properties.description = (
            "Report hasil validasi, perbaikan, anomaly, dan output TXT "
            "Att Data Repair."
        )
        properties.keywords = "OAS-K, Att Data Repair, Attendance, HRIS"

    def _write_all_sheets(
        self,
        workbook: Workbook,
        target: Path,
        job_id: str,
        status: str,
        request: AttDataRepairJobRequest,
        paths: JobPaths,
        analysis: AttDataRepairAnalysisResult,
        txt_artifacts: tuple[TxtArtifact, ...],
        record_txt_assignment: dict[str, str],
        generated_at: datetime,
    ) -> "_ReportContext":
        context = _ReportContext(
            target=target,
            job_id=job_id,
            status=status,
            request=request,
            paths=paths,
            analysis=analysis,
            txt_artifacts=txt_artifacts,
            record_txt_assignment=record_txt_assignment,
            generated_at=generated_at,
        )
        self._guide(workbook["Guide_Status"], context)
        self._process_summary(workbook["Process_Summary"], context)
        self._summary_per_karyawan(workbook["Summary_Per_Karyawan"], context)
        self._final_records(workbook["Final_Records"], context)
        self._changed_records(workbook["Changed_Records"], context)
        self._anomaly(workbook["Anomaly"], context)
        self._change_log(workbook["Change_Log"], context)
        self._source_inventory(workbook["Source_Inventory"], context)
        return context

    def _guide(self, sheet: Any, context: "_ReportContext") -> None:
        _prepare_sheet(
            sheet,
            [
                "Status_Code",
                "Category",
                "Penjelasan_Bahasa_Indonesia",
                "Tindakan_Sistem",
                "Masuk_TXT",
                "Warna",
            ],
        )
        _append_header(sheet, ["Report Identity", "Value"])
        source_sheets = ", ".join((VALID_RECORDS_SHEET, INVALID_RECORDS_SHEET))
        rows = (
            ("Report Name", context.target.name),
            ("Module", FEATURE_NAME),
            ("Job ID", context.job_id),
            ("Job Status", context.status),
            ("Engine Version", ENGINE_VERSION),
            ("Generated At", context.generated_at),
            ("Period Start", format_date(context.analysis_request.period_start)),
            ("Period End", format_date(context.analysis_request.period_end)),
            ("Source Workbook", str(context.analysis_request.source_report)),
            ("Source Sheet", source_sheets),
            ("Output Job Folder", str(context.paths.job_folder)),
            ("TXT Unique Code", context.txt_unique_code or ""),
            ("Jumlah TXT", len(context.txt_artifacts)),
            ("Jumlah Final Record", len(context.analysis.final_records)),
            ("Jumlah Changed Record", len(context.analysis.changed_records)),
            ("Jumlah Anomaly", len(context.analysis.anomalies)),
        )
        for row in rows:
            _append_row(sheet, row)
        _append_section(sheet, "Prinsip Utama")
        for text in _GUIDE_PRINCIPLES:
            _append_row(sheet, [text])
        _append_section(sheet, "Aturan NIK")
        for text in _GUIDE_NIK:
            _append_row(sheet, [text])
        _append_section(sheet, "Aturan Workflow")
        for text in _GUIDE_WORKFLOW:
            _append_row(sheet, [text])
        _append_section(sheet, "Aturan Tanggal")
        for text in _GUIDE_DATE:
            _append_row(sheet, [text])
        _append_section(sheet, "Aturan Waktu")
        for text in _GUIDE_TIME:
            _append_row(sheet, [text])
        _append_section(sheet, "Default Waktu")
        for text in _GUIDE_DEFAULTS:
            _append_row(sheet, [text])
        _append_section(sheet, "Naming Output")
        for text in _GUIDE_NAMING:
            _append_row(sheet, [text])
        _append_section(sheet, "Tabel Referensi Status")
        _append_header(
            sheet,
            [
                "Status_Code",
                "Category",
                "Penjelasan_Bahasa_Indonesia",
                "Tindakan_Sistem",
                "Masuk_TXT",
                "Warna",
            ],
        )
        for row in _status_reference_rows():
            _append_data_row(sheet, row, status_columns=(0,))

    def _process_summary(self, sheet: Any, context: "_ReportContext") -> None:
        _prepare_sheet(
            sheet,
            [
                "Artifact_Type",
                "Workflow",
                "File_Name",
                "Row_Count",
                "File_Size_Bytes",
                "Relative_Path",
            ],
        )
        _append_section(sheet, "SECTION 1 - Job Information")
        _append_header(sheet, ["Field", "Value"])
        for row in (
            ("Job ID", context.job_id),
            ("Job Status", context.status),
            ("Engine Version", ENGINE_VERSION),
            ("Period Start", context.analysis_request.period_start),
            ("Period End", context.analysis_request.period_end),
            ("Source Workbook", str(context.analysis_request.source_report)),
            ("Generated At", context.generated_at),
            ("TXT Unique Code", context.txt_unique_code or ""),
            ("Report File", context.target.name),
        ):
            _append_row(sheet, row)
        _append_section(sheet, "SECTION 2 - Overall Metrics")
        _append_header(sheet, ["Metric", "Value"])
        for row in context.overall_metrics():
            _append_row(sheet, row)
        _append_section(sheet, "SECTION 3 - Workflow Summary")
        _append_header(sheet, ["Metric", "HO", "Branch", "Total"])
        for row in context.workflow_metrics():
            _append_row(sheet, row)
        _append_section(sheet, "SECTION 4 - Anomaly by Code")
        _append_header(sheet, ["Anomaly_Code", "Count", "Penjelasan"])
        for code, count in sorted(context.anomaly_counts.items()):
            _append_row(sheet, [code, count, _ANOMALY_DESCRIPTIONS.get(code, "")])
        _append_section(sheet, "SECTION 5 - Change by Code")
        _append_header(sheet, ["Change_Code", "Count", "Penjelasan"])
        for code, count in sorted(context.change_counts.items()):
            _append_row(sheet, [code, count, _CHANGE_DESCRIPTIONS.get(code, "")])
        _append_section(sheet, "SECTION 6 - Active Configuration")
        _append_header(sheet, ["Setting_Key", "Setting_Value"])
        for row in _active_configuration_rows(context):
            _append_row(sheet, row)
        _append_section(sheet, "SECTION 7 - Output Artifacts")
        _append_header(
            sheet,
            [
                "Artifact_Type",
                "Workflow",
                "File_Name",
                "Row_Count",
                "File_Size_Bytes",
                "Relative_Path",
            ],
        )
        for artifact in context.txt_artifacts:
            _append_row(
                sheet,
                [
                    "TXT",
                    artifact.workflow,
                    artifact.file_name,
                    artifact.row_count,
                    _file_size(artifact.file_path),
                    _relative_to_job(artifact.file_path, context.paths.job_folder),
                ]
            )
        _append_row(
            sheet,
            [
                "EXCEL_REPORT",
                "",
                context.target.name,
                "",
                "",
                _relative_to_job(context.target, context.paths.job_folder),
            ]
        )

    def _summary_per_karyawan(
        self,
        sheet: Any,
        context: "_ReportContext",
    ) -> None:
        headers = [
            "No",
            "Workflow",
            "NIK",
            "Total_Source_Record",
            "Final_Record",
            "Unchanged_Record",
            "Repaired_Record",
            "Anomaly_Record",
            "First_Date",
            "Last_Date",
            "TXT_Output_Status",
        ]
        _prepare_sheet(sheet, headers, data_sheet=True)
        _append_header(sheet, headers)
        groups = _employee_groups(context)
        for number, item in enumerate(groups, 1):
            _append_data_row(
                sheet,
                [
                    number,
                    item["workflow"],
                    item["nik"],
                    item["source"],
                    item["final"],
                    item["unchanged"],
                    item["repaired"],
                    item["anomaly"],
                    item["first_date"],
                    item["last_date"],
                    item["txt_status"],
                ],
                number_formats={2: "@", 8: "mm/dd/yyyy", 9: "mm/dd/yyyy"},
                status_columns=(10,),
            )
        _finish_sheet(sheet, len(headers), len(groups) + 1)

    def _final_records(self, sheet: Any, context: "_ReportContext") -> None:
        headers = [
            "No",
            "Record_ID",
            "Workflow",
            "NIK",
            "Date_In",
            "Time_In",
            "Date_Out",
            "Time_Out",
            "Duration_Minutes",
            "Source_Sheet",
            "Source_File",
            "Relative_Path",
            "Source_Row",
            "Source_Record_No",
            "Source_Status",
            "Final_Status",
            "TXT_File_Name",
        ]
        _prepare_sheet(sheet, headers, data_sheet=True)
        _append_header(sheet, headers)
        for number, record in enumerate(context.analysis.final_records, 1):
            source = context.source_by_id.get(record.record_id)
            _append_data_row(
                sheet,
                [
                    number,
                    record.record_id,
                    record.workflow,
                    record.nik,
                    record.date_in,
                    record.time_in,
                    record.date_out,
                    record.time_out,
                    record.duration_minutes,
                    record.source_sheet,
                    record.source_file,
                    source.relative_path if source else "",
                    record.source_row,
                    source.source_record_no if source else "",
                    source.source_status if source else "",
                    record.final_status,
                    context.record_txt_assignment.get(record.record_id, ""),
                ],
                number_formats={
                    3: "@",
                    4: "mm/dd/yyyy",
                    5: "hh:mm",
                    6: "mm/dd/yyyy",
                    7: "hh:mm",
                },
                status_columns=(15,),
            )
        _finish_sheet(sheet, len(headers), len(context.analysis.final_records) + 1)

    def _changed_records(self, sheet: Any, context: "_ReportContext") -> None:
        headers = [
            "No",
            "Record_ID",
            "Workflow",
            "Original_NIK",
            "Final_NIK",
            "Source_Sheet",
            "Source_File",
            "Relative_Path",
            "Source_Row",
            "Original_Date_In",
            "Final_Date_In",
            "Original_Time_In",
            "Final_Time_In",
            "Original_Date_Out",
            "Final_Date_Out",
            "Original_Time_Out",
            "Final_Time_Out",
            "Change_Count",
            "Change_Codes",
            "Final_Status",
            "TXT_File_Name",
        ]
        _prepare_sheet(sheet, headers, data_sheet=True)
        _append_header(sheet, headers)
        for number, record in enumerate(context.analysis.changed_records, 1):
            source = context.source_by_id.get(record.record_id)
            _append_data_row(
                sheet,
                [
                    number,
                    record.record_id,
                    record.workflow,
                    _raw(source.nik_raw) if source else "",
                    record.nik,
                    record.source_sheet,
                    record.source_file,
                    source.relative_path if source else "",
                    record.source_row,
                    _raw(source.date_in_raw) if source else "",
                    format_date(record.date_in),
                    _raw(source.time_in_raw) if source else "",
                    format_time(record.time_in),
                    _raw(source.date_out_raw) if source else "",
                    format_date(record.date_out),
                    _raw(source.time_out_raw) if source else "",
                    format_time(record.time_out),
                    len(record.changes),
                    "; ".join(change.change_code for change in record.changes),
                    record.final_status,
                    context.record_txt_assignment.get(record.record_id, ""),
                ],
                number_formats={3: "@", 4: "@"},
                status_columns=(19,),
            )
        _finish_sheet(sheet, len(headers), len(context.analysis.changed_records) + 1)

    def _anomaly(self, sheet: Any, context: "_ReportContext") -> None:
        headers = [
            "No",
            "Record_ID",
            "Workflow_Raw",
            "NIK_Raw",
            "Date_In_Raw",
            "Time_In_Raw",
            "Date_Out_Raw",
            "Time_Out_Raw",
            "Source_Sheet",
            "Source_File",
            "Relative_Path",
            "Source_Row",
            "Source_Record_No",
            "Source_Status",
            "Source_Reason",
            "Raw_Value",
            "Anomaly_Code",
            "Reason",
            "Action",
        ]
        _prepare_sheet(sheet, headers, data_sheet=True)
        _append_header(sheet, headers)
        for number, anomaly in enumerate(context.analysis.anomalies, 1):
            source = anomaly.source
            _append_data_row(
                sheet,
                [
                    number,
                    anomaly.record_id,
                    _raw(source.workflow_raw),
                    _raw(source.nik_raw),
                    _raw(source.date_in_raw),
                    _raw(source.time_in_raw),
                    _raw(source.date_out_raw),
                    _raw(source.time_out_raw),
                    source.source_sheet,
                    source.source_file,
                    source.relative_path,
                    source.source_row,
                    source.source_record_no,
                    source.source_status,
                    source.source_reason,
                    source.raw_value,
                    anomaly.anomaly_code,
                    anomaly.reason,
                    _anomaly_action(anomaly.anomaly_code),
                ],
                status_columns=(16,),
            )
        _finish_sheet(sheet, len(headers), len(context.analysis.anomalies) + 1)

    def _change_log(self, sheet: Any, context: "_ReportContext") -> None:
        headers = [
            "No",
            "Record_ID",
            "Workflow",
            "NIK",
            "Source_Sheet",
            "Source_File",
            "Relative_Path",
            "Source_Row",
            "Field_Name",
            "Original_Value",
            "Final_Value",
            "Change_Code",
            "Reason",
        ]
        _prepare_sheet(sheet, headers, data_sheet=True)
        _append_header(sheet, headers)
        for number, change in enumerate(context.analysis.change_log, 1):
            record = context.final_by_id.get(change.record_id)
            source = context.source_by_id.get(change.record_id)
            _append_data_row(
                sheet,
                [
                    number,
                    change.record_id,
                    record.workflow if record else "",
                    record.nik if record else "",
                    source.source_sheet if source else "",
                    source.source_file if source else "",
                    source.relative_path if source else "",
                    source.source_row if source else "",
                    change.field_name,
                    change.original_value,
                    change.final_value,
                    change.change_code,
                    change.reason,
                ],
                number_formats={3: "@"},
                status_columns=(11,),
            )
        _finish_sheet(sheet, len(headers), len(context.analysis.change_log) + 1)

    def _source_inventory(self, sheet: Any, context: "_ReportContext") -> None:
        headers = [
            "No",
            "Source_Workbook",
            "File_Size_Bytes",
            "Last_Modified",
            "Sheet_Name",
            "Rows_Read",
            "Final_Records",
            "Changed_Records",
            "Anomaly_Records",
            "Processed_At",
            "Status",
        ]
        _prepare_sheet(sheet, headers, data_sheet=True)
        _append_header(sheet, headers)
        source_path = context.analysis_request.source_report
        stat = source_path.stat()
        final_by_sheet = Counter(record.source_sheet for record in context.analysis.final_records)
        changed_by_sheet = Counter(record.source_sheet for record in context.analysis.changed_records)
        anomaly_by_sheet = Counter(item.source.source_sheet for item in context.analysis.anomalies)
        for number, sheet_name in enumerate((VALID_RECORDS_SHEET, INVALID_RECORDS_SHEET), 1):
            _append_data_row(
                sheet,
                [
                    number,
                    str(source_path),
                    stat.st_size,
                    datetime.fromtimestamp(stat.st_mtime),
                    sheet_name,
                    context.analysis.source_counts.get(sheet_name, 0),
                    final_by_sheet[sheet_name],
                    changed_by_sheet[sheet_name],
                    anomaly_by_sheet[sheet_name],
                    context.generated_at,
                    "PROCESSED",
                ],
                number_formats={3: "yyyy-mm-dd hh:mm:ss", 9: "yyyy-mm-dd hh:mm:ss"},
            )
        _finish_sheet(sheet, len(headers), 3)


class _ReportContext:
    def __init__(
        self,
        *,
        target: Path,
        job_id: str,
        status: str,
        request: AttDataRepairJobRequest,
        paths: JobPaths,
        analysis: AttDataRepairAnalysisResult,
        txt_artifacts: tuple[TxtArtifact, ...],
        record_txt_assignment: dict[str, str],
        generated_at: datetime,
    ) -> None:
        self.target = target
        self.job_id = job_id
        self.status = status
        self.request = request
        self.analysis_request = request.analysis_request
        self.paths = paths
        self.analysis = analysis
        self.txt_artifacts = txt_artifacts
        self.record_txt_assignment = record_txt_assignment
        self.generated_at = generated_at
        self.source_by_id = {record.record_id: record for record in analysis.source_records}
        self.final_by_id = {record.record_id: record for record in analysis.final_records}
        self.resolved_nik_by_id = {
            record.record_id: normalize_nik(record.nik_raw)
            for record in analysis.source_records
        }
        for change in analysis.change_log:
            if change.field_name == "NIK":
                self.resolved_nik_by_id[change.record_id] = change.final_value

    @property
    def txt_unique_code(self) -> int | None:
        return self.txt_artifacts[0].unique_code if self.txt_artifacts else None

    def overall_metrics(self) -> tuple[tuple[str, int], ...]:
        return (
            ("Source Records Total", self.analysis.source_counts.get("total", 0)),
            (
                "Source Valid_Records",
                self.analysis.source_counts.get(VALID_RECORDS_SHEET, 0),
            ),
            (
                "Source Invalid_Records",
                self.analysis.source_counts.get(INVALID_RECORDS_SHEET, 0),
            ),
            ("Final Records", len(self.analysis.final_records)),
            (
                "Unchanged Records",
                sum(
                    record.final_status == FinalStatus.VALID_UNCHANGED
                    for record in self.analysis.final_records
                ),
            ),
            (
                "Repaired Records",
                sum(
                    record.final_status == FinalStatus.REPAIRED
                    for record in self.analysis.final_records
                ),
            ),
            ("Changed Records", len(self.analysis.changed_records)),
            ("Anomaly Records", len(self.analysis.anomalies)),
            ("TXT Files", len(self.txt_artifacts)),
            ("TXT Rows", sum(artifact.row_count for artifact in self.txt_artifacts)),
            ("Excel Report Files", 1),
        )

    def workflow_metrics(self) -> tuple[tuple[str, int, int, int], ...]:
        return tuple(
            (metric, ho, branch, ho + branch)
            for metric, ho, branch in (
                (
                    "Final Records",
                    self._workflow_final(WORKFLOW_HO),
                    self._workflow_final(WORKFLOW_BRANCH),
                ),
                (
                    "Unchanged Records",
                    self._workflow_status(WORKFLOW_HO, FinalStatus.VALID_UNCHANGED),
                    self._workflow_status(WORKFLOW_BRANCH, FinalStatus.VALID_UNCHANGED),
                ),
                (
                    "Repaired Records",
                    self._workflow_status(WORKFLOW_HO, FinalStatus.REPAIRED),
                    self._workflow_status(WORKFLOW_BRANCH, FinalStatus.REPAIRED),
                ),
                (
                    "Changed Records",
                    self._workflow_changed(WORKFLOW_HO),
                    self._workflow_changed(WORKFLOW_BRANCH),
                ),
                (
                    "TXT Files",
                    self._workflow_txt_files(WORKFLOW_HO),
                    self._workflow_txt_files(WORKFLOW_BRANCH),
                ),
                (
                    "TXT Rows",
                    self._workflow_txt_rows(WORKFLOW_HO),
                    self._workflow_txt_rows(WORKFLOW_BRANCH),
                ),
            )
        )

    @property
    def anomaly_counts(self) -> Counter[str]:
        return Counter(item.anomaly_code for item in self.analysis.anomalies)

    @property
    def change_counts(self) -> Counter[str]:
        return Counter(item.change_code for item in self.analysis.change_log)

    def _workflow_final(self, workflow: str) -> int:
        return sum(record.workflow == workflow for record in self.analysis.final_records)

    def _workflow_status(self, workflow: str, status: str) -> int:
        return sum(
            record.workflow == workflow and record.final_status == status
            for record in self.analysis.final_records
        )

    def _workflow_changed(self, workflow: str) -> int:
        return sum(record.workflow == workflow for record in self.analysis.changed_records)

    def _workflow_txt_files(self, workflow: str) -> int:
        return sum(artifact.workflow == workflow for artifact in self.txt_artifacts)

    def _workflow_txt_rows(self, workflow: str) -> int:
        return sum(
            artifact.row_count
            for artifact in self.txt_artifacts
            if artifact.workflow == workflow
        )


def _employee_groups(context: _ReportContext) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str], dict[str, Any]] = {}

    def group(workflow: str, nik: str) -> dict[str, Any]:
        key = (workflow, nik)
        return groups.setdefault(
            key,
            {
                "workflow": workflow,
                "nik": nik,
                "source": 0,
                "final": 0,
                "unchanged": 0,
                "repaired": 0,
                "anomaly": 0,
                "dates": [],
            },
        )

    for source in context.analysis.source_records:
        nik = context.resolved_nik_by_id.get(source.record_id, "")
        if not nik:
            continue
        workflow = normalize_workflow(source.workflow_raw) or "Invalid"
        item = group(workflow, nik)
        item["source"] += 1
        parsed = normalize_date(
            source.date_in_raw,
            context.analysis_request.period_start,
            context.analysis_request.period_end,
        )
        if parsed.value is not None:
            item["dates"].append(parsed.value)

    for record in context.analysis.final_records:
        item = group(record.workflow, record.nik)
        item["final"] += 1
        if record.final_status == FinalStatus.VALID_UNCHANGED:
            item["unchanged"] += 1
        if record.final_status == FinalStatus.REPAIRED:
            item["repaired"] += 1
        item["dates"].append(record.date_in)

    for anomaly in context.analysis.anomalies:
        nik = context.resolved_nik_by_id.get(
            anomaly.record_id,
            normalize_nik(anomaly.source.nik_raw),
        )
        if not nik:
            continue
        workflow = normalize_workflow(anomaly.source.workflow_raw) or "Invalid"
        item = group(workflow, nik)
        item["anomaly"] += 1

    rows = []
    for item in groups.values():
        if item["final"] and item["anomaly"]:
            status = "PARTIAL"
        elif item["final"]:
            status = "INCLUDED"
        else:
            status = "EXCLUDED"
        dates = item["dates"]
        rows.append(
            {
                **item,
                "first_date": min(dates) if dates else "",
                "last_date": max(dates) if dates else "",
                "txt_status": status,
            }
        )
    rank = {WORKFLOW_HO: 0, WORKFLOW_BRANCH: 1, "Invalid": 2}
    return sorted(rows, key=lambda item: (rank.get(item["workflow"], 99), item["nik"]))


def _active_configuration_rows(
    context: _ReportContext,
) -> tuple[tuple[str, object], ...]:
    request = context.analysis_request
    return (
        ("minimum_duration_minutes", request.minimum_duration_minutes),
        ("weekday_default_in", format_time(request.weekday_default_in)),
        ("weekday_default_out", format_time(request.weekday_default_out)),
        ("saturday_default_in", format_time(request.saturday_default_in)),
        ("saturday_default_out", format_time(request.saturday_default_out)),
        (
            "saturday_missing_out_default",
            format_time(request.saturday_missing_out_default),
        ),
        (
            "sunday_invalid_default_in",
            format_time(request.sunday_invalid_default_in),
        ),
        (
            "sunday_invalid_default_out",
            format_time(request.sunday_invalid_default_out),
        ),
        (
            "midnight_time_out_default",
            format_time(request.midnight_time_out_default),
        ),
        ("txt_max_rows", context.request.txt_max_rows_per_file),
        ("generate_txt", str(context.request.generate_txt).upper()),
        (
            "generate_excel_report",
            str(context.request.generate_excel_report).upper(),
        ),
    )


def _append_section(sheet: Any, title: str) -> None:
    _append_row(sheet, [""])
    cell = WriteOnlyCell(sheet, value=title)
    cell.fill = SECTION_FILL
    cell.font = Font(bold=True)
    _append_row(sheet, [cell])


def _prepare_sheet(
    sheet: Any,
    headers: list[str],
    *,
    data_sheet: bool = False,
) -> None:
    if data_sheet:
        sheet.freeze_panes = "A2"
    for column, header in enumerate(headers, 1):
        width = min(max(len(str(header)) + 2, 10), 32)
        sheet.column_dimensions[get_column_letter(column)].width = width


def _append_header(sheet: Any, values: list[str]) -> None:
    cells: list[WriteOnlyCell] = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cells.append(cell)
    _append_row(sheet, cells)


def _append_data_row(
    sheet: Any,
    values: Any,
    *,
    number_formats: dict[int, str] | None = None,
    status_columns: tuple[int, ...] = (),
) -> None:
    row = list(values)
    formats = number_formats or {}
    for index in sorted(set(formats) | set(status_columns)):
        cell = WriteOnlyCell(sheet, value=row[index])
        number_format = formats.get(index)
        if number_format:
            cell.number_format = number_format
        if index in status_columns:
            fill = _status_fill(row[index])
            if fill is not None:
                cell.fill = fill
        row[index] = cell
    _append_row(sheet, row)


def _finish_sheet(sheet: Any, column_count: int, row_count: int) -> None:
    if column_count and row_count:
        last_column = get_column_letter(column_count)
        sheet.auto_filter.ref = f"A1:{last_column}{row_count}"


def _status_fill(value: Any) -> PatternFill | None:
    text = str(value or "")
    fill = STATUS_FILLS.get(text)
    if fill is None and (
        text in _ANOMALY_DESCRIPTIONS
        or text == AnomalyCode.INVALID_NIK
        or text == AnomalyCode.INVALID_DATE
        or text == AnomalyCode.OUTSIDE_REPORT_PERIOD
        or text == AnomalyCode.INVALID_WORKFLOW
    ):
        fill = STATUS_FILLS["FAILED"]
    if fill is None and text in _CHANGE_DESCRIPTIONS:
        fill = STATUS_FILLS["REPAIRED"]
    return fill


def _append_row(sheet: Any, values: Any) -> None:
    sheet.append(values)
    sheet._oas_row_count = int(getattr(sheet, "_oas_row_count", 0)) + 1


def _file_size(path: Path) -> int:
    return path.stat().st_size if path.exists() else 0


def _relative_to_job(path: Path, job_folder: Path | None) -> str:
    if job_folder is None:
        return path.name
    try:
        return str(path.relative_to(job_folder))
    except ValueError:
        return str(path)


def _raw(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _anomaly_action(code: str) -> str:
    return {
        str(AnomalyCode.INVALID_NIK): (
            "Record dikeluarkan dari TXT karena NIK tidak valid."
        ),
        str(AnomalyCode.INVALID_DATE): (
            "Record dikeluarkan dari TXT karena tanggal tidak valid."
        ),
        str(AnomalyCode.OUTSIDE_REPORT_PERIOD): (
            "Record dikeluarkan dari TXT karena tanggal berada di luar periode."
        ),
        str(AnomalyCode.INVALID_WORKFLOW): (
            "Record dikeluarkan dari TXT karena Workflow tidak valid."
        ),
        str(AnomalyCode.DUPLICATE_RECORD): (
            "Record duplikat dikeluarkan; record lengkap pertama dipertahankan."
        ),
        str(AnomalyCode.SUNDAY_NOT_MONTH_END): (
            "Record Minggu dikeluarkan karena bukan hari terakhir bulan."
        ),
    }.get(code, "Record dikeluarkan dari TXT dan perlu review manual.")


def _status_reference_rows() -> tuple[tuple[str, str, str, str, str, str], ...]:
    rows = [
        (
            str(FinalStatus.VALID_UNCHANGED),
            "Final Status",
            "Record valid dan tidak membutuhkan perbaikan.",
            "Masuk ke Final Records dan TXT.",
            "Ya",
            "Hijau",
        ),
        (
            str(FinalStatus.REPAIRED),
            "Final Status",
            "Record berhasil diperbaiki otomatis dengan rule Att Data Repair.",
            "Masuk ke Final Records dan TXT; detail perubahan ada di Change Log.",
            "Ya",
            "Kuning",
        ),
    ]
    for code, description in _CHANGE_DESCRIPTIONS.items():
        rows.append(
            (
                code,
                "Change Code",
                description,
                "Perubahan dicatat; record tetap masuk TXT bila final valid.",
                "Ya",
                "Kuning/Oranye",
            )
        )
    for code, description in _ANOMALY_DESCRIPTIONS.items():
        rows.append(
            (
                code,
                "Anomaly Code",
                description,
                _anomaly_action(code),
                "Tidak",
                "Merah",
            )
        )
    for code, description in _JOB_DESCRIPTIONS.items():
        rows.append(
            (
                code,
                "Job Status",
                description,
                "Gunakan Process.log, summary.json, dan sheet audit untuk review.",
                "Sesuai status record",
                "Hijau/Kuning/Merah",
            )
        )
    return tuple(rows)


_CHANGE_DESCRIPTIONS = {
    str(ChangeCode.NIK_LEADING_TWO_RESTORED): (
        "Angka 2 yang hilang pada awal NIK panjang ditambahkan kembali."
    ),
    str(ChangeCode.NIK_MISSING_ZERO_RESTORED): (
        "Satu angka nol yang hilang setelah angka 2 pada NIK panjang ditambahkan kembali."
    ),
    str(ChangeCode.NIK_EXTRA_LEADING_TWO_REMOVED): (
        "Angka 2 tambahan pada pola NIK 9 digit dihapus."
    ),
    str(ChangeCode.NIK_LEADING_TWO_REPLACED): (
        "Nol pertama pada NIK 11 digit diganti menjadi angka 2."
    ),
    str(ChangeCode.DATE_YEAR_ALIGNED_TO_CURRENT_YEAR): (
        "Tahun Date_In/Date_Out diselaraskan ke tahun berjalan."
    ),
    str(ChangeCode.DATE_OUT_ALIGNED_TO_DATE_IN): (
        "Date_Out disamakan dengan Date_In karena night shift tidak didukung."
    ),
    str(ChangeCode.TIME_FORMAT_NORMALIZED): (
        "Format jam dibersihkan menjadi HH:MM, misalnya 7;20 menjadi 07:20."
    ),
    str(ChangeCode.MIDNIGHT_TIME_IN_DEFAULTED): (
        "Time_In 00:00 diganti dengan jam masuk default hari yang sudah ada."
    ),
    str(ChangeCode.MIDNIGHT_TIME_OUT_DEFAULTED): (
        "Time_Out 00:00 diganti dengan konfigurasi midnight_time_out_default."
    ),
    str(ChangeCode.SATURDAY_MISSING_OUT_DEFAULTED): (
        "Time_Out Sabtu kosong diisi dari saturday_missing_out_default."
    ),
    str(ChangeCode.TIME_IN_OUT_SWAPPED): (
        "Time_In dan Time_Out ditukar karena urutan jam terbalik."
    ),
    str(ChangeCode.PARTIAL_TIME_REPAIRED): (
        "Salah satu jam rusak lalu diperbaiki memakai jam valid dan durasi minimum."
    ),
    str(ChangeCode.DEFAULT_TIME_APPLIED): (
        "Jam default hari digunakan karena jam sumber tidak dapat dipakai aman."
    ),
    str(ChangeCode.MINIMUM_DURATION_APPLIED): (
        "Durasi dinaikkan sampai minimal 61 menit."
    ),
}
_ANOMALY_DESCRIPTIONS = {
    str(AnomalyCode.INVALID_NIK): (
        "NIK kosong, error spreadsheet, bukan digit ASCII, atau panjang/polanya tidak valid."
    ),
    str(AnomalyCode.INVALID_DATE): "Date_In kosong, error, atau tidak dapat dipastikan.",
    str(AnomalyCode.OUTSIDE_REPORT_PERIOD): "Date_In berada di luar periode report.",
    str(AnomalyCode.INVALID_WORKFLOW): "Workflow kosong atau bukan HO/Branch.",
    str(AnomalyCode.DUPLICATE_RECORD): (
        "Record sumber lengkap memiliki NIK, tanggal, Time_In, dan Time_Out yang sama."
    ),
    str(AnomalyCode.SUNDAY_NOT_MONTH_END): (
        "Tanggal jatuh pada hari Minggu dan bukan hari terakhir bulan."
    ),
}
_JOB_DESCRIPTIONS = {
    "SUCCESS": "Job selesai tanpa anomaly dan artifact berhasil dibuat.",
    "PARTIAL_SUCCESS": "Job selesai dengan sebagian record dikeluarkan sebagai anomaly.",
    "NO_VALID_RECORDS": "Tidak ada record final; report audit tetap dibuat.",
    "FAILED": "Output pipeline gagal dan perlu diperiksa dari log/error.",
}
_GUIDE_PRINCIPLES = (
    "Sumber input adalah Excel Report Attachment Consolidation.",
    "Sheet sumber adalah Valid_Records dan Invalid_Records.",
    "Seluruh record dari kedua sheet dianalisis ulang.",
    "Record dari Invalid_Records yang berhasil diperbaiki dapat masuk TXT.",
    "Record dari Valid_Records tetap diperiksa ulang.",
    "Workbook sumber tidak pernah diubah.",
    "Tidak ada salinan sumber pada folder output.",
    "Tidak ada folder Original pada output Att Data Repair.",
    "NIK adalah identitas utama.",
    "Nama karyawan tidak wajib dan tidak tersedia pada source report.",
    "Output TXT dipisahkan menjadi HO dan Branch.",
    "Maksimum 10.000 baris per TXT.",
    "Duplikat lengkap dipindahkan ke Anomaly; record pertama dipertahankan.",
    "Record dengan tanggal atau jam sumber kosong tidak ikut pencocokan duplikat.",
    "Tidak ada approval PIC per record.",
    "Semua perubahan tercatat pada report.",
)
_GUIDE_NIK = (
    "NIK wajib tersedia.",
    "Leading zero dipertahankan.",
    "NIK wajib berisi digit ASCII dengan panjang tepat 9, 10, atau 11 digit.",
    "NIK 9 digit dan NIK 10 digit resmi dipertahankan.",
    "NIK 10 digit dengan 5-8 nol awal ditambah angka 2 di depan.",
    "NIK 10 digit yang diawali 2 dan 4-7 nol diperbaiki dengan satu nol tambahan.",
    "NIK 10 digit yang diawali 2 dan tepat 3 nol dikembalikan ke NIK 9 digit.",
    "NIK 11 digit valid diawali 2 dan diikuti 5-8 nol.",
    "NIK 11 digit dengan 6-9 nol awal diperbaiki dengan mengganti nol pertama menjadi 2.",
    "NIK kosong, spreadsheet error, atau pola 11 digit meragukan masuk Anomaly.",
    "Perbaikan NIK tidak menggunakan master data karyawan.",
    "Record dengan NIK invalid tidak masuk TXT.",
)
_GUIDE_WORKFLOW = (
    "Workflow valid adalah HO dan Branch.",
    "Nilai BRANCH dapat dinormalisasi menjadi Branch.",
    "Workflow kosong atau invalid masuk Anomaly.",
    "Workflow invalid tidak masuk TXT.",
)
_GUIDE_DATE = (
    "Output tanggal memakai MM/DD/YYYY.",
    "Date_In menjadi tanggal utama.",
    "Date_Out selalu mengikuti Date_In.",
    "Jika Date_Out berbeda, Date_Out disamakan ke Date_In.",
    "Perubahan dicatat sebagai DATE_OUT_ALIGNED_TO_DATE_IN.",
    "Tanggal menggunakan periode report sebagai guard.",
    "Tanggal di luar periode masuk Anomaly.",
    "Tanggal yang tidak dapat ditentukan masuk Anomaly.",
    "Tahun tanggal yang berbeda diselaraskan ke tahun berjalan.",
    "Hari Minggu masuk Anomaly kecuali merupakan hari terakhir bulan.",
    "OAS-K tidak mendukung night shift.",
)
_GUIDE_TIME = (
    "Format seperti 7;20 dapat dinormalisasi menjadi 07:20.",
    "Format seperti 7.20 dapat dinormalisasi menjadi 07:20.",
    "Jika In lebih besar dari Out dan keduanya valid, keduanya ditukar.",
    "Jika hanya satu jam rusak, pertahankan jam valid bila hasil tetap logis.",
    "Durasi minimum adalah 61 menit.",
    "Aturan durasi minimum berlaku untuk semua hari.",
    "Jika perhitungan melewati 23:59, gunakan default hari.",
    "Jika kedua jam rusak, langsung gunakan default hari.",
    "Time_In 00:00 memakai jam masuk default hari yang sudah ada.",
    "Time_Out 00:00 memakai midnight_time_out_default.",
    "Tidak memerlukan approval PIC.",
)
_GUIDE_DEFAULTS = (
    "Senin-Jumat: 09:30-17:00.",
    "Sabtu: pasangan default tetap 09:30-12:05.",
    "Sabtu dengan Time_Out kosong: gunakan saturday_missing_out_default.",
    "Minggu hanya diproses jika merupakan hari terakhir bulan.",
)
_GUIDE_NAMING = (
    "Job Folder: YYYY-MM-DD_XX.",
    "TXT: Att_Data_Repair_[Workflow]-[NNN]_[RRRR].txt.",
    "Excel: Att_Data_Repair_Report_[Job_ID].xlsx.",
    "NNN adalah urutan split file per workflow.",
    "RRRR adalah kode unik acak empat digit.",
    "Kode RRRR sama untuk seluruh TXT dalam satu job.",
    "Tanggal tidak dimasukkan ke nama TXT.",
    "HO dan Branch mempunyai sequence masing-masing.",
)
